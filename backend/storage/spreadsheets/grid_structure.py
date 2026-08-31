from __future__ import annotations

from collections import deque
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any, Dict, List, Sequence, Set, Tuple

from openpyxl.cell.cell import MergedCell

from .cell_visibility import WorksheetVisibility
from .table_geometry import CellBounds

Coordinate = Tuple[int, int]


@dataclass(frozen=True)
class HeaderNode:
    name: str
    col_start: int
    col_end: int
    row_start: int
    row_end: int
    children: Tuple["HeaderNode", ...] = ()


def _has_value(value: Any) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _direct_occupied_cells(
    worksheet,
    visibility: WorksheetVisibility,
    row_limit: int,
    column_limit: int,
) -> Set[Coordinate]:
    occupied: Set[Coordinate] = set()
    for row in range(1, row_limit + 1):
        if visibility.row_hidden(row):
            continue
        for column in range(1, column_limit + 1):
            if visibility.column_hidden(column):
                continue
            cell = worksheet.cell(row=row, column=column)
            if cell.data_type != "e" and _has_value(cell.value):
                occupied.add((row, column))
    return occupied


def _merged_occupied_cells(
    worksheet,
    visibility: WorksheetVisibility,
    row_limit: int,
    column_limit: int,
) -> Set[Coordinate]:
    occupied: Set[Coordinate] = set()
    for merged_range in worksheet.merged_cells.ranges:
        if merged_range.min_row > row_limit or merged_range.min_col > column_limit:
            continue
        if not visibility.cell_visible(merged_range.min_row, merged_range.min_col):
            continue
        anchor = worksheet.cell(merged_range.min_row, merged_range.min_col)
        if anchor.data_type == "e" or not _has_value(anchor.value):
            continue
        for row in range(merged_range.min_row, min(merged_range.max_row, row_limit) + 1):
            if visibility.row_hidden(row):
                continue
            for column in range(
                merged_range.min_col,
                min(merged_range.max_col, column_limit) + 1,
            ):
                if not visibility.column_hidden(column):
                    occupied.add((row, column))
    return occupied


def occupied_cells(
    worksheet,
    max_rows: int,
    max_columns: int,
) -> Set[Coordinate]:
    """Build the non-empty grid used by the PDF's four-neighbour BFS."""

    row_limit = min(worksheet.max_row or 1, max_rows)
    column_limit = min(worksheet.max_column or 1, max_columns)
    visibility = WorksheetVisibility.from_worksheet(worksheet)
    return _direct_occupied_cells(
        worksheet,
        visibility,
        row_limit,
        column_limit,
    ) | _merged_occupied_cells(
        worksheet,
        visibility,
        row_limit,
        column_limit,
    )


def connected_components(occupied: Set[Coordinate]) -> List[Set[Coordinate]]:
    """Return four-directionally connected non-empty cell groups."""

    unvisited = set(occupied)
    groups: List[Set[Coordinate]] = []
    while unvisited:
        start = min(unvisited)
        unvisited.remove(start)
        queue = deque([start])
        group = {start}
        while queue:
            row, column = queue.popleft()
            for neighbour in (
                (row - 1, column),
                (row + 1, column),
                (row, column - 1),
                (row, column + 1),
            ):
                if neighbour in unvisited:
                    unvisited.remove(neighbour)
                    group.add(neighbour)
                    queue.append(neighbour)
        groups.append(group)
    return groups


def component_bounds(component: Set[Coordinate]) -> CellBounds:
    rows = [coordinate[0] for coordinate in component]
    columns = [coordinate[1] for coordinate in component]
    return CellBounds(min(rows), max(rows), min(columns), max(columns))


def _axis_gap(first_min: int, first_max: int, second_min: int, second_max: int) -> int:
    if first_max < second_min:
        return second_min - first_max - 1
    if second_max < first_min:
        return first_min - second_max - 1
    return 0


def _mergeable(first: CellBounds, second: CellBounds, gap: int) -> bool:
    return (
        _axis_gap(first.min_row, first.max_row, second.min_row, second.max_row) <= gap
        and _axis_gap(
            first.min_column,
            first.max_column,
            second.min_column,
            second.max_column,
        )
        <= gap
    )


def _union(first: CellBounds, second: CellBounds) -> CellBounds:
    return CellBounds(
        min(first.min_row, second.min_row),
        max(first.max_row, second.max_row),
        min(first.min_column, second.min_column),
        max(first.max_column, second.max_column),
    )


def merge_adjacent_bounds(
    bounds: Sequence[CellBounds],
    gap: int,
) -> List[CellBounds]:
    """Merge adjacent boxes in all eight directions until convergence."""

    pending = list(bounds)
    changed = True
    while changed:
        changed = False
        merged: List[CellBounds] = []
        while pending:
            current = pending.pop(0)
            index = 0
            while index < len(pending):
                if _mergeable(current, pending[index], gap):
                    current = _union(current, pending.pop(index))
                    changed = True
                    index = 0
                else:
                    index += 1
            merged.append(current)
        pending = merged
    return sorted(
        pending,
        key=lambda item: (item.min_row, item.min_column, item.max_row, item.max_column),
    )


def detect_table_bounds(
    worksheet,
    max_rows: int,
    max_columns: int,
    merge_gap: int,
    min_non_empty_cells: int,
) -> List[CellBounds]:
    occupied = occupied_cells(worksheet, max_rows, max_columns)
    components = connected_components(occupied)
    merged = merge_adjacent_bounds(
        [component_bounds(component) for component in components],
        merge_gap,
    )
    return [
        bounds
        for bounds in merged
        if sum(
            1
            for row, column in occupied
            if bounds.min_row <= row <= bounds.max_row
            and bounds.min_column <= column <= bounds.max_column
        )
        >= min_non_empty_cells
    ]


def _display_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _merged_coordinates(
    worksheet,
    visibility: WorksheetVisibility,
) -> Set[Coordinate]:
    coordinates: Set[Coordinate] = set()
    for merged_range in worksheet.merged_cells.ranges:
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            if visibility.row_hidden(row):
                continue
            for column in range(merged_range.min_col, merged_range.max_col + 1):
                if visibility.column_hidden(column):
                    continue
                coordinates.add((row, column))
    return coordinates


def header_candidate_rows(
    value_worksheet,
    formula_worksheet,
    bounds: CellBounds,
    row_count: int,
) -> List[Dict[str, Any]]:
    """Serialize at most the top N rows with the cell metadata from the PDF."""

    visibility = WorksheetVisibility.from_worksheet(formula_worksheet)
    merged = _merged_coordinates(formula_worksheet, visibility)
    rows: List[Dict[str, Any]] = []
    for row in range(bounds.min_row, bounds.max_row + 1):
        if visibility.row_hidden(row):
            continue
        cells: List[Dict[str, Any]] = []
        for column in range(bounds.min_column, bounds.max_column + 1):
            if visibility.column_hidden(column):
                continue
            raw_cell = formula_worksheet.cell(row=row, column=column)
            value_cell = value_worksheet.cell(row=row, column=column)
            value = value_cell.value
            if raw_cell.data_type == "e" or value_cell.data_type == "e":
                value = None
                cell_type = "empty"
            elif raw_cell.data_type == "f":
                value = value if value is not None else raw_cell.value
                cell_type = "formula"
            elif value is None or isinstance(raw_cell, MergedCell):
                cell_type = "empty"
            elif isinstance(value, (int, float)) and not isinstance(value, bool):
                cell_type = "number"
            else:
                cell_type = "text"
            cells.append(
                {
                    "row": row,
                    "col": column,
                    "value": _display_value(value),
                    "is_merged": (row, column) in merged,
                    "cell_type": cell_type,
                }
            )
        rows.append({"row": row, "cells": cells})
        if len(rows) >= row_count:
            break
    return rows


def validate_title_end(
    worksheet,
    bounds: CellBounds,
    title_row_end: int | None,
) -> int | None:
    """Apply the document's post-check: a title row cannot contain 2+ values."""

    if title_row_end is None:
        return None
    visibility = WorksheetVisibility.from_worksheet(worksheet)
    valid_end = min(title_row_end, bounds.max_row)
    last_visible_row: int | None = None
    previous_visible_row: int | None = None
    for row in range(bounds.min_row, valid_end + 1):
        if visibility.row_hidden(row):
            continue
        last_visible_row = row
        non_empty = 0
        anchors: Set[Coordinate] = set()
        for column in range(bounds.min_column, bounds.max_column + 1):
            if visibility.column_hidden(column):
                continue
            cell = worksheet.cell(row=row, column=column)
            if isinstance(cell, MergedCell):
                continue
            if _has_value(cell.value) and (row, column) not in anchors:
                anchors.add((row, column))
                non_empty += 1
        if non_empty >= 2:
            return previous_visible_row
        previous_visible_row = last_visible_row
    return last_visible_row


def _merge_origin_map(worksheet) -> Dict[Coordinate, Tuple[int, int, int, int]]:
    origins: Dict[Coordinate, Tuple[int, int, int, int]] = {}
    for merged_range in worksheet.merged_cells.ranges:
        origin = (merged_range.min_row, merged_range.min_col)
        origins[origin] = (
            merged_range.min_row,
            merged_range.max_row,
            merged_range.min_col,
            merged_range.max_col,
        )
    return origins


def _header_node(
    worksheet,
    visibility: WorksheetVisibility,
    merge_origins: Dict[Coordinate, Tuple[int, int, int, int]],
    bounds: CellBounds,
    header_end_row: int,
    data_start_column: int,
    row: int,
    column: int,
) -> Dict[str, Any] | None:
    cell = worksheet.cell(row=row, column=column)
    if isinstance(cell, MergedCell) or not _has_value(cell.value):
        return None
    merged = merge_origins.get((row, column), (row, row, column, column))
    visible_columns = [
        candidate
        for candidate in range(
            max(data_start_column, merged[2]),
            min(bounds.max_column, merged[3]) + 1,
        )
        if not visibility.column_hidden(candidate)
    ]
    visible_rows = [
        candidate
        for candidate in range(row, min(header_end_row, merged[1]) + 1)
        if not visibility.row_hidden(candidate)
    ]
    if not visible_columns or not visible_rows:
        return None
    return {
        "name": str(cell.value).strip(),
        "col_start": visible_columns[0],
        "col_end": visible_columns[-1],
        "row_start": row,
        "row_end": visible_rows[-1],
        "children": [],
        "parent": None,
    }


def _attach_header_parents(nodes: List[Dict[str, Any]]) -> None:
    nodes.sort(key=lambda node: (node["row_start"], node["col_start"], node["col_end"]))
    for child in nodes:
        parents = [
            candidate
            for candidate in nodes
            if candidate["row_start"] < child["row_start"]
            and candidate["col_start"] <= child["col_start"]
            and candidate["col_end"] >= child["col_end"]
        ]
        if not parents:
            continue
        parent = max(
            parents,
            key=lambda candidate: (
                candidate["row_start"],
                -(candidate["col_end"] - candidate["col_start"]),
            ),
        )
        child["parent"] = parent
        parent["children"].append(child)


def _serialize_header_node(node: Dict[str, Any]) -> Dict[str, Any]:
    return {
        "name": node["name"],
        "col_start": node["col_start"],
        "col_end": node["col_end"],
        "row_start": node["row_start"],
        "row_end": node["row_end"],
        "children": [_serialize_header_node(child) for child in node["children"]],
    }


def build_column_header_tree(
    worksheet,
    bounds: CellBounds,
    header_start_row: int,
    header_end_row: int,
    data_start_column: int,
) -> List[Dict[str, Any]]:
    """Create a deterministic parent-child tree from row order and col spans."""

    if header_start_row > header_end_row or data_start_column > bounds.max_column:
        return []
    visibility = WorksheetVisibility.from_worksheet(worksheet)
    merge_origins = _merge_origin_map(worksheet)
    mutable_nodes: List[Dict[str, Any]] = []
    for row in range(header_start_row, header_end_row + 1):
        if visibility.row_hidden(row):
            continue
        for column in range(data_start_column, bounds.max_column + 1):
            if visibility.column_hidden(column):
                continue
            node = _header_node(
                worksheet,
                visibility,
                merge_origins,
                bounds,
                header_end_row,
                data_start_column,
                row,
                column,
            )
            if node is not None:
                mutable_nodes.append(node)

    _attach_header_parents(mutable_nodes)
    return [_serialize_header_node(node) for node in mutable_nodes if node["parent"] is None]
