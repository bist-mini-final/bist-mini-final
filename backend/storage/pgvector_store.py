"""Native psycopg storage adapter for PostgreSQL and pgvector."""

from __future__ import annotations

import logging
import re
import threading
from datetime import datetime, timezone
from numbers import Real
from typing import Any, Callable, Dict, List, Optional, Sequence, Tuple
from urllib.parse import urlparse
from uuid import UUID, uuid4, uuid5

import psycopg2.extras
from langchain_core.documents import Document

from backend.contracts.vector import PgVectorReplacePlan
from backend.core.settings import PGVECTOR_URL
from backend.providers.embeddings.ports import EmbeddingEncoder
from backend.storage.embedding_artifacts import EmbeddingArtifactVectors
from backend.storage.pgvector_binary_copy import copy_documents
from backend.storage.spreadsheets.langchain_document import (
    cell_items_to_langchain_documents,
    langchain_document_to_cell_item,
)
from modules.common.config import DEFAULT_EMBEDDING_DIMENSION, DEFAULT_EMBEDDING_MODEL

from .connection_pool import get_pooled_async_connection, get_pooled_raw_connection

logger = logging.getLogger(__name__)


class PgVectorStoreError(RuntimeError):
    """Raised when a pgvector database operation fails."""


PGVECTOR_INSERT_BATCH_SIZE = 1000
VECTOR_INDEX_STRATEGY = "binary_quantized_hnsw_exact_rerank"
VECTOR_PARTITION_STRATEGY = "collection_local_partial_indexes"
_CONCURRENT_OPTIMIZED_INDEX_NAMES = (
    "idx_langchain_pg_embedding_cell_id",
    "idx_langchain_pg_embedding_cell_coord_upper",
    "idx_langchain_pg_embedding_workbook_hash",
    "idx_langchain_pg_embedding_company_name",
    "idx_langchain_pg_embedding_sheet_name",
    "idx_langchain_pg_embedding_sheet_row",
    "idx_langchain_pg_embedding_collection_id",
    "idx_langchain_pg_embedding_document_fts",
)


def _is_numeric_vector_collection(candidate: Any) -> bool:
    """Return whether a candidate is a non-empty sequence of numeric vectors."""
    if isinstance(candidate, (dict, str, bytes)) or not hasattr(candidate, "__len__"):
        return False
    try:
        vectors = list(candidate)
    except TypeError:
        return False
    if not vectors:
        return False
    for vector in vectors:
        if isinstance(vector, (dict, str, bytes)) or not hasattr(vector, "__len__"):
            return False
        try:
            values = list(vector)
        except TypeError:
            return False
        if not values or not all(
            isinstance(value, Real) and not isinstance(value, bool) for value in values
        ):
            return False
    return True


def _escape_like_term(text: str) -> str:
    """Escape user-derived wildcard characters for a literal ILIKE substring."""
    return text.replace("!", "!!").replace("%", "!%").replace("_", "!_")


class PgVectorStore:
    """Manage pgvector collections through the shared psycopg connection pool."""

    @staticmethod
    def index_id(artifact_id: str) -> str:
        """Derive standard pgvector collection index identifier from an artifact ID."""
        return f"idx_{artifact_id}"

    def __init__(self, database_url: str = PGVECTOR_URL) -> None:
        self.database_url = database_url
        self._collection_uuid_cache: Dict[str, str] = {}
        self._collection_uuid_lock = threading.Lock()

    def _raw_connection(self) -> Any:
        raw_url = getattr(self, "database_url", PGVECTOR_URL).replace(
            "postgresql+psycopg://", "postgresql://"
        )
        return get_pooled_raw_connection(raw_url)

    def _read_connection(self) -> Any:
        """Borrow an autocommit connection so SELECTs avoid a rollback round trip."""
        connection = self._raw_connection()
        connection.autocommit = True
        return connection

    def _collection_uuid(self, collection_name: str) -> Optional[str]:
        """Resolve and cache a collection UUID to remove a lookup from hot searches."""
        cached = self._collection_uuid_cache.get(collection_name)
        if cached is not None:
            return cached
        connection = self._read_connection()
        try:
            with connection.cursor() as cursor:
                cursor.execute(
                    "SELECT uuid FROM langchain_pg_collection WHERE name = %s;",
                    (collection_name,),
                )
                row = cursor.fetchone()
        finally:
            connection.close()
        if row is None:
            return None
        # The canonical UUID representation is also safe to embed in a query
        # predicate when selecting a collection-local partial index.
        resolved = str(UUID(str(row[0])))
        with self._collection_uuid_lock:
            self._collection_uuid_cache[collection_name] = resolved
        return resolved

    async def _collection_uuid_async(self, collection_name: str) -> Optional[str]:
        """Resolve collection identity without blocking the workflow event loop."""
        cached = self._collection_uuid_cache.get(collection_name)
        if cached is not None:
            return cached
        raw_url = getattr(self, "database_url", PGVECTOR_URL).replace(
            "postgresql+psycopg://",
            "postgresql://",
        )
        async with get_pooled_async_connection(raw_url) as connection:
            async with connection.cursor() as cursor:
                await cursor.execute(
                    "SELECT uuid FROM langchain_pg_collection WHERE name = %s;",
                    (collection_name,),
                )
                row = await cursor.fetchone()
        if row is None:
            return None
        resolved = str(UUID(str(row[0])))
        with self._collection_uuid_lock:
            self._collection_uuid_cache[collection_name] = resolved
        return resolved

    def _create_collection(
        self,
        collection_name: str,
        metadata: Dict[str, Any],
    ) -> str:
        collection_uuid = str(uuid4())
        connection = self._raw_connection()
        try:
            with connection.cursor() as cursor:
                cursor.execute(
                    """
                    INSERT INTO langchain_pg_collection (uuid, name, cmetadata)
                    VALUES (%s, %s, %s)
                    RETURNING uuid;
                    """,
                    (
                        collection_uuid,
                        collection_name,
                        psycopg2.extras.Json(metadata),
                    ),
                )
                row = cursor.fetchone()
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()
        if row is None:
            raise PgVectorStoreError(
                f"pgvector 컬렉션 생성 결과가 비어 있습니다: {collection_name}"
            )
        resolved = str(UUID(str(row[0])))
        with self._collection_uuid_lock:
            self._collection_uuid_cache[collection_name] = resolved
        return resolved

    def _delete_collection(self, collection_name: str) -> None:
        connection = self._raw_connection()
        try:
            with connection.cursor() as cursor:
                cursor.execute(
                    "DELETE FROM langchain_pg_collection WHERE name = %s;",
                    (collection_name,),
                )
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()
        with self._collection_uuid_lock:
            self._collection_uuid_cache.pop(collection_name, None)

    @staticmethod
    def _collection_metadata(
        metadata: Optional[Dict[str, Any]],
        *,
        model_name: str,
        document_count: int,
    ) -> Dict[str, Any]:
        meta_dict = metadata or {}
        return {
            "file_name": meta_dict.get("file_name", ""),
            "workbook_hash": meta_dict.get("workbook_hash", ""),
            "model": meta_dict.get("model", model_name),
            "dimension": meta_dict.get("dimension", DEFAULT_EMBEDDING_DIMENSION),
            "document_count": document_count,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "pipeline": meta_dict.get("pipeline", "luna_vlm_structured"),
            "duration_seconds": meta_dict.get("duration_seconds"),
            "total_tokens": meta_dict.get("total_tokens"),
            "estimated_cost_usd": meta_dict.get("estimated_cost_usd"),
            "estimated_cost_krw": meta_dict.get("estimated_cost_krw"),
            "batch_size": meta_dict.get("batch_size"),
            "company_name": meta_dict.get("company_name", ""),
            "ticker": meta_dict.get("ticker", ""),
            "operation_id": meta_dict.get("operation_id"),
        }

    def prepare_collection_replace(
        self,
        *,
        index_id: str,
        operation_id: str,
        model_name: str,
        document_count: int,
        metadata: Optional[Dict[str, Any]] = None,
    ) -> PgVectorReplacePlan:
        """Return a retry-stable private collection for distributed shard COPY."""

        clean_meta = self._collection_metadata(
            {**(metadata or {}), "operation_id": operation_id},
            model_name=model_name,
            document_count=document_count,
        )
        dimension = int(clean_meta["dimension"])
        published_uuid = self._collection_uuid(index_id)
        if (
            published_uuid is not None
            and self.collection_document_count(published_uuid) == document_count
        ):
            return PgVectorReplacePlan(
                index_id=index_id,
                operation_id=operation_id,
                staging_name=index_id,
                staging_uuid=published_uuid,
                dimension=dimension,
                metadata=clean_meta,
                published=True,
            )

        staging_name = f"{index_id}__staging__{operation_id[:16]}"
        staging_uuid = self._collection_uuid(staging_name)
        if staging_uuid is None:
            try:
                staging_uuid = self._create_collection(staging_name, clean_meta)
            except Exception:
                # A concurrent coordinator for the same content may have won
                # the unique-name race. Reuse that deterministic staging row.
                staging_uuid = self._collection_uuid(staging_name)
                if staging_uuid is None:
                    raise
        return PgVectorReplacePlan(
            index_id=index_id,
            operation_id=operation_id,
            staging_name=staging_name,
            staging_uuid=staging_uuid,
            dimension=dimension,
            metadata=clean_meta,
        )

    def copy_prepared_collection_shard(
        self,
        plan: PgVectorReplacePlan,
        *,
        documents: Sequence[Document],
        vectors: Sequence[Sequence[float]],
        start_index: int,
    ) -> None:
        """Idempotently replace one disjoint row range in a private collection."""

        if plan.published:
            return
        if len(documents) != len(vectors):
            raise PgVectorStoreError("pgvector shard 문서와 벡터 개수가 다릅니다")
        document_ids = [
            str(uuid5(UUID(plan.staging_uuid), str(start_index + offset)))
            for offset in range(len(documents))
        ]
        connection = self._raw_connection()
        try:
            with connection.cursor() as cursor:
                cursor.execute(
                    "DELETE FROM langchain_pg_embedding WHERE id = ANY(%s::varchar[]);",
                    (document_ids,),
                )
            copy_documents(
                connection,
                collection_uuid=plan.staging_uuid,
                documents=documents,
                vectors=vectors,
                batch_size=max(1, len(documents)),
                document_ids=document_ids,
            )
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()

    def collection_document_count(self, collection_uuid: str) -> int:
        connection = self._read_connection()
        try:
            with connection.cursor() as cursor:
                cursor.execute(
                    "SELECT COUNT(*) FROM langchain_pg_embedding WHERE collection_id = %s;",
                    (collection_uuid,),
                )
                row = cursor.fetchone()
        finally:
            connection.close()
        return int(row[0]) if row else 0

    def publish_prepared_collection(self, plan: PgVectorReplacePlan) -> None:
        """Build one HNSW index and atomically expose a fully copied collection."""

        if plan.published:
            return
        actual_count = self.collection_document_count(plan.staging_uuid)
        expected_count = int(plan.metadata["document_count"])
        if actual_count != expected_count:
            raise PgVectorStoreError(
                "pgvector staging collection 문서 수가 일치하지 않습니다 "
                f"({actual_count}/{expected_count})"
            )
        self.ensure_collection_vector_index(plan.staging_name, plan.dimension)

        retired_name = f"{plan.index_id}__retired__{uuid4().hex}"
        retired_uuid: Optional[str] = None
        retired_dimension = plan.dimension
        connection = self._raw_connection()
        try:
            with connection.cursor() as cursor:
                cursor.execute(
                    "SELECT pg_advisory_xact_lock(hashtext(%s));",
                    (plan.index_id,),
                )
                cursor.execute(
                    """
                    UPDATE langchain_pg_collection
                    SET name = %s
                    WHERE name = %s
                    RETURNING uuid, cmetadata;
                    """,
                    (retired_name, plan.index_id),
                )
                retired = cursor.fetchone()
                if retired is not None:
                    retired_uuid = str(retired[0])
                    if isinstance(retired[1], dict):
                        retired_dimension = int(
                            retired[1].get("dimension") or retired_dimension
                        )
                cursor.execute(
                    """
                    UPDATE langchain_pg_collection
                    SET name = %s, cmetadata = %s
                    WHERE name = %s
                    RETURNING uuid;
                    """,
                    (
                        plan.index_id,
                        psycopg2.extras.Json(plan.metadata),
                        plan.staging_name,
                    ),
                )
                published = cursor.fetchone()
                if published is None:
                    # Another retry may have published the same deterministic
                    # staging collection while this coordinator was waiting.
                    cursor.execute(
                        "SELECT uuid FROM langchain_pg_collection WHERE name = %s;",
                        (plan.index_id,),
                    )
                    published = cursor.fetchone()
                if published is None:
                    raise PgVectorStoreError("완료된 pgvector staging collection이 없습니다")
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()

        published_uuid = str(UUID(str(published[0])))
        with self._collection_uuid_lock:
            self._collection_uuid_cache.pop(plan.staging_name, None)
            self._collection_uuid_cache[plan.index_id] = published_uuid

        if retired_uuid is not None and retired_uuid != published_uuid:
            try:
                self._delete_collection(retired_name)
                self._drop_collection_vector_index(retired_uuid, retired_dimension)
            except Exception:
                logger.warning(
                    "교체된 pgvector collection 정리 실패: %s",
                    retired_name,
                    exc_info=True,
                )
        self.ensure_optimized_indexes()

    @staticmethod
    def _collection_vector_index_name(collection_uuid: str, dimension: int) -> str:
        uuid_hex = UUID(collection_uuid).hex
        return f"idx_lc_hnsw_bq_c_{uuid_hex}_{dimension}"

    def ensure_collection_vector_index(
        self,
        collection_name: str,
        dimension: int,
    ) -> str:
        """Create a compact collection-local binary-quantized HNSW index."""
        collection_uuid = self._collection_uuid(collection_name)
        if collection_uuid is None:
            raise PgVectorStoreError(f"컬렉션별 HNSW 인덱스 대상이 없습니다: {collection_name}")
        if not 0 < dimension <= 64_000:
            raise PgVectorStoreError(f"HNSW 인덱스를 지원하지 않는 벡터 차원입니다: {dimension}")
        index_name = self._collection_vector_index_name(collection_uuid, dimension)
        connection = self._read_connection()
        try:
            with connection.cursor() as cursor:
                cursor.execute(
                    f"""
                    CREATE INDEX CONCURRENTLY IF NOT EXISTS "{index_name}"
                    ON langchain_pg_embedding
                    USING hnsw (
                        (binary_quantize(embedding)::bit({dimension}))
                        bit_hamming_ops
                    )
                    WHERE collection_id = '{collection_uuid}'::uuid
                      AND vector_dims(embedding) = {dimension};
                    """
                )
        finally:
            connection.close()
        return index_name

    def _drop_collection_vector_index(
        self,
        collection_uuid: str,
        dimension: int,
    ) -> None:
        index_name = self._collection_vector_index_name(collection_uuid, dimension)
        connection = self._read_connection()
        try:
            with connection.cursor() as cursor:
                cursor.execute(f'DROP INDEX CONCURRENTLY IF EXISTS "{index_name}";')
        finally:
            connection.close()

    def is_connected(self) -> bool:
        """Check if PostgreSQL + pgvector is reachable."""
        try:
            conn = self._raw_connection()
            try:
                with conn.cursor() as cur:
                    cur.execute("SELECT 1;")
                return True
            finally:
                conn.close()
        except Exception:
            return False

    @staticmethod
    def _keyword_search_query(
        *,
        collection_names: Sequence[str],
        query_text: str,
        k: int,
        company_name: Optional[str],
        sheet_name: Optional[str],
        scoped: bool,
    ) -> tuple[str, List[Any]]:
        where_extra: List[str] = []
        where_params: List[Any] = []
        if scoped and company_name:
            company = company_name.strip()
            where_extra.append(
                "AND (embedding.cmetadata->>'company_name' ILIKE %s ESCAPE '!' "
                "OR embedding.cmetadata->>'company_name' = %s)"
            )
            where_params.extend([f"%{_escape_like_term(company)}%", company])
        if scoped and sheet_name:
            where_extra.append("AND embedding.cmetadata->>'sheet_name' = %s")
            where_params.append(sheet_name)
        extra_sql = " ".join(where_extra)
        return (
            f"""
            SELECT
                embedding.document,
                embedding.cmetadata,
                ts_rank_cd(
                    to_tsvector('simple', embedding.document),
                    plainto_tsquery('simple', %s)
                ) AS fts_score,
                collection.name
            FROM langchain_pg_embedding AS embedding
            JOIN langchain_pg_collection AS collection
              ON collection.uuid = embedding.collection_id
            WHERE collection.name = ANY(%s)
              AND to_tsvector('simple', embedding.document)
                  @@ plainto_tsquery('simple', %s)
              {extra_sql}
            ORDER BY fts_score DESC, embedding.id
            LIMIT %s;
            """,
            [query_text, list(collection_names), query_text, *where_params, k],
        )

    @staticmethod
    def _keyword_search_rows(
        rows: Sequence[Sequence[Any]],
    ) -> List[Tuple[str, Dict[str, Any], float, str]]:
        return [
            (
                str(document),
                metadata if isinstance(metadata, dict) else {},
                float(score) if score is not None else 0.0,
                str(index_id),
            )
            for document, metadata, score, index_id in rows
        ]

    def keyword_search(
        self,
        *,
        collection_names: Sequence[str],
        query_text: str,
        k: int,
        company_name: Optional[str] = None,
        sheet_name: Optional[str] = None,
    ) -> List[Tuple[str, Dict[str, Any], float, str]]:
        """Run one collection-scoped PostgreSQL full-text search."""
        if not collection_names or not query_text.strip() or k <= 0:
            return []

        def fetch(*, scoped: bool) -> List[Tuple[str, Dict[str, Any], float, str]]:
            query, params = self._keyword_search_query(
                collection_names=collection_names,
                query_text=query_text,
                k=k,
                company_name=company_name,
                sheet_name=sheet_name,
                scoped=scoped,
            )
            connection = self._read_connection()
            try:
                with connection.cursor() as cursor:
                    cursor.execute(query, params)
                    rows = cursor.fetchall()
            finally:
                connection.close()
            return self._keyword_search_rows(rows)

        rows = fetch(scoped=True)
        if rows or not (company_name or sheet_name):
            return rows
        return fetch(scoped=False)

    async def keyword_search_async(
        self,
        *,
        collection_names: Sequence[str],
        query_text: str,
        k: int,
        company_name: Optional[str] = None,
        sheet_name: Optional[str] = None,
    ) -> List[Tuple[str, Dict[str, Any], float, str]]:
        """Run collection-scoped FTS through the process async connection pool."""
        if not collection_names or not query_text.strip() or k <= 0:
            return []

        async def fetch(*, scoped: bool) -> List[Tuple[str, Dict[str, Any], float, str]]:
            query, params = self._keyword_search_query(
                collection_names=collection_names,
                query_text=query_text,
                k=k,
                company_name=company_name,
                sheet_name=sheet_name,
                scoped=scoped,
            )
            raw_url = getattr(self, "database_url", PGVECTOR_URL).replace(
                "postgresql+psycopg://",
                "postgresql://",
            )
            async with get_pooled_async_connection(raw_url) as connection:
                async with connection.cursor() as cursor:
                    await cursor.execute(query, params)  # type: ignore[arg-type]
                    rows = await cursor.fetchall()
            return self._keyword_search_rows(rows)

        rows = await fetch(scoped=True)
        if rows or not (company_name or sheet_name):
            return rows
        return await fetch(scoped=False)

    def get_db_info(self) -> Dict[str, Any]:
        """Return connection details, versions, and vector collection statistics."""
        raw_url = getattr(self, "database_url", PGVECTOR_URL).replace(
            "postgresql+psycopg://", "postgresql://"
        )
        parsed = urlparse(raw_url)
        safe_host = parsed.hostname or "localhost"
        port = parsed.port or 5432
        db_name = parsed.path.lstrip("/") or "rag_flow"

        try:
            conn = self._raw_connection()
            try:
                with conn.cursor() as cur:
                    cur.execute("SELECT version();")
                    pg_version = cur.fetchone()[0]

                    cur.execute("SELECT extversion FROM pg_extension WHERE extname = 'vector';")
                    ext_row = cur.fetchone()
                    vector_version = ext_row[0] if ext_row else "not installed"

                    # The physical table names are retained for data compatibility.
                    cur.execute(
                        "SELECT COUNT(*) FROM information_schema.tables WHERE table_name = 'langchain_pg_collection';"
                    )
                    has_collections = cur.fetchone()[0] > 0

                    total_indexes = 0
                    total_chunks = 0
                    binary_quantized_indexes = 0
                    embedding_table_partitioned = False
                    if has_collections:
                        cur.execute("SELECT COUNT(*) FROM langchain_pg_collection;")
                        total_indexes = cur.fetchone()[0]
                        cur.execute("SELECT COUNT(*) FROM langchain_pg_embedding;")
                        total_chunks = cur.fetchone()[0]
                        cur.execute(
                            """
                            SELECT relkind = 'p'
                            FROM pg_catalog.pg_class
                            WHERE oid = 'langchain_pg_embedding'::regclass;
                            """
                        )
                        layout_row = cur.fetchone()
                        embedding_table_partitioned = bool(layout_row and layout_row[0])
                        cur.execute(
                            """
                            SELECT COUNT(*)
                            FROM pg_catalog.pg_indexes
                            WHERE schemaname = current_schema()
                              AND tablename = 'langchain_pg_embedding'
                              AND indexname LIKE 'idx_lc_hnsw_bq_c_%';
                            """
                        )
                        binary_quantized_indexes = int(cur.fetchone()[0])
            finally:
                conn.close()

            return {
                "connected": True,
                "host": safe_host,
                "port": port,
                "database": db_name,
                "postgres_version": pg_version.split()[1] if pg_version else "16",
                "pgvector_version": vector_version,
                "framework": "psycopg + pgvector",
                "total_indexes": total_indexes,
                "total_chunks": total_chunks,
                "vector_index_strategy": VECTOR_INDEX_STRATEGY,
                "vector_index_count": binary_quantized_indexes,
                "partition_strategy": VECTOR_PARTITION_STRATEGY,
                "embedding_table_partitioned": embedding_table_partitioned,
            }
        except Exception as err:
            return {
                "connected": False,
                "host": safe_host,
                "port": port,
                "database": db_name,
                "framework": "psycopg + pgvector",
                "error": str(err),
                "total_indexes": 0,
                "total_chunks": 0,
            }

    def put_documents(
        self,
        index_id: str,
        documents: Sequence[Document],
        model_name: str = DEFAULT_EMBEDDING_MODEL,
        embedding_encoder: Optional[EmbeddingEncoder] = None,
        metadata: Optional[Dict[str, Any]] = None,
        vectors: Optional[Any] = None,
        progress_callback: Optional[Callable[[Dict[str, int]], None]] = None,
    ) -> None:
        """
        Store documents in a pgvector collection, replacing any existing collection with the same identifier.

        Parameters:
            index_id (str): Identifier of the collection to replace.
            documents (List[Document]): Documents to store.
            model_name (str): Embedding model to use when generating vectors.
            embedding_encoder (Optional[EmbeddingEncoder]): Encoder for generating embeddings.
            metadata (Optional[Dict[str, Any]]): Metadata to associate with the collection.
            vectors (Optional[Any]): Precomputed vectors corresponding to all documents.
            progress_callback (Optional[Callable[[Dict[str, int]], None]]): Callback receiving batch and item progress.

        Raises:
            PgVectorStoreError: If inserting a document batch fails.
        """
        if not documents:
            return

        meta_dict = metadata or {}
        clean_meta = {
            "file_name": meta_dict.get("file_name", ""),
            "workbook_hash": meta_dict.get("workbook_hash", ""),
            "model": meta_dict.get("model", model_name),
            "dimension": meta_dict.get("dimension", DEFAULT_EMBEDDING_DIMENSION),
            "document_count": len(documents),
            "created_at": datetime.now(timezone.utc).isoformat(),
            "pipeline": meta_dict.get("pipeline", "luna_vlm_structured"),
            "duration_seconds": meta_dict.get("duration_seconds"),
            "total_tokens": meta_dict.get("total_tokens"),
            "estimated_cost_usd": meta_dict.get("estimated_cost_usd"),
            "estimated_cost_krw": meta_dict.get("estimated_cost_krw"),
            "batch_size": meta_dict.get("batch_size"),
            "company_name": meta_dict.get("company_name", ""),
            "ticker": meta_dict.get("ticker", ""),
        }

        # Validate vectors length BEFORE modifying the collection
        total_items = len(documents)
        use_precomputed_vectors = False
        if vectors is not None:
            if len(vectors) != total_items:
                raise PgVectorStoreError(
                    f"사전 계산된 벡터 개수({len(vectors)})가 문서 개수({total_items})와 일치하지 않습니다."
                )
            use_precomputed_vectors = True
            if isinstance(vectors, EmbeddingArtifactVectors) and vectors.dimension != int(
                clean_meta["dimension"]
            ):
                raise PgVectorStoreError(
                    "float32 아티팩트 차원과 컬렉션 메타데이터 차원이 일치하지 않습니다"
                )

        swap_suffix = uuid4().hex
        staging_name = f"{index_id}__staging__{swap_suffix}"
        retired_name = f"{index_id}__retired__{swap_suffix}"
        # Build a private staging collection so the current queryable index
        # remains intact until every batch has been persisted successfully.
        try:
            staging_uuid = self._create_collection(staging_name, clean_meta)
        except Exception as create_error:
            logger.error(
                "pgvector 컬렉션('%s') 생성 실패: %s",
                staging_name,
                create_error,
                exc_info=True,
            )
            raise PgVectorStoreError(
                f"pgvector 스테이징 컬렉션('{staging_name}') 생성 실패: {create_error}"
            ) from create_error

        total_batches = max(
            1,
            (total_items + PGVECTOR_INSERT_BATCH_SIZE - 1) // PGVECTOR_INSERT_BATCH_SIZE,
        )
        if progress_callback is not None:
            progress_callback(
                {
                    "completed_batches": 0,
                    "total_batches": total_batches,
                    "completed_items": 0,
                    "total_items": total_items,
                }
            )
        batch_index = 0
        copy_connection = self._raw_connection()
        try:
            if isinstance(vectors, EmbeddingArtifactVectors):
                copy_documents(
                    copy_connection,
                    collection_uuid=staging_uuid,
                    documents=documents,
                    vectors=vectors,
                    batch_size=PGVECTOR_INSERT_BATCH_SIZE,
                    progress_callback=progress_callback,
                )
                batch_index = total_batches
            else:
                if embedding_encoder is None:
                    raise PgVectorStoreError(
                        "동적 문서 임베딩에는 embedding_encoder 주입이 필요합니다"
                    )
                for batch_index, start in enumerate(
                    range(0, total_items, PGVECTOR_INSERT_BATCH_SIZE),
                    start=1,
                ):
                    stop = min(start + PGVECTOR_INSERT_BATCH_SIZE, total_items)
                    document_batch = list(documents[start:stop])
                    if use_precomputed_vectors and vectors is not None:
                        vector_batch = [
                            [float(value) for value in vector] for vector in vectors[start:stop]
                        ]
                    else:
                        vector_batch = embedding_encoder.encode(
                            [document.page_content for document in document_batch]
                        )
                    copy_documents(
                        copy_connection,
                        collection_uuid=staging_uuid,
                        documents=document_batch,
                        vectors=vector_batch,
                        batch_size=PGVECTOR_INSERT_BATCH_SIZE,
                    )
                    if progress_callback is not None:
                        progress_callback(
                            {
                                "completed_batches": batch_index,
                                "total_batches": total_batches,
                                "completed_items": stop,
                                "total_items": total_items,
                            }
                        )
            copy_connection.commit()
        except Exception as error:
            copy_connection.rollback()
            # A failed write must not leave a queryable partial collection.
            try:
                self._delete_collection(staging_name)
            except Exception:
                pass
            error_message = str(error)
            for marker in ("\n[SQL:", " [SQL:"):
                if marker in error_message:
                    error_message = error_message.split(marker, 1)[0].rstrip()
                    break
            raise PgVectorStoreError(
                "pgvector 문서 배치 적재 실패 "
                f"({batch_index}/{total_batches}): {error_message[:2000]}"
            ) from error
        finally:
            copy_connection.close()

        try:
            self.ensure_collection_vector_index(
                staging_name,
                int(clean_meta["dimension"]),
            )
        except Exception as index_error:
            staging_uuid = self._collection_uuid_cache.get(staging_name)
            try:
                self._delete_collection(staging_name)
            except Exception:
                pass
            if staging_uuid is not None:
                try:
                    self._drop_collection_vector_index(
                        staging_uuid,
                        int(clean_meta["dimension"]),
                    )
                except Exception:
                    pass
            raise PgVectorStoreError(
                f"컬렉션별 HNSW 인덱스 생성 실패: {index_error}"
            ) from index_error

        # Atomically publish the completed staging collection. Renaming the old
        # collection first keeps reads available and makes a failed rebuild
        # non-destructive.
        had_previous_collection = False
        retired_uuid: Optional[str] = None
        retired_dimension = int(clean_meta["dimension"])
        published_uuid: Optional[str] = None
        try:
            conn = self._raw_connection()
            try:
                with conn.cursor() as cur:
                    cur.execute(
                        """
                        SELECT pg_advisory_xact_lock(hashtext(%s));
                        """,
                        (index_id,),
                    )
                    cur.execute(
                        """
                        UPDATE langchain_pg_collection
                        SET name = %s
                        WHERE name = %s
                        RETURNING uuid, cmetadata;
                        """,
                        (retired_name, index_id),
                    )
                    retired_row = cur.fetchone()
                    had_previous_collection = retired_row is not None
                    if retired_row is not None:
                        retired_uuid = str(retired_row[0])
                        retired_metadata = retired_row[1]
                        if isinstance(retired_metadata, dict):
                            retired_dimension = int(
                                retired_metadata.get("dimension") or retired_dimension
                            )
                    cur.execute(
                        """
                        UPDATE langchain_pg_collection
                        SET name = %s, cmetadata = %s
                        WHERE name = %s
                        RETURNING uuid;
                        """,
                        (
                            index_id,
                            psycopg2.extras.Json(clean_meta),
                            staging_name,
                        ),
                    )
                    published_row = cur.fetchone()
                    if published_row is None:
                        raise PgVectorStoreError(
                            "완료된 pgvector 스테이징 컬렉션을 찾을 수 없습니다"
                        )
                    published_uuid = str(published_row[0])
                conn.commit()
            except Exception:
                conn.rollback()
                raise
            finally:
                conn.close()
        except Exception as publish_error:
            staging_uuid = self._collection_uuid_cache.get(staging_name)
            try:
                self._delete_collection(staging_name)
            except Exception:
                pass
            if staging_uuid is not None:
                try:
                    self._drop_collection_vector_index(
                        staging_uuid,
                        int(clean_meta["dimension"]),
                    )
                except Exception:
                    pass
            raise PgVectorStoreError(
                f"pgvector 컬렉션 원자적 교체 실패: {publish_error}"
            ) from publish_error

        if published_uuid is not None:
            with self._collection_uuid_lock:
                self._collection_uuid_cache.pop(staging_name, None)
                self._collection_uuid_cache[index_id] = published_uuid

        if had_previous_collection:
            try:
                cleanup_conn = self._raw_connection()
                try:
                    with cleanup_conn.cursor() as cur:
                        cur.execute(
                            "DELETE FROM langchain_pg_collection WHERE name = %s;",
                            (retired_name,),
                        )
                    cleanup_conn.commit()
                finally:
                    cleanup_conn.close()
                if retired_uuid is not None:
                    self._drop_collection_vector_index(
                        retired_uuid,
                        retired_dimension,
                    )
            except Exception:
                logger.warning(
                    "교체된 이전 pgvector 컬렉션 정리 실패: %s",
                    retired_name,
                    exc_info=True,
                )

        self.ensure_optimized_indexes()

    def put(
        self,
        index_id: str,
        vectors_or_items: Any = None,
        metadata: Optional[Dict[str, Any]] = None,
        embedding_encoder: Optional[EmbeddingEncoder] = None,
        vectors: Any = None,
        progress_callback: Optional[Callable[[Dict[str, int]], None]] = None,
    ) -> None:
        """
        Insert spreadsheet cell items into a vector collection.

        Parameters:
            vectors_or_items (Any): Cell items or precomputed vectors to use when inserting documents.
            metadata (Optional[Dict[str, Any]]): Collection and workbook metadata, including the cell items.
            progress_callback (Optional[Callable[[Dict[str, int]], None]]): Callback receiving insertion progress updates.
        """
        meta_dict = metadata or {}
        raw_items = meta_dict.get("items") or []
        model_name = meta_dict.get("model", DEFAULT_EMBEDDING_MODEL)
        file_name = meta_dict.get("file_name", "")
        workbook_hash = meta_dict.get("workbook_hash", "")
        company_name = meta_dict.get("company_name", "")

        resolved_vectors = vectors if vectors is not None else vectors_or_items
        usable_vectors = (
            resolved_vectors if _is_numeric_vector_collection(resolved_vectors) else None
        )

        docs = cell_items_to_langchain_documents(
            items=raw_items,
            file_name=file_name,
            workbook_hash=workbook_hash,
            index_id=index_id,
            company_name=company_name,
        )
        self.put_documents(
            index_id=index_id,
            documents=docs,
            model_name=model_name,
            embedding_encoder=embedding_encoder,
            metadata=metadata,
            vectors=usable_vectors,
            progress_callback=progress_callback,
        )

    def list_indexes(self) -> List[Dict[str, Any]]:
        """
        List all collections registered in the pgvector storage.

        Returns:
            List[Dict[str, Any]]: Collection summaries with metadata and document counts.
        """
        conn = self._raw_connection()
        try:
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT COUNT(*) FROM information_schema.tables WHERE table_name = 'langchain_pg_collection';"
                )
                if cur.fetchone()[0] == 0:
                    return []

                cur.execute(
                    """
                    SELECT c.name, c.cmetadata, COALESCE(e.chunk_count, 0) AS chunk_count
                    FROM langchain_pg_collection c
                    LEFT JOIN (
                        SELECT collection_id, COUNT(*) AS chunk_count
                        FROM langchain_pg_embedding
                        GROUP BY collection_id
                    ) e ON c.uuid = e.collection_id;
                    """
                )
                rows = cur.fetchall()
        finally:
            conn.close()

        results = []
        for r in rows:
            name = r[0]
            meta = r[1] or {}
            if isinstance(meta, str):
                import json

                try:
                    meta = json.loads(meta)
                except Exception:
                    meta = {}
            count = r[2]
            results.append(
                {
                    "index_id": name,
                    "file_name": meta.get("file_name", "unknown"),
                    "workbook_hash": meta.get("workbook_hash", ""),
                    "company_name": meta.get("company_name", ""),
                    "ticker": meta.get("ticker", ""),
                    "model": meta.get("model", DEFAULT_EMBEDDING_MODEL),
                    "dimension": meta.get("dimension", DEFAULT_EMBEDDING_DIMENSION),
                    "document_count": count,
                    "created_at": meta.get("created_at") or datetime.now(timezone.utc).isoformat(),
                    "storage": "PostgreSQL + pgvector",
                    "duration_seconds": meta.get("duration_seconds"),
                    "total_tokens": meta.get("total_tokens"),
                    "estimated_cost_usd": meta.get("estimated_cost_usd"),
                    "estimated_cost_krw": meta.get("estimated_cost_krw"),
                    "batch_size": meta.get("batch_size"),
                }
            )
        return results

    def list_data_scopes(self) -> List[Dict[str, Any]]:
        """Return the compact collection catalog used by query routing.

        Document counts come from collection metadata and sheet names come from
        the normalized ``sheets`` table, so routing never scans the embedding
        table merely to discover available data.
        """
        conn = self._raw_connection()
        try:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    SELECT
                        collection.name,
                        collection.cmetadata,
                        COALESCE(
                            ARRAY(
                                SELECT sheet.sheet_name
                                FROM sheets AS sheet
                                WHERE sheet.file_id = collection.cmetadata->>'workbook_hash'
                                  AND sheet.is_visible = TRUE
                                ORDER BY sheet.sheet_index, sheet.sheet_name
                            ),
                            ARRAY[]::text[]
                        ) AS sheet_names
                    FROM langchain_pg_collection AS collection
                    ORDER BY collection.name;
                    """
                )
                rows = cur.fetchall()
        finally:
            conn.close()

        return self._data_scope_rows(rows)

    @staticmethod
    def _data_scope_rows(
        rows: Sequence[Sequence[Any]],
    ) -> List[Dict[str, Any]]:
        scopes: List[Dict[str, Any]] = []
        for index_id, raw_metadata, raw_sheet_names in rows:
            metadata = raw_metadata if isinstance(raw_metadata, dict) else {}
            scopes.append(
                {
                    "index_id": str(index_id),
                    "file_name": str(metadata.get("file_name") or index_id),
                    "workbook_hash": str(metadata.get("workbook_hash") or index_id),
                    "company_name": str(metadata.get("company_name") or ""),
                    "sheet_names": [
                        str(name) for name in (raw_sheet_names or []) if str(name).strip()
                    ],
                    "model": str(metadata.get("model") or DEFAULT_EMBEDDING_MODEL),
                    "dimension": int(metadata.get("dimension") or DEFAULT_EMBEDDING_DIMENSION),
                    "document_count": int(metadata.get("document_count") or 0),
                }
            )
        return scopes

    async def list_data_scopes_async(self) -> List[Dict[str, Any]]:
        """Return routing scopes through the native async PostgreSQL pool."""
        raw_url = getattr(self, "database_url", PGVECTOR_URL).replace(
            "postgresql+psycopg://",
            "postgresql://",
        )
        async with get_pooled_async_connection(raw_url) as connection:
            async with connection.cursor() as cursor:
                await cursor.execute(
                    """
                    SELECT
                        collection.name,
                        collection.cmetadata,
                        COALESCE(
                            ARRAY(
                                SELECT sheet.sheet_name
                                FROM sheets AS sheet
                                WHERE sheet.file_id = collection.cmetadata->>'workbook_hash'
                                  AND sheet.is_visible = TRUE
                                ORDER BY sheet.sheet_index, sheet.sheet_name
                            ),
                            ARRAY[]::text[]
                        ) AS sheet_names
                    FROM langchain_pg_collection AS collection
                    ORDER BY collection.name;
                    """
                )
                rows = await cursor.fetchall()
        return self._data_scope_rows(rows)

    def get_index_metadata(self, index_id: str) -> Dict[str, Any]:
        """Retrieve metadata dictionary for a pgvector collection."""
        try:
            return self.get_index_detail(index_id, limit=1)
        except Exception:
            return {}

    def get_index_detail(self, index_id: str, limit: int = 15) -> Dict[str, Any]:
        """Retrieve collection detail and sample document chunks."""
        conn = self._raw_connection()
        try:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    SELECT c.uuid, c.name, c.cmetadata, COUNT(e.id) AS chunk_count
                    FROM langchain_pg_collection c
                    LEFT JOIN langchain_pg_embedding e ON c.uuid = e.collection_id
                    WHERE c.name = %s
                    GROUP BY c.uuid, c.name, c.cmetadata::text;
                    """,
                    (index_id,),
                )
                row = cur.fetchone()
                if not row:
                    raise PgVectorStoreError(f"pgvector 컬렉션을 찾을 수 없습니다: {index_id}")

                meta = row[2] or {}
                if isinstance(meta, str):
                    import json

                    try:
                        meta = json.loads(meta)
                    except Exception:
                        meta = {}

                cur.execute(
                    """
                    SELECT document, cmetadata
                    FROM langchain_pg_embedding
                    WHERE collection_id = %s
                    LIMIT %s;
                    """,
                    (row[0], limit),
                )
                chunk_rows = cur.fetchall()

                detected_tables_list = []
                sheet_names_list = []
                if meta.get("workbook_hash"):
                    cur.execute(
                        "SELECT sheet_name, detected_tables FROM sheets WHERE file_id = %s;",
                        (meta.get("workbook_hash"),),
                    )
                    sheet_rows = cur.fetchall()
                    for s_row in sheet_rows:
                        s_name = s_row[0]
                        sheet_names_list.append(s_name)
                        if s_row[1] and isinstance(s_row[1], list):
                            detected_tables_list.extend(s_row[1])
        finally:
            conn.close()

        sample_items = []
        for c in chunk_rows:
            text = c[0]
            cmeta = c[1] or {}
            if isinstance(cmeta, str):
                import json

                try:
                    cmeta = json.loads(cmeta)
                except Exception:
                    cmeta = {}
            sample_items.append(
                {
                    "cell_id": cmeta.get("cell_id", ""),
                    "sheet_name": cmeta.get("sheet_name", ""),
                    "cell_coord": cmeta.get("cell_coord", ""),
                    "row_header": cmeta.get("row_header", []),
                    "column_header": cmeta.get("column_header", []),
                    "cell_value": cmeta.get("cell_value", ""),
                    "text": text,
                }
            )

        luna_output = None
        if detected_tables_list:
            luna_output = {
                "file_name": meta.get("file_name", ""),
                "workbook_hash": meta.get("workbook_hash", ""),
                "sheet_names": sheet_names_list,
                "tables": detected_tables_list,
            }

        return {
            "index_id": index_id,
            "file_name": meta.get("file_name", ""),
            "workbook_hash": meta.get("workbook_hash", ""),
            "company_name": meta.get("company_name", ""),
            "ticker": meta.get("ticker", ""),
            "model": meta.get("model", ""),
            "dimension": meta.get("dimension", DEFAULT_EMBEDDING_DIMENSION),
            "document_count": row[3],
            "created_at": meta.get("created_at") or datetime.now(timezone.utc).isoformat(),
            "storage": "PostgreSQL + pgvector",
            "duration_seconds": meta.get("duration_seconds"),
            "total_tokens": meta.get("total_tokens"),
            "estimated_cost_usd": meta.get("estimated_cost_usd"),
            "estimated_cost_krw": meta.get("estimated_cost_krw"),
            "batch_size": meta.get("batch_size"),
            "sample_items": sample_items,
            "sheet_names": sheet_names_list,
            "tables": detected_tables_list,
            "luna_output": luna_output,
        }

    def update_index_company(self, index_id: str, company_name: str) -> Dict[str, Any]:
        """Update company_name in collection metadata and cascade to all chunks and DB tables."""
        conn = None
        try:
            conn = self._raw_connection()
            with conn.cursor() as cur:
                # 1. Update langchain_pg_collection (cmetadata is column type json)
                cur.execute(
                    """
                    UPDATE langchain_pg_collection
                    SET cmetadata = jsonb_set(
                        COALESCE(cmetadata::jsonb, '{}'::jsonb),
                        '{company_name}',
                        to_jsonb(%s::text)
                    )::json
                    WHERE name = %s
                    RETURNING uuid, cmetadata;
                    """,
                    (company_name, index_id),
                )
                row = cur.fetchone()
                if not row:
                    raise PgVectorStoreError(f"pgvector 컬렉션을 찾을 수 없습니다: {index_id}")

                collection_uuid = row[0]
                col_meta = row[1] or {}
                if isinstance(col_meta, str):
                    import json

                    try:
                        col_meta = json.loads(col_meta)
                    except Exception:
                        col_meta = {}

                # 2. Update langchain_pg_embedding (cascade to all chunks in this collection)
                cur.execute(
                    """
                    UPDATE langchain_pg_embedding
                    SET cmetadata = jsonb_set(
                            COALESCE(cmetadata, '{}'::jsonb),
                            '{company_name}',
                            to_jsonb(%s::text)
                        ),
                        document = regexp_replace(
                            CASE
                                WHEN document ~ '^Company:\\s*[^|]*\\|\\s*'
                                    THEN document
                                ELSE 'Company: ? | ' || document
                            END,
                            '^Company:\\s*[^|]*\\|\\s*',
                            'Company: ' || %s || ' | '
                        )
                    WHERE collection_id = %s;
                    """,
                    (company_name, company_name, collection_uuid),
                )

            conn.commit()
        except PgVectorStoreError:
            raise
        except Exception as err:
            raise PgVectorStoreError(f"기업명 수정 실패: {err}") from err
        finally:
            if conn is not None:
                try:
                    conn.close()
                except Exception:
                    pass

        return self.get_index_detail(index_id)

    def ensure_optimized_indexes(self) -> None:
        """Create shared metadata and full-text indexes when they don't exist."""
        conn = None
        try:
            conn = self._raw_connection()
            if hasattr(conn, "autocommit"):
                conn.autocommit = True
            try:
                with conn.cursor() as cur:
                    cur.execute(
                        "SELECT COUNT(*) FROM information_schema.tables WHERE table_name = 'langchain_pg_embedding';"
                    )
                    if cur.fetchone()[0] > 0:
                        self._drop_invalid_optimized_indexes(cur)
                        index_stmts = [
                            "CREATE INDEX CONCURRENTLY IF NOT EXISTS idx_langchain_pg_embedding_cell_id ON langchain_pg_embedding ((cmetadata->>'cell_id'));",
                            "CREATE INDEX CONCURRENTLY IF NOT EXISTS idx_langchain_pg_embedding_cell_coord_upper ON langchain_pg_embedding ((UPPER(cmetadata->>'cell_coord')));",
                            "CREATE INDEX CONCURRENTLY IF NOT EXISTS idx_langchain_pg_embedding_workbook_hash ON langchain_pg_embedding ((cmetadata->>'workbook_hash'));",
                            "CREATE INDEX CONCURRENTLY IF NOT EXISTS idx_langchain_pg_embedding_company_name ON langchain_pg_embedding ((cmetadata->>'company_name'));",
                            "CREATE INDEX CONCURRENTLY IF NOT EXISTS idx_langchain_pg_embedding_sheet_name ON langchain_pg_embedding ((cmetadata->>'sheet_name'));",
                            """
                            CREATE INDEX CONCURRENTLY IF NOT EXISTS idx_langchain_pg_embedding_sheet_row
                            ON langchain_pg_embedding (
                                (cmetadata->>'sheet_name'),
                                ((cmetadata->>'row_index')::int)
                            )
                            WHERE cmetadata->>'row_index' ~ '^\\d+$';
                            """,
                            "CREATE INDEX CONCURRENTLY IF NOT EXISTS idx_langchain_pg_embedding_collection_id ON langchain_pg_embedding (collection_id);",
                            """
                            CREATE INDEX CONCURRENTLY IF NOT EXISTS idx_langchain_pg_embedding_document_fts
                            ON langchain_pg_embedding
                            USING gin (to_tsvector('simple', document));
                            """,
                            """
                            CREATE INDEX IF NOT EXISTS idx_langchain_pg_embedding_cmetadata
                            ON langchain_pg_embedding
                            USING gin (cmetadata jsonb_path_ops);
                            """,
                        ]
                        for stmt in index_stmts:
                            try:
                                cur.execute(stmt)
                            except Exception as stmt_err:
                                logger.debug("Index creation notice: %s (%s)", stmt[:60], stmt_err)
            except Exception as error:
                logger.warning("pgvector 최적화 인덱스 생성 실패: %s", error)
                try:
                    with conn.cursor() as cur:
                        self._drop_invalid_optimized_indexes(cur)
                except Exception as cleanup_error:
                    logger.error(
                        "유효하지 않은 pgvector 최적화 인덱스 정리 실패: %s",
                        cleanup_error,
                    )
            finally:
                conn.close()
        except Exception as pool_error:
            logger.warning("ensure_optimized_indexes: 연결 풀 오류: %s", pool_error)
            if conn is not None:
                try:
                    conn.close()
                except Exception:
                    pass

    @staticmethod
    def _drop_invalid_optimized_indexes(cur: Any) -> None:
        """Drop invalid concurrent indexes so a later setup can rebuild them."""
        cur.execute(
            """
            SELECT index_class.relname
            FROM pg_catalog.pg_index AS index_state
            JOIN pg_catalog.pg_class AS index_class
              ON index_class.oid = index_state.indexrelid
            JOIN pg_catalog.pg_namespace AS index_namespace
              ON index_namespace.oid = index_class.relnamespace
            WHERE index_namespace.nspname = current_schema()
              AND index_class.relname = ANY(%s)
              AND NOT index_state.indisvalid;
            """,
            (list(_CONCURRENT_OPTIMIZED_INDEX_NAMES),),
        )
        invalid_index_names = [row[0] for row in cur.fetchall()]
        for index_name in invalid_index_names:
            if index_name not in _CONCURRENT_OPTIMIZED_INDEX_NAMES:
                continue
            cur.execute(f'DROP INDEX CONCURRENTLY IF EXISTS "{index_name}";')

    def delete(self, index_id: str) -> bool:
        """Delete a collection from pgvector.

        Returns:
            bool: True if deleted, False if not found.

        Raises:
            PgVectorStoreError: If database operation fails (not including not-found case).
        """
        conn = self._raw_connection()
        try:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    DELETE FROM langchain_pg_collection
                    WHERE name = %s
                    RETURNING uuid, cmetadata;
                    """,
                    (index_id,),
                )
                deleted_row = cur.fetchone()
            conn.commit()
        except Exception as error:
            conn.rollback()
            raise PgVectorStoreError(
                f"pgvector 컬렉션 삭제 중 데이터베이스 오류 발생: {error}"
            ) from error
        finally:
            conn.close()
        if deleted_row is None:
            return False
        with self._collection_uuid_lock:
            self._collection_uuid_cache.pop(index_id, None)
        metadata = deleted_row[1] if isinstance(deleted_row[1], dict) else {}
        dimension = int(metadata.get("dimension") or DEFAULT_EMBEDDING_DIMENSION)
        try:
            self._drop_collection_vector_index(str(deleted_row[0]), dimension)
        except Exception:
            logger.warning(
                "삭제된 컬렉션의 HNSW 인덱스 정리 실패: %s",
                index_id,
                exc_info=True,
            )
        return True

    def delete_by_workbook_hash(self, workbook_hash: str) -> int:
        """Logical cascade deletion: remove all collections matching a deleted workbook hash."""
        if not workbook_hash:
            return 0
        try:
            conn = self._raw_connection()
            try:
                with conn.cursor() as cur:
                    cur.execute(
                        """
                        DELETE FROM langchain_pg_collection
                        WHERE cmetadata->>'workbook_hash' = %s
                        RETURNING name, uuid, cmetadata;
                        """,
                        (workbook_hash,),
                    )
                    deleted_rows = cur.fetchall()
                conn.commit()
            finally:
                conn.close()
            for collection_name, collection_uuid, raw_metadata in deleted_rows:
                with self._collection_uuid_lock:
                    self._collection_uuid_cache.pop(str(collection_name), None)
                metadata = raw_metadata if isinstance(raw_metadata, dict) else {}
                dimension = int(metadata.get("dimension") or DEFAULT_EMBEDDING_DIMENSION)
                try:
                    self._drop_collection_vector_index(
                        str(collection_uuid),
                        dimension,
                    )
                except Exception:
                    logger.warning(
                        "삭제된 컬렉션의 HNSW 인덱스 정리 실패: %s",
                        collection_name,
                        exc_info=True,
                    )
            return len(deleted_rows)
        except Exception:
            return 0

    def search(
        self,
        index_id: str,
        query_text: str,
        model_name: str = DEFAULT_EMBEDDING_MODEL,
        embedding_encoder: Optional[EmbeddingEncoder] = None,
        limit: int = 5,
    ) -> List[Tuple[float, Dict[str, Any]]]:
        """Perform similarity search with quantized HNSW candidates and exact reranking."""
        try:
            if embedding_encoder is None:
                raise PgVectorStoreError("유사도 검색에는 embedding_encoder 주입이 필요합니다")
            vectors = embedding_encoder.encode_for_model([query_text], model_name)
            if not vectors:
                raise PgVectorStoreError("검색 질의 임베딩 결과가 비어 있습니다")
            query_vector = vectors[0]
            hits = self.similarity_search_by_vector_with_score(
                collection_name=index_id,
                embedding=query_vector,
                k=limit,
            )
            results = []
            for doc, score in hits:
                similarity = max(0.0, 1.0 - float(score))
                cell_item = langchain_document_to_cell_item(doc, score=similarity)
                results.append((similarity, cell_item))
            return results
        except Exception as search_err:
            raise PgVectorStoreError(
                f"PostgreSQL pgvector 유사도 검색 실패 ({index_id}): {search_err}"
            ) from search_err

    def similarity_search_by_vector_with_score(
        self,
        collection_name: str,
        embedding: List[float],
        k: int = 10,
        sheet_names: Optional[List[str]] = None,
        company_name: Optional[str] = None,
    ) -> List[Tuple[Any, float]]:
        """
        Search a collection for documents nearest to an embedding vector with optional company and sheet filters.

        Parameters:
            collection_name (str): Name of the collection to search.
            embedding (List[float]): Query embedding vector.
            k (int): Maximum number of results to return.
            sheet_names (Optional[List[str]]): Specific worksheet names to restrict search scope.
            company_name (Optional[str]): Company/Entity identifier to restrict search scope.

        Returns:
            List[Tuple[Any, float]]: Document and cosine-distance pairs, or an empty list when the collection does not exist.

        Raises:
            PgVectorStoreError: If the direct indexed search fails.
        """
        conn = None
        try:
            collection_uuid = self._collection_uuid(collection_name)
            if collection_uuid is None:
                return []
            query_sql, full_params = self._dense_search_query(
                collection_uuid=collection_uuid,
                embedding=embedding,
                k=k,
                sheet_names=sheet_names,
                company_name=company_name,
            )
            conn = self._read_connection()
            with conn.cursor() as cur:
                cur.execute(query_sql, full_params)
                rows = cur.fetchall()
            return self._dense_search_rows(rows)
        except Exception as error:
            raise PgVectorStoreError(
                f"PostgreSQL pgvector 유사도 검색 실패 ({collection_name}): {error}"
            ) from error
        finally:
            if conn is not None:
                try:
                    conn.close()
                except Exception:
                    pass

    @staticmethod
    def _dense_search_query(
        *,
        collection_uuid: str,
        embedding: Sequence[float],
        k: int,
        sheet_names: Optional[Sequence[str]],
        company_name: Optional[str],
    ) -> tuple[str, List[Any]]:
        """Build the shared sync/async exact-rerank query and bound parameters."""
        canonical_uuid = str(UUID(collection_uuid))
        dim = len(embedding)
        vector_param = "[" + ",".join(format(float(value), ".17g") for value in embedding) + "]"
        where_clauses = [f"collection_id = '{canonical_uuid}'::uuid"]
        params: List[Any] = []
        if 0 < dim <= 64_000:
            candidate_limit = min(max(k * 20, 100), 1000)
            index_where_sql = (
                f"collection_id = '{canonical_uuid}'::uuid AND vector_dims(embedding) = {dim}"
            )
            outer_where_clauses: List[str] = []
            outer_params: List[Any] = []
            if company_name and company_name.strip():
                clean_company = company_name.strip()
                escaped_company = _escape_like_term(clean_company)
                outer_where_clauses.append(
                    "(cmetadata->>'company_name' ILIKE %s ESCAPE '!' "
                    "OR cmetadata->>'company_name' = %s)"
                )
                outer_params.extend([f"%{escaped_company}%", clean_company])
            if sheet_names:
                valid_sheets = [
                    sheet.strip()
                    for sheet in sheet_names
                    if isinstance(sheet, str) and sheet.strip()
                ]
                if valid_sheets:
                    outer_where_clauses.append("(cmetadata->>'sheet_name' = ANY(%s))")
                    outer_params.append(valid_sheets)
            outer_where_sql = (
                f"WHERE {' AND '.join(outer_where_clauses)}" if outer_where_clauses else ""
            )
            return (
                f"""
                WITH query_vector AS MATERIALIZED (
                    SELECT %s::vector AS embedding
                ), candidates AS MATERIALIZED (
                    SELECT id, document, cmetadata, source.embedding
                    FROM langchain_pg_embedding AS source
                    WHERE {index_where_sql}
                    ORDER BY
                        binary_quantize(source.embedding)::bit({dim})
                        <~> binary_quantize(
                            (SELECT embedding FROM query_vector)
                        )::bit({dim})
                    LIMIT %s
                )
                SELECT id, document, cmetadata,
                       candidates.embedding <=> query_vector.embedding AS distance
                FROM candidates
                CROSS JOIN query_vector
                {outer_where_sql}
                ORDER BY candidates.embedding <=> query_vector.embedding
                LIMIT %s;
                """,
                [vector_param, candidate_limit, *outer_params, k],
            )

        if company_name and company_name.strip():
            clean_company = company_name.strip()
            escaped_company = _escape_like_term(clean_company)
            where_clauses.append(
                "(cmetadata->>'company_name' ILIKE %s ESCAPE '!' "
                "OR cmetadata->>'company_name' = %s)"
            )
            params.extend([f"%{escaped_company}%", clean_company])
        if sheet_names:
            valid_sheets = [
                sheet.strip() for sheet in sheet_names if isinstance(sheet, str) and sheet.strip()
            ]
            if valid_sheets:
                where_clauses.append("(cmetadata->>'sheet_name' = ANY(%s))")
                params.append(valid_sheets)
        where_sql = " AND ".join(where_clauses)
        return (
            f"""
            SELECT id, document, cmetadata, (embedding <=> %s::vector) AS distance
            FROM langchain_pg_embedding
            WHERE {where_sql}
            ORDER BY embedding <=> %s::vector
            LIMIT %s;
            """,
            [vector_param, *params, vector_param, k],
        )

    @staticmethod
    def _dense_search_rows(
        rows: Sequence[Sequence[Any]],
    ) -> List[Tuple[Document, float]]:
        import json

        results: List[Tuple[Document, float]] = []
        for row in rows:
            row_id, text, metadata, distance = row
            if isinstance(metadata, str):
                try:
                    metadata = json.loads(metadata)
                except Exception:
                    metadata = {}
            results.append(
                (
                    Document(
                        page_content=str(text or ""),
                        metadata=metadata if isinstance(metadata, dict) else {},
                        id=str(row_id),
                    ),
                    float(distance) if distance is not None else 0.0,
                )
            )
        return results

    async def similarity_search_by_vector_with_score_async(
        self,
        collection_name: str,
        embedding: Sequence[float],
        k: int = 10,
        *,
        sheet_names: Optional[Sequence[str]] = None,
        company_name: Optional[str] = None,
    ) -> List[Tuple[Document, float]]:
        """Run the collection-local HNSW query through native psycopg async I/O."""
        try:
            collection_uuid = await self._collection_uuid_async(collection_name)
            if collection_uuid is None:
                return []
            query_sql, params = self._dense_search_query(
                collection_uuid=collection_uuid,
                embedding=embedding,
                k=k,
                sheet_names=sheet_names,
                company_name=company_name,
            )
            raw_url = getattr(self, "database_url", PGVECTOR_URL).replace(
                "postgresql+psycopg://",
                "postgresql://",
            )
            async with get_pooled_async_connection(raw_url) as connection:
                async with connection.cursor() as cursor:
                    await cursor.execute(query_sql, params)  # type: ignore[arg-type]
                    rows = await cursor.fetchall()
            return self._dense_search_rows(rows)
        except Exception as error:
            raise PgVectorStoreError(
                f"PostgreSQL pgvector 유사도 검색 실패 ({collection_name}): {error}"
            ) from error

    @staticmethod
    def _cell_metadata_query(
        *,
        clean_ids: Sequence[str],
        workbook_hash: Optional[str],
        company_name: Optional[str],
        collection_uuid: Optional[str],
        limit: int,
        cell_references: Optional[Sequence[Dict[str, Optional[str]]]],
    ) -> Optional[tuple[str, List[Any]]]:
        extracted_coords: List[str] = []
        for cell_id in clean_ids:
            parts = cell_id.replace(":", " ").replace("!", " ").split()
            for part in parts:
                candidate = part.strip()
                if (
                    candidate
                    and candidate[0].isalpha()
                    and any(character.isdigit() for character in candidate)
                ):
                    extracted_coords.append(candidate.upper())
        all_search_targets = list(
            dict.fromkeys(item.upper() for item in [*clean_ids, *extracted_coords])
        )
        where_clauses: List[str] = []
        params: List[Any] = []
        if cell_references is None:
            where_clauses.append(
                """(
                    cmetadata->>'cell_id' = ANY(%s)
                    OR cmetadata->>'cell_coord' = ANY(%s)
                    OR UPPER(cmetadata->>'cell_coord') = ANY(%s)
                )"""
            )
            params.extend([list(clean_ids), all_search_targets, all_search_targets])
        else:
            triple_companies: List[str] = []
            triple_sheets: List[str] = []
            triple_coords: List[str] = []
            sheet_pair_sheets: List[str] = []
            sheet_pair_coords: List[str] = []
            company_pair_companies: List[str] = []
            company_pair_coords: List[str] = []
            unqualified_coords: List[str] = []
            for reference in cell_references:
                coord = (reference.get("cell_coord") or "").strip().upper()
                if not coord:
                    continue
                ref_sheet = (reference.get("sheet_name") or "").strip()
                ref_company = (reference.get("company_name") or "").strip()
                if ref_company and ref_sheet:
                    triple_companies.append(ref_company.upper())
                    triple_sheets.append(ref_sheet.upper())
                    triple_coords.append(coord)
                elif ref_sheet:
                    sheet_pair_sheets.append(ref_sheet.upper())
                    sheet_pair_coords.append(coord)
                elif ref_company:
                    company_pair_companies.append(ref_company.upper())
                    company_pair_coords.append(coord)
                else:
                    unqualified_coords.append(coord)
            if unqualified_coords:
                where_clauses.append("UPPER(cmetadata->>'cell_coord') = ANY(%s)")
                params.append(list(dict.fromkeys(unqualified_coords)))
            if sheet_pair_coords:
                where_clauses.append(
                    """EXISTS (
                        SELECT 1
                        FROM unnest(%s::text[], %s::text[])
                            AS reference(sheet_name, cell_coord)
                        WHERE UPPER(cmetadata->>'sheet_name') = reference.sheet_name
                          AND UPPER(cmetadata->>'cell_coord') = reference.cell_coord
                    )"""
                )
                params.extend([sheet_pair_sheets, sheet_pair_coords])
            if company_pair_coords:
                where_clauses.append(
                    """EXISTS (
                        SELECT 1
                        FROM unnest(%s::text[], %s::text[])
                            AS reference(company_name, cell_coord)
                        WHERE UPPER(COALESCE(cmetadata->>'company_name', '')) = reference.company_name
                          AND UPPER(cmetadata->>'cell_coord') = reference.cell_coord
                    )"""
                )
                params.extend([company_pair_companies, company_pair_coords])
            if triple_coords:
                where_clauses.append(
                    """EXISTS (
                        SELECT 1
                        FROM unnest(%s::text[], %s::text[], %s::text[])
                            AS reference(company_name, sheet_name, cell_coord)
                        WHERE UPPER(COALESCE(cmetadata->>'company_name', '')) = reference.company_name
                          AND UPPER(cmetadata->>'sheet_name') = reference.sheet_name
                          AND UPPER(cmetadata->>'cell_coord') = reference.cell_coord
                    )"""
                )
                params.extend([triple_companies, triple_sheets, triple_coords])
            if not where_clauses:
                return None

        query = f"""
            WITH ranked_cells AS (
            SELECT
                id,
                document,
                cmetadata,
                cmetadata->>'cell_id' AS cell_id,
                cmetadata->>'cell_coord' AS cell_coord,
                cmetadata->>'sheet_name' AS sheet_name,
                cmetadata->>'cell_value' AS cell_value,
                cmetadata->'row_header' AS row_header,
                cmetadata->'column_header' AS column_header,
                cmetadata->>'company_name' AS company_name,
                ROW_NUMBER() OVER (
                    PARTITION BY
                        UPPER(COALESCE(cmetadata->>'company_name', '')),
                        UPPER(cmetadata->>'sheet_name'),
                        UPPER(cmetadata->>'cell_coord')
                    ORDER BY
                        CASE
                            WHEN document NOT LIKE '%%Cell Value: ?%%' AND document NOT LIKE '%%Cell Value: NA%%' AND NULLIF(cmetadata->>'cell_value', '') IS NOT NULL AND cmetadata->>'cell_value' NOT IN ('?', 'NA') THEN 0
                            WHEN cmetadata->>'variant' = 'header_with_value' THEN 1
                            WHEN document NOT LIKE '%%Cell Value: ?%%' THEN 2
                            WHEN cmetadata->>'variant' = 'header_only' THEN 3
                            ELSE 4
                        END,
                        CASE WHEN jsonb_typeof(cmetadata->'row_header') = 'array'
                            THEN jsonb_array_length(cmetadata->'row_header') ELSE 0 END DESC,
                        CASE WHEN jsonb_typeof(cmetadata->'column_header') = 'array'
                            THEN jsonb_array_length(cmetadata->'column_header') ELSE 0 END DESC,
                        id
                ) AS cell_rank
            FROM langchain_pg_embedding
            WHERE ({" OR ".join(where_clauses)})
        """
        if collection_uuid:
            query += " AND collection_id = %s"
            params.append(collection_uuid)
        elif workbook_hash:
            query += " AND cmetadata->>'workbook_hash' = %s"
            params.append(workbook_hash)
        if company_name and company_name.strip():
            query += " AND UPPER(cmetadata->>'company_name') = %s"
            params.append(company_name.strip().upper())
        query += """
            )
            SELECT
                id,
                document,
                cmetadata,
                COALESCE(cell_id, sheet_name || ' Cell ' || cell_coord) AS cell_id,
                cell_coord,
                sheet_name,
                cell_value,
                row_header,
                column_header,
                company_name
            FROM ranked_cells
            WHERE cell_rank = 1
            ORDER BY
                CASE WHEN COALESCE(cell_id, sheet_name || ' Cell ' || cell_coord) = ANY(%s) THEN 0 ELSE 1 END,
                UPPER(sheet_name),
                UPPER(cell_coord),
                id
            LIMIT %s;
        """
        params.extend([list(clean_ids), limit])
        return query, params

    @staticmethod
    def _cell_metadata_rows(
        rows: Sequence[Sequence[Any]],
    ) -> List[Dict[str, Any]]:
        import json

        results: List[Dict[str, Any]] = []
        for row in rows:
            (
                _row_id,
                text,
                _metadata,
                cell_id,
                cell_coord,
                sheet_name,
                cell_value,
                row_header,
                column_header,
                company_name,
            ) = row
            metadata = _metadata if isinstance(_metadata, dict) else {}
            if isinstance(_metadata, str):
                try:
                    metadata = json.loads(_metadata)
                except Exception:
                    metadata = {}
            if isinstance(row_header, str):
                try:
                    row_header = json.loads(row_header)
                except Exception:
                    pass
            if isinstance(column_header, str):
                try:
                    column_header = json.loads(column_header)
                except Exception:
                    pass
            results.append(
                {
                    "cell_id": cell_id or f"{sheet_name}:{cell_coord}",
                    "cell_coord": cell_coord,
                    "sheet_name": sheet_name,
                    "cell_value": cell_value,
                    "row_header": (
                        row_header
                        if isinstance(row_header, list)
                        else ([row_header] if row_header else [])
                    ),
                    "column_header": (
                        column_header
                        if isinstance(column_header, list)
                        else ([column_header] if column_header else [])
                    ),
                    "company_name": company_name,
                    "file_name": metadata.get("file_name", ""),
                    "workbook_hash": metadata.get("workbook_hash", ""),
                    "index_id": metadata.get("index_id", ""),
                    "source_text": text,
                }
            )
        return results

    async def fetch_cells_by_metadata_async(
        self,
        cell_identifiers: Sequence[str],
        workbook_hash: Optional[str] = None,
        company_name: Optional[str] = None,
        collection_name: Optional[str] = None,
        limit: int = 50,
        cell_references: Optional[Sequence[Dict[str, Optional[str]]]] = None,
    ) -> List[Dict[str, Any]]:
        """Fetch qualified cells through native async PostgreSQL I/O."""
        if not cell_identifiers or not (workbook_hash or collection_name or company_name):
            return []
        clean_ids = [cell_id.strip() for cell_id in cell_identifiers if cell_id and cell_id.strip()]
        if not clean_ids:
            return []
        collection_uuid = (
            await self._collection_uuid_async(collection_name) if collection_name else None
        )
        if collection_name and collection_uuid is None and not (workbook_hash or company_name):
            return []
        query_spec = self._cell_metadata_query(
            clean_ids=clean_ids,
            workbook_hash=workbook_hash,
            company_name=company_name,
            collection_uuid=collection_uuid,
            limit=limit,
            cell_references=cell_references,
        )
        if query_spec is None:
            return []
        query, params = query_spec
        raw_url = getattr(self, "database_url", PGVECTOR_URL).replace(
            "postgresql+psycopg://",
            "postgresql://",
        )
        try:
            async with get_pooled_async_connection(raw_url) as connection:
                async with connection.cursor() as cursor:
                    await cursor.execute(query, params)  # type: ignore[arg-type]
                    rows = await cursor.fetchall()
            return self._cell_metadata_rows(rows)
        except Exception:
            logger.exception("비동기 직접 셀 메타데이터 조회 실패")
            return []

    def fetch_cells_by_metadata(
        self,
        cell_identifiers: List[str],
        workbook_hash: Optional[str] = None,
        company_name: Optional[str] = None,
        collection_name: Optional[str] = None,
        limit: int = 50,
        cell_references: Optional[List[Dict[str, Optional[str]]]] = None,
    ) -> List[Dict[str, Any]]:
        """
        Fetch cell documents matching the provided identifiers, optionally filtered by workbook, company, or collection.

        Parameters:
            cell_identifiers (List[str]): Cell IDs, coordinates, or strings containing coordinate-like values.
            workbook_hash (Optional[str]): Restricts results to a workbook with this hash.
            company_name (Optional[str]): Restricts results to this company name.
            collection_name (Optional[str]): Restricts results to this collection.
            limit (int): Maximum number of matching cells to return.
            cell_references (Optional[List[Dict[str, Optional[str]]]]): Structured
                coordinate filters whose optional ``sheet_name`` and ``company_name``
                are matched together with the coordinate.

        Returns:
            List[Dict[str, Any]]: Normalized cell records containing identifiers, values, headers, source text, and company metadata. Returns an empty list when the input is empty or retrieval fails.
        """
        if not cell_identifiers or not (workbook_hash or collection_name or company_name):
            if cell_identifiers:
                logger.warning(
                    "직접 셀 메타데이터 조회를 거부했습니다: collection_name, workbook_hash 또는 company_name이 필요합니다"
                )
            return []
        clean_ids = [cell_id.strip() for cell_id in cell_identifiers if cell_id and cell_id.strip()]
        if not clean_ids:
            return []
        collection_uuid = self._collection_uuid(collection_name) if collection_name else None
        if collection_name and collection_uuid is None and not (workbook_hash or company_name):
            return []
        query_spec = self._cell_metadata_query(
            clean_ids=clean_ids,
            workbook_hash=workbook_hash,
            company_name=company_name,
            collection_uuid=collection_uuid,
            limit=limit,
            cell_references=cell_references,
        )
        if query_spec is None:
            return []
        query, params = query_spec
        connection = self._raw_connection()
        try:
            with connection.cursor() as cursor:
                cursor.execute(query, params)
                rows = cursor.fetchall()
        except Exception:
            logger.exception("직접 셀 메타데이터 조회 실패")
            return []
        finally:
            connection.close()
        return self._cell_metadata_rows(rows)

    def fetch_adjacent_row_cells(
        self,
        collection_name: Optional[str] = None,
        workbook_hash: Optional[str] = None,
        sheet_name: Optional[str] = None,
        row_index: Optional[int] = None,
        limit: int = 50,
    ) -> List[Dict[str, Any]]:
        """Fetch all cells on a given row index within a sheet and collection for horizontal timeseries expansion."""
        if not (collection_name or workbook_hash) or row_index is None:
            return []
        return self.fetch_rows_cells(
            collection_name=collection_name,
            workbook_hash=workbook_hash,
            sheet_name=sheet_name,
            row_indices=[row_index],
            limit_per_row=limit,
        ).get(row_index, [])

    @staticmethod
    def _rows_cells_query(
        *,
        collection_name: Optional[str],
        workbook_hash: Optional[str],
        sheet_name: Optional[str],
        row_indices: Sequence[int],
    ) -> Optional[tuple[str, List[Any]]]:
        unique_rows = sorted({row for row in row_indices if row > 0})
        if not unique_rows or not (collection_name or workbook_hash):
            return None
        collection_names = (
            [name.strip() for name in collection_name.split(",") if name.strip()]
            if collection_name
            else []
        )
        scope_clauses: List[str] = []
        scope_params: List[Any] = []
        cte_params: List[Any] = []
        if collection_names:
            collection_cte_sql = (
                "target_collections AS ("
                "SELECT uuid FROM langchain_pg_collection WHERE name = ANY(%s)"
                "),"
            )
            cte_params.append(collection_names)
            scope_clauses.append("collection_id IN (SELECT uuid FROM target_collections)")
        else:
            collection_cte_sql = ""
        if workbook_hash and not collection_names:
            workbook_hashes = [value.strip() for value in workbook_hash.split(",") if value.strip()]
            scope_clauses.append("cmetadata->>'workbook_hash' = ANY(%s)")
            scope_params.append(workbook_hashes)
        if sheet_name:
            scope_clauses.append("cmetadata->>'sheet_name' = %s")
            scope_params.append(sheet_name)
        scope_sql = " AND " + " AND ".join(scope_clauses) if scope_clauses else ""
        return (
            rf"""
            WITH {collection_cte_sql}
            filtered_rows AS (
                SELECT
                    id,
                    document,
                    cmetadata,
                    COALESCE(
                        CASE WHEN cmetadata->>'row_index' ~ '^\d+$' THEN (cmetadata->>'row_index')::int ELSE NULL END,
                        NULLIF(regexp_replace(cmetadata->>'cell_coord', '[^0-9]', '', 'g'), '')::int
                    ) AS resolved_row_index,
                    CASE
                        WHEN cmetadata->>'col_index' ~ '^\d+$'
                            THEN (cmetadata->>'col_index')::int
                        ELSE NULL
                    END AS resolved_col_index,
                    ROW_NUMBER() OVER (
                        PARTITION BY COALESCE(
                            CASE WHEN cmetadata->>'row_index' ~ '^\d+$' THEN (cmetadata->>'row_index')::int ELSE NULL END,
                            NULLIF(regexp_replace(cmetadata->>'cell_coord', '[^0-9]', '', 'g'), '')::int
                        ),
                        cmetadata->>'cell_coord'
                        ORDER BY
                            CASE WHEN cmetadata->>'variant' = 'header_with_value' THEN 0 ELSE 1 END,
                            CASE WHEN jsonb_typeof(cmetadata->'row_header') = 'array'
                                THEN jsonb_array_length(cmetadata->'row_header') ELSE 0 END DESC,
                            CASE WHEN jsonb_typeof(cmetadata->'column_header') = 'array'
                                THEN jsonb_array_length(cmetadata->'column_header') ELSE 0 END DESC,
                            CASE WHEN cmetadata->>'col_index' ~ '^\d+$' THEN (cmetadata->>'col_index')::int ELSE 99999 END,
                            id
                    ) AS coord_rank
                FROM langchain_pg_embedding
                WHERE (
                    (cmetadata->>'row_index' ~ '^\d+$' AND (cmetadata->>'row_index')::int = ANY(%s))
                    OR
                    (NULLIF(regexp_replace(cmetadata->>'cell_coord', '[^0-9]', '', 'g'), '') ~ '^\d+$'
                     AND (NULLIF(regexp_replace(cmetadata->>'cell_coord', '[^0-9]', '', 'g'), ''))::int = ANY(%s))
                )
                {scope_sql}
            )
            SELECT id, document, cmetadata, resolved_row_index, resolved_col_index
            FROM filtered_rows
            WHERE coord_rank = 1
            ORDER BY resolved_row_index,
                     cmetadata->>'cell_coord',
                     id;
            """,
            [*cte_params, unique_rows, unique_rows, *scope_params],
        )

    @staticmethod
    def _rows_cells_rows(
        rows: Sequence[Sequence[Any]],
        limit_per_row: int,
    ) -> Dict[int, List[Dict[str, Any]]]:
        from openpyxl.utils.cell import column_index_from_string

        results: Dict[int, List[Dict[str, Any]]] = {}
        for cell_id, document, metadata, resolved_row, resolved_col in rows:
            cmetadata = metadata if isinstance(metadata, dict) else {}
            col_idx = resolved_col
            if col_idx is None and cmetadata.get("cell_coord"):
                col_match = re.match(r"^([A-Za-z]+)", cmetadata.get("cell_coord", ""))
                if col_match:
                    try:
                        col_idx = column_index_from_string(col_match.group(1))
                    except Exception:
                        pass
            results.setdefault(int(resolved_row), []).append(
                {
                    "cell_id": cell_id or cmetadata.get("cell_id", ""),
                    "cell_coord": cmetadata.get("cell_coord", ""),
                    "cell_value": cmetadata.get("cell_value", ""),
                    "sheet_name": cmetadata.get("sheet_name", ""),
                    "row_header": cmetadata.get("row_header", []),
                    "column_header": cmetadata.get("column_header", []),
                    "company_name": cmetadata.get("company_name", ""),
                    "row_index": resolved_row,
                    "col_index": col_idx,
                    "source_text": document or "",
                }
            )
        for row_index, row_cells in results.items():
            row_cells.sort(
                key=lambda cell: (
                    cell.get("col_index") is None,
                    cell.get("col_index") or 0,
                )
            )
            results[row_index] = row_cells[: max(1, limit_per_row)]
        return results

    async def fetch_rows_cells_async(
        self,
        *,
        collection_name: Optional[str],
        workbook_hash: Optional[str],
        sheet_name: Optional[str],
        row_indices: Sequence[int],
        limit_per_row: int = 50,
    ) -> Dict[int, List[Dict[str, Any]]]:
        """Fetch row context with native async PostgreSQL I/O."""
        query_spec = self._rows_cells_query(
            collection_name=collection_name,
            workbook_hash=workbook_hash,
            sheet_name=sheet_name,
            row_indices=row_indices,
        )
        if query_spec is None:
            return {}
        query_sql, params = query_spec
        raw_url = getattr(self, "database_url", PGVECTOR_URL).replace(
            "postgresql+psycopg://",
            "postgresql://",
        )
        try:
            async with get_pooled_async_connection(raw_url) as connection:
                async with connection.cursor() as cursor:
                    await cursor.execute(query_sql, params)  # type: ignore[arg-type]
                    rows = await cursor.fetchall()
            return self._rows_cells_rows(rows, limit_per_row)
        except Exception as error:
            logger.warning("fetch_rows_cells_async 실패: %s", error)
            return {}

    def fetch_rows_cells(
        self,
        collection_name: Optional[str],
        workbook_hash: Optional[str],
        sheet_name: Optional[str],
        row_indices: List[int],
        limit_per_row: int = 50,
    ) -> Dict[int, List[Dict[str, Any]]]:
        """Fetch several rows in one SQL round trip, bounded per requested row."""
        unique_rows = sorted({row for row in row_indices if row > 0})
        if not unique_rows or not (collection_name or workbook_hash):
            return {}
        conn = self._read_connection()
        try:
            with conn.cursor() as cur:
                if collection_name:
                    collection_names = [
                        name.strip() for name in collection_name.split(",") if name.strip()
                    ]
                else:
                    collection_names = []

                scope_clauses: List[str] = []
                scope_params: List[Any] = []
                cte_params: List[Any] = []
                if collection_names:
                    collection_cte_sql = (
                        "target_collections AS ("
                        "SELECT uuid FROM langchain_pg_collection WHERE name = ANY(%s)"
                        "),"
                    )
                    cte_params.append(collection_names)
                    scope_clauses.append("collection_id IN (SELECT uuid FROM target_collections)")
                else:
                    collection_cte_sql = ""
                if workbook_hash and not collection_names:
                    workbook_hashes = [
                        value.strip() for value in workbook_hash.split(",") if value.strip()
                    ]
                    scope_clauses.append("cmetadata->>'workbook_hash' = ANY(%s)")
                    scope_params.append(workbook_hashes)
                if sheet_name:
                    scope_clauses.append("cmetadata->>'sheet_name' = %s")
                    scope_params.append(sheet_name)

                scope_sql = " AND " + " AND ".join(scope_clauses) if scope_clauses else ""

                query_sql = f"""
                    WITH {collection_cte_sql}
                    filtered_rows AS (
                        SELECT
                            id,
                            document,
                            cmetadata,
                            COALESCE(
                                CASE WHEN cmetadata->>'row_index' ~ '^\\d+$' THEN (cmetadata->>'row_index')::int ELSE NULL END,
                                NULLIF(regexp_replace(cmetadata->>'cell_coord', '[^0-9]', '', 'g'), '')::int
                            ) AS resolved_row_index,
                            CASE
                                WHEN cmetadata->>'col_index' ~ '^\\d+$'
                                    THEN (cmetadata->>'col_index')::int
                                ELSE NULL
                            END AS resolved_col_index,
                            ROW_NUMBER() OVER (
                                PARTITION BY COALESCE(
                                    CASE WHEN cmetadata->>'row_index' ~ '^\\d+$' THEN (cmetadata->>'row_index')::int ELSE NULL END,
                                    NULLIF(regexp_replace(cmetadata->>'cell_coord', '[^0-9]', '', 'g'), '')::int
                                ),
                                cmetadata->>'cell_coord'
                                ORDER BY
                                    CASE WHEN cmetadata->>'variant' = 'header_with_value' THEN 0 ELSE 1 END,
                                    CASE WHEN jsonb_typeof(cmetadata->'row_header') = 'array'
                                        THEN jsonb_array_length(cmetadata->'row_header') ELSE 0 END DESC,
                                    CASE WHEN jsonb_typeof(cmetadata->'column_header') = 'array'
                                        THEN jsonb_array_length(cmetadata->'column_header') ELSE 0 END DESC,
                                    CASE WHEN cmetadata->>'col_index' ~ '^\\d+$' THEN (cmetadata->>'col_index')::int ELSE 99999 END,
                                    id
                            ) AS coord_rank
                        FROM langchain_pg_embedding
                        WHERE (
                            (cmetadata->>'row_index' ~ '^\\d+$' AND (cmetadata->>'row_index')::int = ANY(%s))
                            OR
                            (NULLIF(regexp_replace(cmetadata->>'cell_coord', '[^0-9]', '', 'g'), '') ~ '^\\d+$'
                             AND (NULLIF(regexp_replace(cmetadata->>'cell_coord', '[^0-9]', '', 'g'), ''))::int = ANY(%s))
                        )
                        {scope_sql}
                    )
                    SELECT id, document, cmetadata, resolved_row_index, resolved_col_index
                    FROM filtered_rows
                    WHERE coord_rank = 1
                    ORDER BY resolved_row_index,
                             cmetadata->>'cell_coord',
                             id;
                """
                params = [
                    *cte_params,
                    unique_rows,
                    unique_rows,
                    *scope_params,
                ]
                cur.execute(query_sql, tuple(params))
                results: Dict[int, List[Dict[str, Any]]] = {}
                from openpyxl.utils.cell import column_index_from_string

                for cid, doc, meta, resolved_row, resolved_col in cur.fetchall():
                    cmetadata = meta if isinstance(meta, dict) else {}
                    col_idx = resolved_col
                    if col_idx is None and cmetadata.get("cell_coord"):
                        col_match = re.match(r"^([A-Za-z]+)", cmetadata.get("cell_coord", ""))
                        if col_match:
                            try:
                                col_idx = column_index_from_string(col_match.group(1))
                            except Exception:
                                pass

                    results.setdefault(int(resolved_row), []).append(
                        {
                            "cell_id": cid or cmetadata.get("cell_id", ""),
                            "cell_coord": cmetadata.get("cell_coord", ""),
                            "cell_value": cmetadata.get("cell_value", ""),
                            "sheet_name": cmetadata.get("sheet_name", ""),
                            "row_header": cmetadata.get("row_header", []),
                            "column_header": cmetadata.get("column_header", []),
                            "company_name": cmetadata.get("company_name", ""),
                            "row_index": resolved_row,
                            "col_index": col_idx,
                            "source_text": doc or "",
                        }
                    )
                # Sort each row's cells by col_index ascending
                for r_idx, r_cells in results.items():
                    r_cells.sort(
                        key=lambda c: (c.get("col_index") is None, c.get("col_index") or 0)
                    )
                    results[r_idx] = r_cells[: max(1, limit_per_row)]
                return results
        except Exception as err:
            logger.warning("fetch_rows_cells 실패: %s", err)
            return {}
        finally:
            conn.close()

    def list_registered_collections(self) -> List[Dict[str, Any]]:
        """List all indexed pgvector collections and metadata."""
        conn = self._raw_connection()
        try:
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT uuid, name, cmetadata FROM langchain_pg_collection ORDER BY name ASC;"
                )
                return [
                    {
                        "uuid": str(row[0]),
                        "name": str(row[1]),
                        "cmetadata": row[2] if isinstance(row[2], dict) else {},
                    }
                    for row in cur.fetchall()
                ]
        except Exception as err:
            logger.warning(f"list_registered_collections 실패: {err}")
            return []
        finally:
            conn.close()
