"""Session-scoped chat attachment storage and safe text extraction."""

from __future__ import annotations

import csv
import json
import re
import zipfile
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from xml.etree import ElementTree

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook

from backend.core.settings import CHAT_UPLOAD_DIR

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_EXTRACTED_CHARS = 60_000
MAX_SHEET_EXTRACTED_CHARS = 6_000
MAX_EVIDENCE_CHARS = 12_000
SUPPORTED_SUFFIXES = frozenset({".txt", ".md", ".csv", ".json", ".xlsx", ".xlsm"})
_FINANCIAL_TERM_ALIASES = {
    "매출": ("revenue", "sales"),
    "영업이익": ("operating income", "operating profit"),
    "순이익": ("net income", "net earnings"),
    "자산": ("assets", "total assets"),
    "부채": ("liabilities", "debt", "total debt"),
    "자본": ("equity", "shareholders"),
    "현금흐름": ("cash flow", "free cash flow"),
    "이익률": ("margin",),
    "주가": ("share price", "price"),
}


def _limit(text: str) -> str:
    normalized = text.replace("\x00", "").strip()
    if len(normalized) <= MAX_EXTRACTED_CHARS:
        return normalized
    return normalized[:MAX_EXTRACTED_CHARS] + "\n\n[첨부 파일 내용이 길어 처음 60,000자만 사용했습니다.]"


def compact_evidence(question: str, extracted_text: str) -> str:
    """Select question-relevant spreadsheet rows before sending evidence to the LLM."""
    lines = [line.strip() for line in extracted_text.splitlines() if line.strip()]
    terms = {
        term.casefold()
        for term in re.findall(r"[\w가-힣]{2,}", question)
        if term.casefold() not in {"알려줘", "보여줘", "어떻게", "이것", "그것", "파일", "첨부"}
    }
    for korean, aliases in _FINANCIAL_TERM_ALIASES.items():
        if korean in question:
            terms.update(aliases)
    selected: list[str] = []
    current_sheet = ""
    recent: list[str] = []
    # Keep the sheet name close to every matched row so the source stays interpretable.
    for line in lines:
        if line.startswith("[시트:"):
            current_sheet = line
            recent = []
            continue
        lowered = line.casefold()
        score = sum(1 for term in terms if term in lowered)
        if score:
            selected.extend([current_sheet, *recent[-2:], line])
        recent.append(line)

    if not selected:
        # General questions still receive representative rows from each sheet.
        sheet_rows: dict[str, int] = {}
        current_sheet = ""
        for line in lines:
            if line.startswith("[시트:"):
                current_sheet = line
                sheet_rows.setdefault(current_sheet, 0)
                continue
            if current_sheet and sheet_rows[current_sheet] < 4:
                selected.extend([current_sheet, line])
                sheet_rows[current_sheet] += 1

    deduplicated = list(dict.fromkeys(selected))
    result = "\n".join(deduplicated)
    if len(result) > MAX_EVIDENCE_CHARS:
        result = result[:MAX_EVIDENCE_CHARS] + "\n[질문 관련 근거가 길어 일부만 사용했습니다.]"
    return result or extracted_text[:MAX_EVIDENCE_CHARS]


def _xml_text(element: ElementTree.Element | None) -> str:
    if element is None:
        return ""
    return "".join(element.itertext()).strip()


def _extract_xlsx_xml(content: bytes) -> str:
    """Fallback reader for valid xlsx containers openpyxl cannot fully deserialize."""
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    lines: list[str] = []
    with zipfile.ZipFile(BytesIO(content)) as archive:
        shared: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
            shared = [_xml_text(item) for item in root.findall(f"{namespace}si")]
        sheets = sorted(
            (name for name in archive.namelist() if name.startswith("xl/worksheets/") and name.endswith(".xml")),
            key=str.casefold,
        )
        for sheet in sheets:
            lines.append(f"[시트: {Path(sheet).stem}]")
            root = ElementTree.fromstring(archive.read(sheet))
            for row in root.findall(f".//{namespace}row"):
                values: list[str] = []
                for cell in row.findall(f"{namespace}c"):
                    value = _xml_text(cell.find(f"{namespace}v"))
                    cell_type = cell.get("t")
                    if cell_type == "s" and value.isdigit() and int(value) < len(shared):
                        value = shared[int(value)]
                    elif cell_type == "inlineStr":
                        value = _xml_text(cell.find(f"{namespace}is"))
                    if value:
                        values.append(value)
                if values:
                    lines.append(" | ".join(values))
                if sum(len(line) + 1 for line in lines) > MAX_EXTRACTED_CHARS:
                    return _limit("\n".join(lines))
    return _limit("\n".join(lines))


def extract_text(file_name: str, content: bytes) -> str:
    """Extract a bounded, plain-text evidence context from supported local files."""
    suffix = Path(file_name).suffix.casefold()
    if suffix in {".txt", ".md"}:
        return _limit(content.decode("utf-8-sig", errors="replace"))
    if suffix == ".json":
        try:
            return _limit(json.dumps(json.loads(content.decode("utf-8-sig")), ensure_ascii=False, indent=2))
        except (UnicodeDecodeError, json.JSONDecodeError):
            return _limit(content.decode("utf-8-sig", errors="replace"))
    if suffix == ".csv":
        decoded = content.decode("utf-8-sig", errors="replace")
        rows = csv.reader(decoded.splitlines())
        return _limit("\n".join(" | ".join(cell.strip() for cell in row) for row in rows))
    if suffix in {".xlsx", ".xlsm"}:
        try:
            # External workbook links are irrelevant to chat evidence and some
            # valid financial templates contain link metadata openpyxl cannot parse.
            workbook = load_workbook(
                BytesIO(content),
                read_only=True,
                data_only=True,
                keep_links=False,
            )
            lines: list[str] = []
            try:
                for worksheet in workbook.worksheets:
                    # Finance workbooks often keep a vendor/add-in payload in a
                    # veryHidden sheet. It is not user data and can consume the
                    # entire context window with binary-like text.
                    if worksheet.sheet_state != "visible":
                        continue
                    worksheet_lines = [f"[시트: {worksheet.title}]"]
                    for row in worksheet.iter_rows(values_only=True):
                        values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                        if values:
                            worksheet_lines.append(" | ".join(values))
                        if sum(len(line) + 1 for line in worksheet_lines) > MAX_SHEET_EXTRACTED_CHARS:
                            worksheet_lines.append("[이 시트는 처음 6,000자만 사용했습니다.]")
                            break
                    lines.extend(worksheet_lines)
                    if sum(len(line) + 1 for line in lines) > MAX_EXTRACTED_CHARS:
                        return _limit("\n".join(lines))
            finally:
                workbook.close()
            return _limit("\n".join(lines))
        except (TypeError, ValueError):
            return _extract_xlsx_xml(content)
    raise ValueError("지원하지 않는 파일 형식입니다")


async def save_upload(file: UploadFile) -> tuple[str, str, str | None, int, str, str]:
    if not file.filename:
        raise HTTPException(status_code=400, detail="유효한 파일명이 필요합니다")
    safe_name = Path(file.filename).name
    suffix = Path(safe_name).suffix.casefold()
    if suffix not in SUPPORTED_SUFFIXES:
        allowed = ", ".join(sorted(SUPPORTED_SUFFIXES))
        raise HTTPException(status_code=415, detail=f"지원 파일: {allowed}")
    try:
        content = await file.read(MAX_UPLOAD_BYTES + 1)
        if len(content) > MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=413, detail="첨부 파일은 20MB를 초과할 수 없습니다")
        extracted = extract_text(safe_name, content)
        if not extracted:
            raise HTTPException(status_code=422, detail="질문 근거로 사용할 텍스트를 추출하지 못했습니다")
        attachment_id = f"attachment-{uuid4().hex}"
        CHAT_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
        destination = CHAT_UPLOAD_DIR / f"{attachment_id}{suffix}"
        destination.write_bytes(content)
        return attachment_id, safe_name, file.content_type, len(content), str(destination.resolve()), extracted
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(status_code=422, detail=f"첨부 파일을 읽지 못했습니다: {error}") from error
    finally:
        await file.close()
