from datetime import date, datetime
from pathlib import Path
import textwrap
from typing import Dict, List, Optional, Tuple

from PIL import Image, ImageDraw, ImageFont
from openpyxl.styles.numbers import is_date_format

from .table_geometry import SheetLayout, compute_sheet_layout


THEME_COLORS = {
    0: "#FFFFFF",
    1: "#000000",
    2: "#E7E6E6",
    3: "#44546A",
    4: "#5B9BD5",
    5: "#ED7D31",
    6: "#A5A5A5",
    7: "#FFC000",
    8: "#4472C4",
    9: "#70AD47",
}


def _rgb_color(color, default: str) -> str:
    if color is None:
        return default
    if getattr(color, "type", None) == "rgb":
        value = str(getattr(color, "rgb", "")).lstrip("#")
        if len(value) == 8:
            value = value[2:]
        if len(value) == 6:
            return f"#{value}"
    if getattr(color, "type", None) == "theme":
        base = THEME_COLORS.get(getattr(color, "theme", -1), default)
        tint = float(getattr(color, "tint", 0.0) or 0.0)
        if not tint or not base.startswith("#") or len(base) != 7:
            return base
        channels = [int(base[index : index + 2], 16) for index in (1, 3, 5)]
        adjusted = [
            round(channel * (1 + tint))
            if tint < 0
            else round(channel + (255 - channel) * tint)
            for channel in channels
        ]
        return "#" + "".join(f"{max(0, min(255, channel)):02X}" for channel in adjusted)
    return default


def _cell_fill(cell) -> str:
    fill = getattr(cell, "fill", None)
    if fill is None or fill.fill_type in (None, "none"):
        return "#FFFFFF"
    return _rgb_color(fill.fgColor, "#F2F2F2")


def _cell_text(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (date, datetime)) or (
        isinstance(value, (int, float)) and is_date_format(str(cell.number_format))
    ):
        if not isinstance(value, (date, datetime)):
            return str(value)
        return value.strftime("%Y.%m.%d")
    if isinstance(value, float):
        if "%" in str(cell.number_format):
            decimals = max(
                0,
                min(6, str(cell.number_format).split("%", 1)[0].count("0") - 1),
            )
            return f"{value * 100:.{decimals}f}%"
        negative = value < 0 and "(" in str(cell.number_format)
        rendered = f"{abs(value) if negative else value:,.2f}".rstrip("0").rstrip(".")
        return f"({rendered})" if negative else rendered
    if isinstance(value, int):
        negative = value < 0 and "(" in str(cell.number_format)
        rendered = f"{abs(value) if negative else value:,}"
        return f"({rendered})" if negative else rendered
    return str(value).strip()


def _font(size: int, bold: bool, italic: bool = False) -> ImageFont.ImageFont:
    style_candidates = {
        (False, False): (
            "/System/Library/Fonts/Supplemental/Arial.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        ),
        (True, False): (
            "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        ),
        (False, True): (
            "/System/Library/Fonts/Supplemental/Arial Italic.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf",
        ),
        (True, True): (
            "/System/Library/Fonts/Supplemental/Arial Bold Italic.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-BoldOblique.ttf",
        ),
    }
    candidates = (*style_candidates[(bold, italic)], "/System/Library/Fonts/Helvetica.ttc")
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size=size)
        except (OSError, ValueError):
            continue
    return ImageFont.load_default()


def _text_size(
    draw: ImageDraw.ImageDraw,
    text: str,
    font: ImageFont.ImageFont,
) -> Tuple[int, int]:
    box = draw.multiline_textbbox((0, 0), text, font=font, spacing=1)
    return box[2] - box[0], box[3] - box[1]


def _wrap_text(
    draw: ImageDraw.ImageDraw,
    text: str,
    font: ImageFont.ImageFont,
    max_width: float,
) -> str:
    if not text or max_width <= 4:
        return ""
    lines: List[str] = []
    for source_line in text.splitlines() or [text]:
        words = source_line.split()
        if not words:
            lines.append("")
            continue
        current = words[0]
        for word in words[1:]:
            candidate = f"{current} {word}"
            if _text_size(draw, candidate, font)[0] <= max_width:
                current = candidate
            else:
                lines.append(current)
                current = word
        lines.append(current)

    wrapped: List[str] = []
    for line in lines:
        if _text_size(draw, line, font)[0] <= max_width:
            wrapped.append(line)
            continue
        average_width = max(1.0, _text_size(draw, line, font)[0] / max(1, len(line)))
        character_count = max(1, int(max_width / average_width))
        wrapped.extend(
            textwrap.wrap(
                line,
                width=character_count,
                break_long_words=True,
                break_on_hyphens=False,
            )
            or [""]
        )
    return "\n".join(wrapped)


def _fit_text(
    draw: ImageDraw.ImageDraw,
    text: str,
    font_size: int,
    bold: bool,
    italic: bool,
    max_width: float,
    max_height: float,
) -> Tuple[str, ImageFont.ImageFont, int, int]:
    for candidate_size in range(font_size, 4, -1):
        font = _font(candidate_size, bold, italic)
        wrapped = _wrap_text(draw, text, font, max_width)
        width, height = _text_size(draw, wrapped, font)
        if width <= max_width and height <= max_height:
            return wrapped, font, width, height
    font = _font(5, bold, italic)
    wrapped = _wrap_text(draw, text, font, max_width)
    width, height = _text_size(draw, wrapped, font)
    return wrapped, font, width, height


def _border_width(style: Optional[str]) -> int:
    if style in {"medium", "mediumDashed", "mediumDashDot", "mediumDashDotDot"}:
        return 2
    if style in {"thick", "double"}:
        return 3
    return 1


def _merge_map(worksheet) -> Dict[Tuple[int, int], Optional[Tuple[int, int]]]:
    merged: Dict[Tuple[int, int], Optional[Tuple[int, int]]] = {}
    for cell_range in worksheet.merged_cells.ranges:
        for row in range(cell_range.min_row, cell_range.max_row + 1):
            for column in range(cell_range.min_col, cell_range.max_col + 1):
                merged[(row, column)] = (
                    (cell_range.max_row, cell_range.max_col)
                    if row == cell_range.min_row and column == cell_range.min_col
                    else None
                )
    return merged


class ExcelSheetRenderer:
    """Render one styled Excel worksheet to a deterministic PNG image."""

    def render(
        self,
        worksheet,
        output_path: Path,
        max_rows: int,
        max_columns: int,
    ) -> SheetLayout:
        layout = compute_sheet_layout(worksheet, max_rows, max_columns)
        image = Image.new("RGB", (layout.width, layout.height), "white")
        draw = ImageDraw.Draw(image)
        merged = _merge_map(worksheet)

        for row in range(1, layout.max_row + 1):
            if layout.row_heights[row - 1] <= 0:
                continue
            for column in range(1, layout.max_column + 1):
                if layout.column_widths[column - 1] <= 0:
                    continue
                merge_end = merged.get((row, column), (row, column))
                if merge_end is None:
                    continue
                max_row, max_column = merge_end
                max_row = min(max_row, layout.max_row)
                max_column = min(max_column, layout.max_column)
                x1 = layout.x_offsets[column - 1]
                y1 = layout.y_offsets[row - 1]
                x2 = layout.x_offsets[max_column]
                y2 = layout.y_offsets[max_row]
                cell = worksheet.cell(row=row, column=column)

                draw.rectangle((x1, y1, x2, y2), fill=_cell_fill(cell))
                border = getattr(cell, "border", None)
                if border:
                    if border.top and border.top.style:
                        draw.line(
                            (x1, y1, x2, y1),
                            fill=_rgb_color(border.top.color, "#64748B"),
                            width=_border_width(border.top.style),
                        )
                    if border.bottom and border.bottom.style:
                        draw.line(
                            (x1, y2, x2, y2),
                            fill=_rgb_color(border.bottom.color, "#64748B"),
                            width=_border_width(border.bottom.style),
                        )
                    if border.left and border.left.style:
                        draw.line(
                            (x1, y1, x1, y2),
                            fill=_rgb_color(border.left.color, "#64748B"),
                            width=_border_width(border.left.style),
                        )
                    if border.right and border.right.style:
                        draw.line(
                            (x2, y1, x2, y2),
                            fill=_rgb_color(border.right.color, "#64748B"),
                            width=_border_width(border.right.style),
                        )

                text = _cell_text(cell)
                if not text:
                    continue
                font_size = max(6, min(18, round(float(cell.font.sz or 10))))
                text, font, text_width, text_height = _fit_text(
                    draw,
                    text,
                    font_size,
                    bool(cell.font.bold),
                    bool(cell.font.italic),
                    max(1.0, x2 - x1 - 6),
                    max(1.0, y2 - y1 - 4),
                )
                alignment = getattr(cell.alignment, "horizontal", None)
                if alignment == "right" or isinstance(cell.value, (int, float)):
                    text_x = x2 - text_width - 3
                elif alignment in {"center", "centerContinuous"}:
                    text_x = x1 + ((x2 - x1) - text_width) / 2
                else:
                    text_x = x1 + 3
                vertical = getattr(cell.alignment, "vertical", None)
                if vertical == "top":
                    text_y = y1 + 2
                elif vertical == "bottom":
                    text_y = y2 - text_height - 2
                else:
                    text_y = y1 + max(1.0, ((y2 - y1) - text_height) / 2)
                draw.multiline_text(
                    (text_x, text_y),
                    text,
                    fill=_rgb_color(cell.font.color, "#111827"),
                    font=font,
                    spacing=1,
                    align="center" if alignment in {"center", "centerContinuous"} else "left",
                )
                if cell.font.underline:
                    underline_y = min(y2 - 1, text_y + text_height)
                    draw.line(
                        (text_x, underline_y, text_x + text_width, underline_y),
                        fill=_rgb_color(cell.font.color, "#111827"),
                        width=1,
                    )

        output_path.parent.mkdir(parents=True, exist_ok=True)
        image.save(output_path, format="PNG", dpi=(150, 150))
        return layout
