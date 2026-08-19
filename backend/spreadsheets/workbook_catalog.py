import hashlib
from pathlib import Path
from typing import List

import openpyxl

from .cell_visibility import worksheet_visible


SUPPORTED_WORKBOOK_SUFFIXES = {".xlsx", ".xlsm"}
SKIP_SHEET_PREFIXES = (
    "__snloffice",
    "___snloffice",
    "_ciqhidden",
    "snl",
    "intermediate",
    "chart_data",
)


class WorkbookCatalogError(ValueError):
    """Raised when a requested processed workbook is unavailable or unsafe."""


class WorkbookCatalog:
    """Discover and securely resolve workbooks under one processed-data root."""

    def __init__(self, root: Path) -> None:
        self.root = root.resolve()

    def file_names(self) -> List[str]:
        if not self.root.is_dir():
            return []
        return sorted(
            path.name
            for path in self.root.iterdir()
            if path.is_file()
            and not path.name.startswith("~$")
            and path.suffix.lower() in SUPPORTED_WORKBOOK_SUFFIXES
        )

    def resolve(self, file_name: str) -> Path:
        selected = file_name.strip()
        available = self.file_names()
        if not selected:
            if len(available) == 1:
                selected = available[0]
            else:
                raise WorkbookCatalogError("사용할 processed Excel 파일을 선택해야 합니다")
        if selected not in available:
            raise WorkbookCatalogError(
                f"data/processed에서 선택할 수 없는 Excel 파일입니다: {selected}"
            )
        path = (self.root / selected).resolve()
        if path.parent != self.root:
            raise WorkbookCatalogError("processed 디렉터리 밖의 파일은 선택할 수 없습니다")
        return path

    @staticmethod
    def sha256(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as workbook:
            for chunk in iter(lambda: workbook.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()

    @staticmethod
    def sheet_names(path: Path) -> List[str]:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            sheet_names: List[str] = []
            for sheet in workbook.worksheets:
                if not worksheet_visible(sheet) or sheet.title.lower().startswith(
                    SKIP_SHEET_PREFIXES
                ):
                    continue

                # Some valid producers omit the worksheet ``dimension`` element.
                # In openpyxl read-only mode that leaves max_row/max_column unset
                # until the bounds are calculated from the sheet data.
                if sheet.max_row is None or sheet.max_column is None:
                    sheet.calculate_dimension(force=True)

                if sheet.max_row and sheet.max_column:
                    sheet_names.append(sheet.title)
            return sheet_names
        finally:
            workbook.close()
