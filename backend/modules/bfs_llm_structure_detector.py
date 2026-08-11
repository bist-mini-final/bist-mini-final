from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, cast

import openpyxl
from pydantic import BaseModel, Field

from ..chat_completion import ChatCompletionClient, ChatCompletionError
from ..config import PROCESSED_DATA_DIR, SPREADSHEET_ARTIFACT_DIR
from ..spreadsheets.cell_visibility import WorksheetVisibility, worksheet_visible
from ..spreadsheets.grid_structure import (
    build_column_header_tree,
    detect_table_bounds,
    header_candidate_rows,
    validate_title_end,
)
from ..spreadsheets.prompt_guidance import TEXT_CELL_ROLE_GUIDANCE
from ..spreadsheets.sheet_renderer import ExcelSheetRenderer
from ..spreadsheets.table_geometry import CellBounds, cell_bounds_bbox, compute_sheet_layout
from ..spreadsheets.workbook_catalog import WorkbookCatalog, WorkbookCatalogError
from .base import ExecutableModule, ModuleDefinition, ModuleDTO, ModuleExecutionError
from .docling_table_detector import _safe_name
from .processed_file_selector import WorkbookSelectionDTO
from .spreadsheet_structure import SpreadsheetStructureOutput


BOUNDARY_SYSTEM_PROMPT = f"""You are an expert in Excel table structure analysis.
Hidden sheets, rows, and columns are intentionally absent. Never infer or restore them.
{TEXT_CELL_ROLE_GUIDANCE}

For every supplied table region, decide only these boundaries from its top candidate rows:
1. title_row_end: absolute last row of conceptual table titles, or null. A title describes the whole table. A row with values in two or more different columns is a header, not a title.
2. data_start_row: absolute first data row. Consecutive text rows are usually headers; the first row where numeric/formula values become the main content is usually data.
3. index_column_count: count of leftmost label/category/index columns. Text labels on the left with numeric data to the right indicate index columns. It must be at least 0 and strictly smaller than column_count.

Return one decision for every region_id. Use absolute Excel coordinates. Do not infer any cells not shown. Output only valid JSON in this exact shape:
{{"decisions":[{{"region_id":"...","title_row_end":null,"data_start_row":1,"index_column_count":1}}]}}"""


BOUNDARY_USER_TEMPLATE = """Analyze the following BFS-detected Excel table regions. Empty/error cells are represented as null. Merged-cell coordinates are marked explicitly.
{regions_json}"""


class BfsLlmStructureDetectorInput(WorkbookSelectionDTO):
    model: str = Field(default="gpt-5.6-luna", description="표 경계 판단에 사용할 LLM ID")
    max_rows: int = Field(default=400, ge=1, le=2000, description="시트에서 분석할 최대 행 수")
    max_columns: int = Field(default=60, ge=1, le=200, description="시트에서 분석할 최대 열 수")
    merge_gap: int = Field(default=2, ge=0, le=5, description="BFS 영역을 병합할 빈 셀 간격")
    min_non_empty_cells: int = Field(default=2, ge=1, le=100, description="표 후보로 유지할 최소 비어 있지 않은 셀 수")
    min_table_columns: int = Field(default=2, ge=1, le=20, description="직렬화 대상 표 후보의 최소 열 수")
    header_candidate_rows: int = Field(default=10, ge=1, le=30, description="LLM 경계 판정에 전달할 표 상단 행 수")
    llm_batch_size: int = Field(default=8, ge=1, le=20, description="한 LLM 요청에서 함께 판정할 표 후보 수")
    system_prompt: str = Field(default=BOUNDARY_SYSTEM_PROMPT, description="제목·헤더·데이터 경계 판단 지시")
    user_prompt_template: str = Field(default=BOUNDARY_USER_TEMPLATE, description="{regions_json} 변수를 지원하는 사용자 프롬프트")


class BoundaryDecisionDTO(ModuleDTO):
    region_id: str
    title_row_end: Optional[int]
    data_start_row: int
    index_column_count: int


class BoundaryBatchDTO(ModuleDTO):
    decisions: List[BoundaryDecisionDTO]


class BfsLlmStructureDetectorOutput(SpreadsheetStructureOutput):
    pass


def _strip_code_fence(value: str) -> str:
    text = value.strip()
    if not text.startswith("```"):
        return text
    lines = text.splitlines()[1:]
    if lines and lines[-1].strip() == "```":
        lines.pop()
    return "\n".join(lines).strip()


class BfsLlmStructureDetectorModule(ExecutableModule):
    definition = ModuleDefinition(
        type="bfs_llm_structure_detector",
        label="BFS + LLM Table Structure Detector",
        category="Logic",
        description="셀 연결요소로 표를 분리하고 상단 행만 LLM으로 판단해 제목·계층 헤더·데이터 영역을 구성합니다.",
        inputs=["input"],
        outputs=["output"],
        config_fields=[
            "model",
            "max_rows",
            "max_columns",
            "merge_gap",
            "min_non_empty_cells",
            "min_table_columns",
            "header_candidate_rows",
            "llm_batch_size",
            "system_prompt",
            "user_prompt_template",
        ],
        raw_output=True,
        version="3",
    )
    input_model = BfsLlmStructureDetectorInput
    output_model = BfsLlmStructureDetectorOutput

    def __init__(
        self,
        completion_client: Optional[ChatCompletionClient] = None,
        catalog: WorkbookCatalog | None = None,
        renderer: ExcelSheetRenderer | None = None,
        processed_dir: Path = PROCESSED_DATA_DIR,
        artifact_dir: Path = SPREADSHEET_ARTIFACT_DIR,
    ) -> None:
        self.completion_client = completion_client or ChatCompletionClient()
        self.catalog = catalog or WorkbookCatalog(processed_dir)
        self.renderer = renderer or ExcelSheetRenderer()
        self.artifact_dir = artifact_dir

    def _decide_boundaries(
        self,
        regions: List[Dict[str, Any]],
        settings: BfsLlmStructureDetectorInput,
    ) -> Dict[str, BoundaryDecisionDTO]:
        decisions: Dict[str, BoundaryDecisionDTO] = {}
        for offset in range(0, len(regions), settings.llm_batch_size):
            batch = regions[offset : offset + settings.llm_batch_size]
            prompt = settings.user_prompt_template.replace(
                "{regions_json}",
                json.dumps(batch, ensure_ascii=False, separators=(",", ":")),
            )
            try:
                response = self.completion_client.complete_structured(
                    settings.model,
                    [
                        {"role": "system", "content": settings.system_prompt},
                        {"role": "user", "content": prompt},
                    ],
                    "excel_table_boundaries",
                    BoundaryBatchDTO.model_json_schema(),
                )
                parsed = BoundaryBatchDTO.model_validate_json(_strip_code_fence(response))
            except ChatCompletionError as error:
                raise ModuleExecutionError(str(error)) from error
            except (ValueError, TypeError, json.JSONDecodeError) as error:
                raise ModuleExecutionError(
                    "LLM 표 경계 응답이 유효한 decisions JSON이 아닙니다"
                ) from error

            expected = {str(region["region_id"]) for region in batch}
            returned = {decision.region_id for decision in parsed.decisions}
            if expected != returned or len(returned) != len(parsed.decisions):
                raise ModuleExecutionError(
                    "LLM 표 경계 응답의 region_id가 요청한 표 후보와 일치하지 않습니다"
                )
            decisions.update({decision.region_id: decision for decision in parsed.decisions})
        return decisions

    @staticmethod
    def _validate_decision(
        decision: BoundaryDecisionDTO,
        bounds: CellBounds,
        visibility: WorksheetVisibility,
    ) -> None:
        if not bounds.min_row <= decision.data_start_row <= bounds.max_row:
            raise ModuleExecutionError(
                f"{decision.region_id}의 data_start_row가 표 범위를 벗어났습니다"
            )
        if visibility.row_hidden(decision.data_start_row):
            raise ModuleExecutionError(
                f"{decision.region_id}의 data_start_row가 숨겨진 행을 가리킵니다"
            )
        if decision.title_row_end is not None and not (
            bounds.min_row <= decision.title_row_end < decision.data_start_row
        ):
            raise ModuleExecutionError(
                f"{decision.region_id}의 title_row_end가 올바른 경계가 아닙니다"
            )
        if (
            decision.title_row_end is not None
            and visibility.row_hidden(decision.title_row_end)
        ):
            raise ModuleExecutionError(
                f"{decision.region_id}의 title_row_end가 숨겨진 행을 가리킵니다"
            )
        width = sum(
            1
            for column in range(bounds.min_column, bounds.max_column + 1)
            if not visibility.column_hidden(column)
        )
        if not 0 <= decision.index_column_count < width:
            raise ModuleExecutionError(
                f"{decision.region_id}의 index_column_count가 표 너비와 맞지 않습니다"
            )

    @staticmethod
    def _region(
        region_id: str,
        region_type: str,
        bounds: CellBounds,
        layout,
        confidence: float,
        parent_ids: List[str],
    ) -> Dict[str, Any]:
        return {
            "region_id": region_id,
            "type": region_type,
            "excel_range": bounds.excel_range,
            "bbox_px": cell_bounds_bbox(bounds, layout),
            "rows": (bounds.min_row, bounds.max_row),
            "columns": (bounds.min_column, bounds.max_column),
            "confidence": confidence,
            "parent_ids": parent_ids,
        }

    def execute(self, payload: BaseModel) -> Dict[str, Any]:
        settings = cast(BfsLlmStructureDetectorInput, payload)
        try:
            workbook_path = self.catalog.resolve(settings.file_name)
            current_hash = self.catalog.sha256(workbook_path)
        except (OSError, ValueError, WorkbookCatalogError) as error:
            raise ModuleExecutionError(str(error)) from error
        if current_hash != settings.workbook_hash:
            raise ModuleExecutionError(
                "선택 이후 Excel 파일이 변경되었습니다. 파일 선택 모듈을 다시 실행하세요"
            )

        formula_workbook = None
        value_workbook = None
        try:
            common_options = {
                "read_only": False,
                "keep_vba": workbook_path.suffix.lower() == ".xlsm",
            }
            formula_workbook = openpyxl.load_workbook(
                workbook_path,
                data_only=False,
                **common_options,
            )
            value_workbook = openpyxl.load_workbook(
                workbook_path,
                data_only=True,
                **common_options,
            )
            candidates: List[Tuple[str, int, CellBounds, Any]] = []
            prompt_regions: List[Dict[str, Any]] = []
            layouts: Dict[str, Any] = {}
            output_root = self.artifact_dir / current_hash[:16] / "rendered"

            for sheet_name in settings.sheet_names:
                if sheet_name not in formula_workbook.sheetnames:
                    raise ModuleExecutionError(f"Excel 시트를 찾을 수 없습니다: {sheet_name}")
                formula_sheet = formula_workbook[sheet_name]
                value_sheet = value_workbook[sheet_name]
                if not worksheet_visible(formula_sheet) or not worksheet_visible(value_sheet):
                    raise ModuleExecutionError(
                        f"숨겨진 Excel 시트는 분석할 수 없습니다: {sheet_name}"
                    )
                visibility = WorksheetVisibility.from_worksheet(formula_sheet)
                layout = self.renderer.render(
                    value_sheet,
                    output_root / f"{_safe_name(sheet_name)}.png",
                    settings.max_rows,
                    settings.max_columns,
                )
                layouts[sheet_name] = layout
                bounds_list = detect_table_bounds(
                    formula_sheet,
                    settings.max_rows,
                    settings.max_columns,
                    settings.merge_gap,
                    settings.min_non_empty_cells,
                )
                bounds_list = [
                    bounds
                    for bounds in bounds_list
                    if sum(
                        1
                        for column in range(bounds.min_column, bounds.max_column + 1)
                        if not visibility.column_hidden(column)
                    )
                    >= settings.min_table_columns
                ]
                for table_index, bounds in enumerate(bounds_list, start=1):
                    visible_rows = [
                        row
                        for row in range(bounds.min_row, bounds.max_row + 1)
                        if not visibility.row_hidden(row)
                    ]
                    visible_columns = [
                        column
                        for column in range(bounds.min_column, bounds.max_column + 1)
                        if not visibility.column_hidden(column)
                    ]
                    region_id = f"{sheet_name}:table_{table_index}"
                    candidates.append((sheet_name, table_index, bounds, value_sheet))
                    prompt_regions.append(
                        {
                            "region_id": region_id,
                            "sheet_name": sheet_name,
                            "excel_range": bounds.excel_range,
                            "row_count": len(visible_rows),
                            "column_count": len(visible_columns),
                            "candidate_rows": header_candidate_rows(
                                value_sheet,
                                formula_sheet,
                                bounds,
                                settings.header_candidate_rows,
                            ),
                        }
                    )

            decisions = self._decide_boundaries(prompt_regions, settings)
            output_tables: List[Dict[str, Any]] = []
            for sheet_name, table_index, bounds, value_sheet in candidates:
                region_key = f"{sheet_name}:table_{table_index}"
                decision = decisions[region_key]
                visibility = WorksheetVisibility.from_worksheet(value_sheet)
                self._validate_decision(decision, bounds, visibility)
                visible_rows = [
                    row
                    for row in range(bounds.min_row, bounds.max_row + 1)
                    if not visibility.row_hidden(row)
                ]
                visible_columns = [
                    column
                    for column in range(bounds.min_column, bounds.max_column + 1)
                    if not visibility.column_hidden(column)
                ]
                title_end = validate_title_end(
                    value_sheet,
                    bounds,
                    decision.title_row_end,
                )
                header_rows = [
                    row
                    for row in visible_rows
                    if (title_end is None or row > title_end)
                    and row < decision.data_start_row
                ]
                header_start = header_rows[0] if header_rows else decision.data_start_row
                header_end = header_rows[-1] if header_rows else decision.data_start_row - 1
                data_start_column = visible_columns[decision.index_column_count]
                layout = layouts[sheet_name]
                prefix = f"table_{table_index}"
                regions: List[Dict[str, Any]] = []
                parent_ids: List[str] = []

                if title_end is not None:
                    title_id = f"{prefix}_title"
                    regions.append(
                        self._region(
                            title_id,
                            "title",
                            CellBounds(
                                visible_rows[0],
                                title_end,
                                visible_columns[0],
                                visible_columns[-1],
                            ),
                            layout,
                            0.85,
                            [],
                        )
                    )
                    parent_ids.append(title_id)

                if header_start <= header_end:
                    header_id = f"{prefix}_column_header"
                    regions.append(
                        self._region(
                            header_id,
                            "column_header",
                            CellBounds(
                                header_start,
                                header_end,
                                visible_columns[0],
                                visible_columns[-1],
                            ),
                            layout,
                            0.85,
                            list(parent_ids),
                        )
                    )
                    parent_ids.append(header_id)

                row_header_id = None
                if decision.index_column_count > 0:
                    row_header_id = f"{prefix}_row_header"
                    regions.append(
                        self._region(
                            row_header_id,
                            "row_header",
                            CellBounds(
                                decision.data_start_row,
                                visible_rows[-1],
                                visible_columns[0],
                                visible_columns[decision.index_column_count - 1],
                            ),
                            layout,
                            0.85,
                            list(parent_ids),
                        )
                    )

                data_parents = list(parent_ids)
                if row_header_id is not None:
                    data_parents.append(row_header_id)
                regions.append(
                    self._region(
                        f"{prefix}_data",
                        "data",
                        CellBounds(
                            decision.data_start_row,
                            visible_rows[-1],
                            data_start_column,
                            visible_columns[-1],
                        ),
                        layout,
                        0.90,
                        data_parents,
                    )
                )

                output_tables.append(
                    {
                        "sheet_name": sheet_name,
                        "table_index": table_index,
                        "excel_range": CellBounds(
                            visible_rows[0],
                            visible_rows[-1],
                            visible_columns[0],
                            visible_columns[-1],
                        ).excel_range,
                        "regions": regions,
                        "header_tree": build_column_header_tree(
                            value_sheet,
                            CellBounds(
                                visible_rows[0],
                                visible_rows[-1],
                                visible_columns[0],
                                visible_columns[-1],
                            ),
                            header_start,
                            header_end,
                            data_start_column,
                        ),
                    }
                )
        finally:
            if formula_workbook is not None:
                formula_workbook.close()
            if value_workbook is not None:
                value_workbook.close()

        return {
            "file_name": workbook_path.name,
            "workbook_hash": current_hash,
            "tables": output_tables,
        }
