"""Exhaustive whole-sheet table structure analysis with GPT-5.6 Luna vision."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Dict, List, Literal, Optional, Protocol, cast

import openpyxl
from pydantic import BaseModel, Field, model_validator

from ..config import PROCESSED_DATA_DIR, SPREADSHEET_ARTIFACT_DIR
from ..openai_responses_vision import (
    OpenAIResponsesVisionClient,
    OpenAIResponsesVisionError,
    OpenAIResponsesVisionResult,
)
from ..spreadsheets.cell_semantics import collect_non_empty_cells, compact_sheet_context
from ..spreadsheets.cell_type_overlay import render_cell_type_overlay
from ..spreadsheets.cell_visibility import WorksheetVisibility, worksheet_visible
from ..spreadsheets.prompt_guidance import TEXT_CELL_ROLE_GUIDANCE
from ..spreadsheets.sheet_renderer import ExcelSheetRenderer
from ..spreadsheets.table_fragment_merge import parse_excel_range
from ..spreadsheets.table_geometry import CellBounds, SheetLayout
from ..spreadsheets.workbook_catalog import WorkbookCatalog, WorkbookCatalogError
from .base import ExecutableModule, ModuleDefinition, ModuleDTO, ModuleExecutionError
from .docling_table_detector import _safe_name
from .local_vlm_structure_detector import (
    LocalVlmStructureDetectorModule,
    LocalVlmTableDecisionDTO,
    _contains,
)
from .processed_file_selector import WorkbookSelectionDTO
from .spreadsheet_structure import SpreadsheetStructureOutput


LUNA_VLM_SYSTEM_PROMPT = f"""You identify every table and its internal regions in an Excel worksheet.
Hidden sheets, hidden rows, and hidden columns are intentionally absent and must never be inferred.

You receive exactly one high-detail image containing the complete visible worksheet. It is never a crop, tile, or preselected candidate region. Inspect the entire image globally before deciding table boundaries, including sparse areas and tables separated by blank rows or columns.

The image preserves the original worksheet style while populated cells are tinted and labeled by exact Excel coordinate. Colors are: text amber, number green, date blue, boolean purple, error red, unresolved formula gray. A magenta inner border marks a formula. Exact coordinate, type, value, formula, and merged-range facts for the complete sheet are supplied as text and are authoritative.

{TEXT_CELL_ROLE_GUIDANCE}

Return every independent rectangular table in the worksheet. A table must contain data cells; do not emit decorations, isolated notes, or empty rectangles. Every returned coordinate must be inside sheet_range. Use:
- excel_range: the complete visible table.
- title_range: title rows above this table's data, or null.
- column_header_range: every visible column-header level above the data, or null.
- row_header_range: label columns left of the data, or null.
- data_range: the complete value matrix.

All ranges use inclusive Excel coordinates and semantic regions must not overlap. For example, when row_header_range ends at D, data_range must start at E or later. Never repeat a row-header column inside data_range.

Use the full-sheet context to keep one logical table whole instead of splitting it at image boundaries. Do not propose search candidates and do not omit a table because its location is uncertain. Return only JSON matching the supplied schema."""


LUNA_VLM_USER_TEMPLATE = """Analyze the complete visible worksheet {sheet_name}.
Complete visible sheet range: {sheet_range}
The compact tuple format is [excel_coord, value_type, value, optional_flags].
Exact full-sheet context:
{sheet_context}"""


LUNA_SHEET_RESPONSE_SCHEMA: Dict[str, Any] = {
    "type": "object",
    "properties": {
        "tables": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "excel_range": {"type": "string"},
                    "title_range": {"type": ["string", "null"]},
                    "column_header_range": {"type": ["string", "null"]},
                    "row_header_range": {"type": ["string", "null"]},
                    "data_range": {"type": "string"},
                    "confidence": {"type": "number", "minimum": 0, "maximum": 1},
                },
                "required": [
                    "excel_range",
                    "title_range",
                    "column_header_range",
                    "row_header_range",
                    "data_range",
                    "confidence",
                ],
                "additionalProperties": False,
            },
        }
    },
    "required": ["tables"],
    "additionalProperties": False,
}


class LunaVisionClient(Protocol):
    def complete_structured(
        self,
        *,
        model: str,
        system_prompt: str,
        user_prompt: str,
        image_path: Path,
        schema_name: str,
        json_schema: Dict[str, Any],
        reasoning_effort: Literal["none", "low", "medium", "high"],
        max_output_tokens: int,
        timeout_seconds: int,
    ) -> OpenAIResponsesVisionResult | str: ...


class LunaVlmStructureDetectorInput(WorkbookSelectionDTO):
    model: str = Field(
        default="gpt-5.6-luna",
        min_length=1,
        description="OpenAI Responses API 멀티모달 모델 ID",
    )
    max_rows: int = Field(default=400, ge=1, le=2000, description="시트에서 분석할 최대 행 수")
    max_columns: int = Field(default=60, ge=1, le=200, description="시트에서 분석할 최대 열 수")
    max_context_cells: int = Field(
        default=50000,
        ge=100,
        le=100000,
        description="한 시트에서 좌표 컨텍스트로 전달할 최대 값 셀 수",
    )
    reasoning_effort: Literal["none", "low", "medium", "high"] = Field(
        default="low",
        description="Luna 추론 강도",
    )
    max_output_tokens: int = Field(default=6000, ge=1000, le=32000, description="시트별 최대 출력 토큰")
    timeout_seconds: int = Field(default=240, ge=30, le=900, description="시트별 API 요청 제한 시간(초)")
    validation_retries: int = Field(default=1, ge=0, le=2, description="좌표 규칙 위반 응답의 교정 재시도 횟수")
    system_prompt: str = Field(default=LUNA_VLM_SYSTEM_PROMPT, min_length=1, description="전체 시트 구조 식별 시스템 프롬프트")
    user_prompt_template: str = Field(
        default=LUNA_VLM_USER_TEMPLATE,
        min_length=1,
        description="sheet_name, sheet_range, sheet_context 변수를 지원하는 전체 시트 프롬프트",
    )

    @model_validator(mode="before")
    @classmethod
    def discard_legacy_tiling_settings(cls, value: Any) -> Any:
        """Accept workflows saved before whole-sheet inference replaced tiling."""

        if not isinstance(value, dict):
            return value
        migrated = dict(value)
        for field_name in (
            "tile_rows",
            "tile_columns",
            "row_overlap",
            "column_overlap",
            "overview_max_edge",
            "max_concurrency",
        ):
            migrated.pop(field_name, None)
        legacy_system_prompt = migrated.get("system_prompt")
        if (
            isinstance(legacy_system_prompt, str)
            and "You receive two images in this order" in legacy_system_prompt
            and "current tile" in legacy_system_prompt
        ):
            migrated.pop("system_prompt", None)
        legacy_user_prompt = migrated.get("user_prompt_template")
        if (
            isinstance(legacy_user_prompt, str)
            and "{tile_index}" in legacy_user_prompt
            and "{tile_context}" in legacy_user_prompt
        ):
            migrated.pop("user_prompt_template", None)
        return migrated


class LunaSheetDecisionDTO(ModuleDTO):
    tables: List[LocalVlmTableDecisionDTO]


class LunaVlmStructureDetectorOutput(SpreadsheetStructureOutput):
    pass


def _replace_prompt_variables(template: str, values: Dict[str, str]) -> str:
    result = template
    for name, value in values.items():
        result = result.replace("{" + name + "}", value)
    return result


def _validated_sheet_decision(
    table: LocalVlmTableDecisionDTO,
    sheet_bounds: CellBounds,
) -> LocalVlmTableDecisionDTO:
    whole = parse_excel_range(table.excel_range)
    if not _contains(sheet_bounds, whole):
        raise ModuleExecutionError(
            f"Luna 응답 범위가 시트 {sheet_bounds.excel_range}를 벗어났습니다: {table.excel_range}"
        )
    for field_name in (
        "title_range",
        "column_header_range",
        "row_header_range",
        "data_range",
    ):
        value = getattr(table, field_name)
        if value and not _contains(whole, parse_excel_range(value)):
            raise ModuleExecutionError(
                f"Luna {field_name}이 excel_range를 벗어났습니다: {value}"
            )
    return table


class LunaVlmStructureDetectorModule(ExecutableModule):
    definition = ModuleDefinition(
        type="luna_vlm_structure_detector",
        label="Luna Full-Sheet Structure Detector",
        category="Logic",
        description="후보 영역이나 타일 분할 없이 표시된 시트 전체 이미지와 좌표 컨텍스트를 한 번에 분석합니다.",
        inputs=["input"],
        outputs=["output"],
        config_fields=[
            "model",
            "max_rows",
            "max_columns",
            "max_context_cells",
            "reasoning_effort",
            "max_output_tokens",
            "timeout_seconds",
            "validation_retries",
            "system_prompt",
            "user_prompt_template",
        ],
        raw_output=True,
        version="4",
    )
    input_model = LunaVlmStructureDetectorInput
    output_model = LunaVlmStructureDetectorOutput

    def __init__(
        self,
        vision_client: Optional[LunaVisionClient] = None,
        catalog: Optional[WorkbookCatalog] = None,
        renderer: Optional[ExcelSheetRenderer] = None,
        processed_dir: Path = PROCESSED_DATA_DIR,
        artifact_dir: Path = SPREADSHEET_ARTIFACT_DIR,
    ) -> None:
        self.vision_client = vision_client or OpenAIResponsesVisionClient()
        self.catalog = catalog or WorkbookCatalog(processed_dir)
        self.renderer = renderer or ExcelSheetRenderer()
        self.artifact_dir = artifact_dir
        # Reuse the canonical, coordinate-validated structure assembler used
        # by the existing local VLM module so downstream DTOs stay identical.
        self.structure_assembler = LocalVlmStructureDetectorModule(
            catalog=self.catalog,
            renderer=self.renderer,
            processed_dir=processed_dir,
            artifact_dir=artifact_dir,
        )

    def _analyze_sheet(
        self,
        settings: LunaVlmStructureDetectorInput,
        sheet_name: str,
        sheet_bounds: CellBounds,
        image_path: Path,
        layout: SheetLayout,
        visibility: WorksheetVisibility,
        cells: List[Dict[str, Any]],
    ) -> List[LocalVlmTableDecisionDTO]:
        context = compact_sheet_context(sheet_name, layout, cells)
        prompt = _replace_prompt_variables(
            settings.user_prompt_template,
            {
                "sheet_name": sheet_name,
                "sheet_range": sheet_bounds.excel_range,
                "sheet_context": json.dumps(
                    context,
                    ensure_ascii=False,
                    separators=(",", ":"),
                ),
            },
        )
        previous_response = ""
        validation_error = ""
        for attempt in range(settings.validation_retries + 1):
            correction = ""
            if attempt:
                correction = (
                    "\n\nThe previous JSON violated the whole-sheet coordinate rules: "
                    f"{validation_error}\nPrevious JSON: {previous_response}\n"
                    "Return a corrected complete object. All ranges must stay inside "
                    f"{sheet_bounds.excel_range}."
                )
            try:
                response = self.vision_client.complete_structured(
                    model=settings.model,
                    system_prompt=settings.system_prompt,
                    user_prompt=prompt + correction,
                    image_path=image_path,
                    schema_name="luna_spreadsheet_sheet",
                    json_schema=LUNA_SHEET_RESPONSE_SCHEMA,
                    reasoning_effort=settings.reasoning_effort,
                    max_output_tokens=settings.max_output_tokens,
                    timeout_seconds=settings.timeout_seconds,
                )
                previous_response = (
                    response.content
                    if isinstance(response, OpenAIResponsesVisionResult)
                    else str(response)
                )
                decision = LunaSheetDecisionDTO.model_validate_json(previous_response)
                tables = [
                    _validated_sheet_decision(table, sheet_bounds)
                    for table in decision.tables
                ]
                for table in tables:
                    self.structure_assembler._validate_table(table, layout, visibility)
                return tables
            except OpenAIResponsesVisionError:
                raise
            except (ValueError, TypeError, json.JSONDecodeError, ModuleExecutionError) as error:
                validation_error = str(error)
                if attempt >= settings.validation_retries:
                    raise ModuleExecutionError(
                        f"Luna 시트 {sheet_name} 응답을 좌표 규칙에 맞게 교정하지 못했습니다: {validation_error}"
                    ) from error
        return []

    def execute(self, payload: BaseModel) -> Dict[str, Any]:
        settings = cast(LunaVlmStructureDetectorInput, payload)
        try:
            workbook_path = self.catalog.resolve(settings.file_name)
            current_hash = self.catalog.sha256(workbook_path)
        except (OSError, ValueError, WorkbookCatalogError) as error:
            raise ModuleExecutionError(str(error)) from error
        if current_hash != settings.workbook_hash:
            raise ModuleExecutionError("선택 이후 Excel 파일이 변경되었습니다. 파일 선택 모듈을 다시 실행하세요")

        formula_workbook = None
        value_workbook = None
        try:
            options = {
                "read_only": False,
                "keep_vba": workbook_path.suffix.lower() == ".xlsm",
            }
            formula_workbook = openpyxl.load_workbook(workbook_path, data_only=False, **options)
            value_workbook = openpyxl.load_workbook(workbook_path, data_only=True, **options)
            outputs: List[Dict[str, Any]] = []
            output_root = self.artifact_dir / current_hash[:16]

            for sheet_name in settings.sheet_names:
                if sheet_name not in formula_workbook.sheetnames:
                    raise ModuleExecutionError(f"Excel 시트를 찾을 수 없습니다: {sheet_name}")
                formula_sheet = formula_workbook[sheet_name]
                value_sheet = value_workbook[sheet_name]
                if not worksheet_visible(formula_sheet) or not worksheet_visible(value_sheet):
                    raise ModuleExecutionError(f"숨겨진 Excel 시트는 분석할 수 없습니다: {sheet_name}")

                safe_sheet = _safe_name(sheet_name)
                rendered_path = output_root / "rendered" / f"{safe_sheet}.png"
                typed_path = output_root / "typed" / f"{safe_sheet}.png"
                layout = self.renderer.render(
                    value_sheet,
                    rendered_path,
                    settings.max_rows,
                    settings.max_columns,
                )
                try:
                    cells = collect_non_empty_cells(
                        formula_sheet,
                        value_sheet,
                        layout,
                        settings.max_context_cells,
                    )
                except ValueError as error:
                    raise ModuleExecutionError(str(error)) from error

                render_cell_type_overlay(rendered_path, typed_path, layout, cells)
                visible_rows = [row for row, height in enumerate(layout.row_heights, start=1) if height > 0]
                visible_columns = [column for column, width in enumerate(layout.column_widths, start=1) if width > 0]
                if not visible_rows or not visible_columns or not cells:
                    raise ModuleExecutionError(f"시트 {sheet_name}에 표시된 값 셀 영역이 없습니다")
                sheet_bounds = CellBounds(
                    visible_rows[0],
                    visible_rows[-1],
                    visible_columns[0],
                    visible_columns[-1],
                )
                visibility = WorksheetVisibility.from_worksheet(value_sheet)
                try:
                    decisions = self._analyze_sheet(
                        settings,
                        sheet_name,
                        sheet_bounds,
                        typed_path,
                        layout,
                        visibility,
                        cells,
                    )
                except OpenAIResponsesVisionError as error:
                    raise ModuleExecutionError(str(error)) from error

                for table_index, decision in enumerate(decisions, start=1):
                    outputs.append(
                        self.structure_assembler._table_output(
                            table_index,
                            decision,
                            value_sheet,
                            layout,
                        )
                    )
        finally:
            if formula_workbook is not None:
                formula_workbook.close()
            if value_workbook is not None:
                value_workbook.close()

        return {
            "file_name": workbook_path.name,
            "workbook_hash": current_hash,
            "tables": outputs,
        }
