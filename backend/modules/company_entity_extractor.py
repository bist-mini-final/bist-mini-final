"""Extract a company entity from a selected workbook."""

from __future__ import annotations

import json
import logging
from pathlib import Path
from typing import Any, ClassVar, Dict, Optional

from pydantic import Field

from ..core.settings import PROCESSED_DATA_DIR
from ..llm.chat_completion import ChatCompletionClient
from ..spreadsheets.workbook_catalog import WorkbookCatalog, WorkbookCatalogError
from .base import (
    ExecutableModule,
    ModuleConfigDTO,
    ModuleDefinition,
    ModuleDTO,
    ModuleExecutionError,
)
from .processed_file_selector import WorkbookSelectionDTO


logger = logging.getLogger(__name__)

COMPANY_EXTRACTION_SYSTEM_PROMPT = """You identify the corporate entity represented in a financial spreadsheet.
Use the file name, worksheet names, and sampled top cells to infer the formal company name and ticker.
Preserve the standard English or Korean company name. Return only JSON matching the supplied schema."""


class CompanyEntityExtractorInputDTO(WorkbookSelectionDTO):
    """Workbook selection received directly from processed_file_selector."""


class CompanyEntityExtractorConfigDTO(ModuleConfigDTO):
    model: str = Field(
        default="gpt-5.6-luna",
        description="기업 엔티티 추출에 사용할 LLM 모델",
    )


class CompanyEntityExtractorExecutionDTO(
    CompanyEntityExtractorInputDTO,
    CompanyEntityExtractorConfigDTO,
):
    """Validated input and config used by the extractor."""


class CompanyEntityExtractorOutputDTO(ModuleDTO):
    company_name: str = Field(default="", description="공식 기업명")
    ticker: str = Field(default="", description="티커 심볼 (없으면 빈 문자열)")
    display_name: str = Field(default="", description="'기업명 (TICKER)' 형식 표시명")
    confidence: str = Field(default="low", description="추출 신뢰도 (high/medium/low)")
    source: str = Field(default="heuristic", description="추출 방법 (llm / heuristic)")


_COMPANY_ENTITY_DEFINITION = ModuleDefinition(
    type="company_entity_extractor",
    label="Company Entity Extractor",
    category="VLM Vision",
    description="선택된 Excel 워크북에서 기업명과 티커를 추출합니다.",
    inputs=["input"],
    outputs=["output"],
    config_fields=["model"],
    raw_output=True,
    version="2",
)


class CompanyEntityExtractorModule(ExecutableModule):
    """Extract company metadata with an LLM and deterministic fallback."""

    definition: ClassVar[ModuleDefinition] = _COMPANY_ENTITY_DEFINITION
    input_model = CompanyEntityExtractorInputDTO
    config_model = CompanyEntityExtractorConfigDTO
    execution_model = CompanyEntityExtractorExecutionDTO
    output_model = CompanyEntityExtractorOutputDTO

    def __init__(
        self,
        catalog: WorkbookCatalog | None = None,
        processed_dir: Path = PROCESSED_DATA_DIR,
        completion_client: Optional[ChatCompletionClient] = None,
    ) -> None:
        """Initialize the extractor with a workbook catalog.
        
        Parameters:
        	catalog (WorkbookCatalog | None): Catalog used to resolve workbooks. A catalog for `processed_dir` is created when omitted.
        	processed_dir (Path): Directory used to create the default workbook catalog.
        """
        self.catalog = catalog or WorkbookCatalog(processed_dir)
        self.completion_client = completion_client or ChatCompletionClient()

    @staticmethod
    def _heuristic(file_name: str) -> Dict[str, Any]:
        """
        Derive company metadata from a workbook filename.
        
        Parameters:
        	file_name (str): Workbook filename used to derive the company name.
        
        Returns:
        	Dict[str, Any]: Metadata containing the derived company name and display name, an empty ticker, low confidence, and a heuristic source.
        """
        name = Path(file_name).stem.replace("_", " ").replace("-", " ").strip()
        return {
            "company_name": name,
            "ticker": "",
            "display_name": name,
            "confidence": "low",
            "source": "heuristic",
        }

    @staticmethod
    def _sample_workbook(path: Path, sheet_names: list[str]) -> list[str]:
        """Extracts representative text from the first two requested workbook sheets.
        
        Parameters:
            path (Path): Path to the Excel workbook.
            sheet_names (list[str]): Names of the sheets to sample.
        
        Returns:
            list[str]: Sampled sheet contents, with populated cells from the first eight rows and ten columns.
        """
        import openpyxl

        sampled: list[str] = []
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet_name in sheet_names[:2]:
                if sheet_name not in workbook.sheetnames:
                    continue
                lines: list[str] = []
                for row in workbook[sheet_name].iter_rows(
                    max_row=8,
                    max_col=10,
                    values_only=True,
                ):
                    cells = [
                        str(cell).strip()
                        for cell in row
                        if cell is not None and str(cell).strip()
                    ]
                    if cells:
                        lines.append(" | ".join(cells[:10]))
                if lines:
                    sampled.append(f"[{sheet_name}]\n" + "\n".join(lines))
        finally:
            workbook.close()
        return sampled

    def execute(
        self,
        payload: CompanyEntityExtractorExecutionDTO,
    ) -> Dict[str, Any]:
        """
        Extract company metadata from the selected workbook.
        
        Parameters:
        	payload (CompanyEntityExtractorExecutionDTO): Workbook selection and model configuration used for extraction.
        
        Returns:
        	Dict[str, Any]: Extracted company name, ticker, display name, confidence, and source. Uses filename-based metadata when workbook sampling or LLM extraction is unavailable or fails.
        
        Raises:
        	ModuleExecutionError: If the workbook cannot be resolved from the catalog.
        """
        try:
            workbook_path = self.catalog.resolve(payload.file_name)
        except (OSError, ValueError, WorkbookCatalogError) as error:
            raise ModuleExecutionError(str(error)) from error

        fallback = self._heuristic(payload.file_name)
        if not getattr(self.completion_client, "api_key", None):
            return fallback

        try:
            sampled_lines = self._sample_workbook(
                workbook_path,
                payload.sheet_names,
            )
        except Exception as error:
            logger.warning("기업 엔티티 셀 샘플링 실패: %s", error)
            return fallback
        if not sampled_lines:
            return fallback

        sampled_text = "\n".join(sampled_lines)[:2500]
        context = (
            f"File Name: {payload.file_name}\n"
            f"Sheet Names: {payload.sheet_names}\n\n"
            f"Top Cells Content:\n{sampled_text}"
        )
        try:
            response = self.completion_client.complete_with_metadata(
                model=payload.model,
                messages=[
                    {"role": "system", "content": COMPANY_EXTRACTION_SYSTEM_PROMPT},
                    {"role": "user", "content": context},
                ],
                response_format={
                    "type": "json_schema",
                    "json_schema": {
                        "name": "company_metadata",
                        "strict": True,
                        "schema": {
                            "type": "object",
                            "properties": {
                                "company_name": {"type": "string"},
                                "ticker": {"type": "string"},
                                "display_name": {"type": "string"},
                                "confidence": {"type": "string"},
                            },
                            "required": [
                                "company_name",
                                "ticker",
                                "display_name",
                                "confidence",
                            ],
                            "additionalProperties": False,
                        },
                    },
                },
            )
            parsed = json.loads(response.content or "{}")
            company_name = str(parsed.get("company_name") or "").strip()
            ticker = str(parsed.get("ticker") or "").strip()
            display_name = str(parsed.get("display_name") or "").strip()
            display_name = display_name or (
                f"{company_name} ({ticker})" if ticker else company_name
            )
            if not display_name:
                return fallback
            return {
                "company_name": company_name or display_name,
                "ticker": ticker,
                "display_name": display_name,
                "confidence": str(parsed.get("confidence") or "high"),
                "source": "llm",
            }
        except Exception as error:
            logger.warning("기업 엔티티 LLM 추출 실패, 파일명 기반 결과 사용: %s", error)
            return fallback
