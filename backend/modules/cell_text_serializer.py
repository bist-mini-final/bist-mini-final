from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Literal, Optional, cast

import openpyxl
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from ..config import PROCESSED_DATA_DIR
from ..spreadsheets.cell_visibility import worksheet_visible
from ..spreadsheets.structured_cell_text import (
    SERIALIZATION_VERSION,
    UNKNOWN_FIELD,
    WorksheetValueReader,
    canonical_sheet_name,
    ensure_period_header,
    generate_header_combinations,
    serialize_structured_cell,
    sheet_code,
)
from ..spreadsheets.workbook_catalog import WorkbookCatalog, WorkbookCatalogError
from .base import (
    EmptyModuleConfigDTO,
    ExecutableModule,
    ModuleDefinition,
    ModuleDTO,
    ModuleExecutionError,
)
from .spreadsheet_structure import (
    ClassifiedRegionDTO,
    ClassifiedTableDTO,
    SpreadsheetStructureOutput,
)


class CellTextSerializerInputDTO(SpreadsheetStructureOutput):
    """Classified workbook data received from a structure detector."""


class CellTextDocumentDTO(ModuleDTO):
    cell_id: str = Field(description="시트 코드와 셀 좌표로 만든 검색 문서 ID")
    sheet_name: str = Field(description="정규화된 원본 시트 이름")
    cell_coord: str = Field(description="Excel 셀 좌표")
    row_header: List[str] = Field(description="상위 수준부터 수집한 행 헤더")
    column_header: List[str] = Field(description="상위 수준부터 수집한 열 헤더")
    cell_value: str = Field(description="Excel 데이터 셀의 표시 값")
    variant: Literal["header_only", "header_with_value"]
    text: str = Field(description="4필드 공통 포맷으로 직렬화한 검색 문서")


class CellTextSerializerOutput(ModuleDTO):
    file_name: str
    workbook_hash: str
    items: List[CellTextDocumentDTO]


class CellTextSerializerModule(ExecutableModule):
    definition = ModuleDefinition(
        type="cell_text_serializer",
        label="Structured Cell Text Serializer",
        category="Transform",
        description="분류된 Excel 셀을 Sheet·Row Header·Column Header·Cell Value 포맷으로 직렬화합니다.",
        inputs=["input"],
        outputs=["output"],
        config_fields=[],
        raw_output=True,
        version=SERIALIZATION_VERSION,
    )
    input_model = CellTextSerializerInputDTO
    config_model = EmptyModuleConfigDTO
    execution_model = CellTextSerializerInputDTO
    output_model = CellTextSerializerOutput

    def __init__(
        self,
        catalog: WorkbookCatalog | None = None,
        processed_dir: Path = PROCESSED_DATA_DIR,
    ) -> None:
        self.catalog = catalog or WorkbookCatalog(processed_dir)

    @staticmethod
    def _region(
        table: ClassifiedTableDTO,
        region_type: str,
    ) -> Optional[ClassifiedRegionDTO]:
        return next(
            (region for region in table.regions if region.type == region_type),
            None,
        )

    @staticmethod
    def _header_values(
        reader: WorksheetValueReader,
        row: int,
        column: int,
        title_region: Optional[ClassifiedRegionDTO],
        row_region: Optional[ClassifiedRegionDTO],
        column_region: Optional[ClassifiedRegionDTO],
    ) -> tuple[List[str], List[str]]:
        row_headers: List[str] = []
        if title_region is not None:
            for title_row in range(title_region.rows[0], title_region.rows[1] + 1):
                if reader.row_hidden(title_row):
                    continue
                for title_column in range(
                    title_region.columns[0],
                    title_region.columns[1] + 1,
                ):
                    if reader.column_hidden(title_column):
                        continue
                    value = reader.value(title_row, title_column)
                    if value is not None and value not in row_headers:
                        row_headers.append(value)
        if row_region is not None:
            for header_column in range(
                row_region.columns[0],
                row_region.columns[1] + 1,
            ):
                if reader.column_hidden(header_column):
                    continue
                value = reader.value(row, header_column)
                if value is not None and value not in row_headers:
                    row_headers.append(value)

        column_headers: List[str] = []
        if column_region is not None:
            for header_row in range(
                column_region.rows[0],
                column_region.rows[1] + 1,
            ):
                if reader.row_hidden(header_row):
                    continue
                value = reader.value(header_row, column)
                if value is not None:
                    column_headers.append(value)
        return row_headers, ensure_period_header(column_headers)

    def _table_documents(
        self,
        reader: WorksheetValueReader,
        table: ClassifiedTableDTO,
        seen_cell_ids: set[str],
    ) -> List[Dict[str, Any]]:
        data_region = self._region(table, "data")
        if data_region is None:
            return []
        row_region = self._region(table, "row_header")
        column_region = self._region(table, "column_header")
        title_region = self._region(table, "title")
        canonical_name = canonical_sheet_name(table.sheet_name)
        code = sheet_code(table.sheet_name)
        documents: List[Dict[str, Any]] = []

        for row in range(data_region.rows[0], data_region.rows[1] + 1):
            if reader.row_hidden(row):
                continue
            for column in range(
                data_region.columns[0],
                data_region.columns[1] + 1,
            ):
                if reader.column_hidden(column):
                    continue
                cell_value = reader.value(row, column)
                if cell_value is None:
                    continue
                row_headers, column_headers = self._header_values(
                    reader,
                    row,
                    column,
                    title_region,
                    row_region,
                    column_region,
                )
                if not row_headers:
                    continue

                cell_coord = f"{get_column_letter(column)}{row}"
                cell_id = f"{code} Cell {cell_coord}"
                if cell_id in seen_cell_ids:
                    raise ModuleExecutionError(
                        f"겹치는 테이블 영역에서 중복 셀이 생성되었습니다: {cell_id}"
                    )
                seen_cell_ids.add(cell_id)

                cell_seen_texts: set[str] = set()
                header_combos = generate_header_combinations(row_headers, column_headers)
                for r_combo, c_combo in header_combos:
                    common = {
                        "cell_id": cell_id,
                        "sheet_name": canonical_name,
                        "cell_coord": cell_coord,
                        "row_header": r_combo,
                        "column_header": c_combo,
                        "cell_value": cell_value,
                    }
                    for variant, val in [("header_only", UNKNOWN_FIELD), ("header_with_value", cell_value)]:
                        text = serialize_structured_cell(
                            canonical_name,
                            r_combo,
                            c_combo,
                            val,
                        )
                        if text not in cell_seen_texts:
                            cell_seen_texts.add(text)
                            documents.append(
                                {
                                    **common,
                                    "variant": variant,
                                    "text": text,
                                }
                            )
        return documents

    def execute(self, payload: BaseModel) -> Dict[str, Any]:
        input_data = cast(CellTextSerializerInputDTO, payload)
        try:
            workbook_path = self.catalog.resolve(input_data.file_name)
            current_hash = self.catalog.sha256(workbook_path)
        except (OSError, ValueError, WorkbookCatalogError) as error:
            raise ModuleExecutionError(str(error)) from error
        if current_hash != input_data.workbook_hash:
            raise ModuleExecutionError(
                "영역 분류 이후 Excel 파일이 변경되었습니다. 앞 모듈부터 다시 실행하세요"
            )

        workbook = None
        try:
            workbook = openpyxl.load_workbook(
                workbook_path,
                read_only=False,
                data_only=True,
                keep_vba=workbook_path.suffix.lower() == ".xlsm",
            )
            items: List[Dict[str, Any]] = []
            seen_cell_ids: set[str] = set()
            readers: Dict[str, WorksheetValueReader] = {}
            for table in input_data.tables:
                if table.sheet_name not in workbook.sheetnames:
                    raise ModuleExecutionError(
                        f"Excel 시트를 찾을 수 없습니다: {table.sheet_name}"
                    )
                reader = readers.get(table.sheet_name)
                if reader is None:
                    worksheet = workbook[table.sheet_name]
                    if not worksheet_visible(worksheet):
                        raise ModuleExecutionError(
                            f"숨겨진 Excel 시트는 직렬화할 수 없습니다: {table.sheet_name}"
                        )
                    reader = WorksheetValueReader(worksheet)
                    readers[table.sheet_name] = reader
                items.extend(
                    self._table_documents(
                        reader,
                        table,
                        seen_cell_ids,
                    )
                )
        finally:
            if workbook is not None:
                workbook.close()

        return {
            "file_name": workbook_path.name,
            "workbook_hash": current_hash,
            "items": items,
        }
