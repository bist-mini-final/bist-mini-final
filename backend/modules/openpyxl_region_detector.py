from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Literal, cast

import openpyxl
from pydantic import BaseModel, Field

from ..config import PROCESSED_DATA_DIR
from ..spreadsheets.cell_visibility import WorksheetVisibility, worksheet_visible
from ..spreadsheets.table_geometry import CellBounds, cell_bounds_bbox, compute_sheet_layout
from ..spreadsheets.grid_structure import build_column_header_tree
from ..spreadsheets.workbook_catalog import WorkbookCatalog, WorkbookCatalogError
from .base import ExecutableModule, ModuleConfigDTO, ModuleDefinition, ModuleExecutionError, ModuleInputDTO
from .docling_table_detector import DoclingTableRegionDTO
from .spreadsheet_structure import (
    ClassifiedRegionDTO,
    ClassifiedTableDTO,
    SpreadsheetStructureOutput,
)


class OpenpyxlRegionDetectorInputDTO(ModuleInputDTO):
    file_name: str
    workbook_hash: str
    tables: List[DoclingTableRegionDTO]


class OpenpyxlRegionDetectorConfigDTO(ModuleConfigDTO):
    header_scan_rows: int = Field(
        default=10,
        ge=1,
        le=50,
        description="각 테이블 상단에서 헤더 스타일을 검사할 최대 행 수",
    )
    bold_ratio_threshold: float = Field(
        default=0.3,
        ge=0,
        le=1,
        description="column_header로 판정할 최소 bold 셀 비율",
    )
    fill_ratio_threshold: float = Field(
        default=0.4,
        ge=0,
        le=1,
        description="column_header로 판정할 최소 배경색 셀 비율",
    )


class OpenpyxlRegionDetectorExecutionDTO(
    OpenpyxlRegionDetectorInputDTO,
    OpenpyxlRegionDetectorConfigDTO,
):
    """Internal union of detected tables and classification policy."""


class OpenpyxlRegionDetectorOutput(SpreadsheetStructureOutput):
    pass


class OpenpyxlRegionDetectorModule(ExecutableModule):
    definition = ModuleDefinition(
        type="openpyxl_region_detector",
        label="OpenPyXL Table Region Classifier",
        category="Logic",
        description="Docling 테이블 경계 안에서 헤더와 데이터 영역을 셀 서식으로 분류합니다.",
        inputs=["input"],
        outputs=["output"],
        config_fields=[
            "header_scan_rows",
            "bold_ratio_threshold",
            "fill_ratio_threshold",
        ],
        raw_output=True,
        version="3",
    )
    input_model = OpenpyxlRegionDetectorInputDTO
    config_model = OpenpyxlRegionDetectorConfigDTO
    execution_model = OpenpyxlRegionDetectorExecutionDTO
    output_model = OpenpyxlRegionDetectorOutput

    def __init__(
        self,
        catalog: WorkbookCatalog | None = None,
        processed_dir: Path = PROCESSED_DATA_DIR,
    ) -> None:
        self.catalog = catalog or WorkbookCatalog(processed_dir)

    @staticmethod
    def _region(
        region_id: str,
        region_type: Literal["title", "column_header", "row_header", "data"],
        bounds: CellBounds,
        layout,
        parent_ids: List[str],
    ) -> Dict[str, Any]:
        return {
            "region_id": region_id,
            "type": region_type,
            "excel_range": bounds.excel_range,
            "bbox_px": cell_bounds_bbox(bounds, layout),
            "rows": (bounds.min_row, bounds.max_row),
            "columns": (bounds.min_column, bounds.max_column),
            "parent_ids": parent_ids,
        }

    def _analyze_table(
        self,
        worksheet,
        table_index: int,
        bounds: CellBounds,
        layout,
        settings: OpenpyxlRegionDetectorExecutionDTO,
    ) -> Dict[str, Any]:
        visibility = WorksheetVisibility.from_worksheet(worksheet)
        visible_rows = [
            row
            for row in range(bounds.min_row, bounds.max_row + 1)
            if not visibility.row_hidden(row)
        ]
        visible_columns = [
            column
            for column in range(bounds.min_column, bounds.max_column + 1)
            if not visibility.column_hidden(column)
        ]
        if not visible_rows or not visible_columns:
            raise ModuleExecutionError(
                f"표시된 셀이 없는 테이블 영역입니다: {bounds.excel_range}"
            )
        visible_bounds = CellBounds(
            visible_rows[0],
            visible_rows[-1],
            visible_columns[0],
            visible_columns[-1],
        )
        total_columns = len(visible_columns)
        scan_rows = visible_rows[: settings.header_scan_rows]

        title_end_row = visible_rows[0] - 1
        if total_columns > 1:
            for row in scan_rows:
                non_empty = sum(
                    1
                    for column in visible_columns
                    if worksheet.cell(row=row, column=column).value not in (None, "")
                )
                if non_empty == 1:
                    title_end_row = row
                else:
                    break

        remaining_scan_rows = [row for row in scan_rows if row > title_end_row]
        header_start_row = (
            remaining_scan_rows[0]
            if remaining_scan_rows
            else visible_rows[0]
        )
        header_end_row = header_start_row - 1

        for row in remaining_scan_rows:
            bold_count = 0
            fill_count = 0
            for column in visible_columns:
                cell = worksheet.cell(row=row, column=column)
                if cell.font and cell.font.bold:
                    bold_count += 1
                fill = cell.fill
                if fill and fill.fill_type and fill.fill_type != "none":
                    fill_count += 1
            bold_ratio = bold_count / max(1, total_columns)
            fill_ratio = fill_count / max(1, total_columns)
            if (
                bold_ratio > settings.bold_ratio_threshold
                or fill_ratio > settings.fill_ratio_threshold
            ):
                header_end_row = row
            elif header_end_row >= bounds.min_row:
                break

        content_data_start = None
        for row in remaining_scan_rows:
            values = [
                worksheet.cell(row=row, column=column).value
                for column in visible_columns
            ]
            non_empty = [value for value in values if value not in (None, "")]
            numeric = [
                value
                for value in non_empty
                if isinstance(value, (int, float)) and not isinstance(value, bool)
            ]
            period_numbers = numeric and all(
                isinstance(value, int) and 1900 <= value <= 2100 for value in numeric
            )
            contains_date = any(isinstance(value, (date, datetime)) for value in non_empty)
            if (
                numeric
                and not period_numbers
                and not contains_date
                and len(numeric) >= max(1, (len(non_empty) - 1 + 1) // 2)
            ):
                content_data_start = row
                break

        if content_data_start is not None:
            data_start_row = content_data_start
            preceding_rows = [
                row for row in remaining_scan_rows if row < data_start_row
            ]
            header_end_row = (
                preceding_rows[-1]
                if preceding_rows
                else header_start_row - 1
            )
        else:
            if header_end_row < header_start_row:
                header_end_row = header_start_row
            data_rows = [row for row in visible_rows if row > header_end_row]
            data_start_row = data_rows[0] if data_rows else bounds.max_row + 1

        prefix = f"table_{table_index}"
        column_header_id = f"{prefix}_column_header"
        row_header_id = f"{prefix}_row_header"
        regions: List[Dict[str, Any]] = []
        parent_ids: List[str] = []
        if title_end_row >= visible_rows[0]:
            title_id = f"{prefix}_title"
            regions.append(
                self._region(
                    title_id,
                    "title",
                    CellBounds(
                        visible_rows[0],
                        title_end_row,
                        visible_columns[0],
                        visible_columns[-1],
                    ),
                    layout,
                    [],
                )
            )
            parent_ids.append(title_id)
        if header_start_row <= header_end_row:
            regions.append(
                self._region(
                    column_header_id,
                    "column_header",
                    CellBounds(
                        header_start_row,
                        header_end_row,
                        visible_columns[0],
                        visible_columns[-1],
                    ),
                    layout,
                    list(parent_ids),
                )
            )
            parent_ids.append(column_header_id)

        index_column_count = 0
        sample_rows = [row for row in visible_rows if row >= data_start_row][:20]
        for column_index, column in enumerate(visible_columns[:-1]):
            column_values = []
            for row in sample_rows:
                value = worksheet.cell(row=row, column=column).value
                if value not in (None, ""):
                    column_values.append(value)
            if not column_values:
                break
            text_ratio = sum(
                1
                for value in column_values
                if not isinstance(value, (int, float, date, datetime))
            ) / len(column_values)
            rows_with_numeric_right = sum(
                1
                for row in sample_rows
                if any(
                    isinstance(worksheet.cell(row=row, column=right).value, (int, float))
                    and not isinstance(worksheet.cell(row=row, column=right).value, bool)
                    for right in visible_columns[column_index + 1 :]
                )
            )
            if text_ratio >= 0.5 and rows_with_numeric_right > 0:
                index_column_count += 1
            else:
                break

        if index_column_count == 0 and len(visible_columns) > 1:
            index_column_count = 1
        has_row_header = index_column_count > 0
        if sample_rows and has_row_header:
            regions.append(
                self._region(
                    row_header_id,
                    "row_header",
                    CellBounds(
                        data_start_row,
                        visible_rows[-1],
                        visible_columns[0],
                        visible_columns[index_column_count - 1],
                    ),
                    layout,
                    list(parent_ids),
                )
            )

        data_start_column = (
            visible_columns[index_column_count]
            if index_column_count < len(visible_columns)
            else visible_columns[-1] + 1
        )
        if sample_rows and data_start_column <= visible_columns[-1]:
            parents = list(parent_ids)
            if has_row_header:
                parents.append(row_header_id)
            regions.append(
                self._region(
                    f"{prefix}_data",
                    "data",
                    CellBounds(
                        data_start_row,
                        visible_rows[-1],
                        data_start_column,
                        visible_columns[-1],
                    ),
                    layout,
                    parents,
                )
            )

        return {
            "table_index": table_index,
            "excel_range": visible_bounds.excel_range,
            "regions": regions,
            "header_tree": build_column_header_tree(
                worksheet,
                visible_bounds,
                header_start_row,
                header_end_row,
                data_start_column,
            ),
        }

    def execute(self, payload: BaseModel) -> Dict[str, Any]:
        input_data = cast(OpenpyxlRegionDetectorExecutionDTO, payload)
        try:
            workbook_path = self.catalog.resolve(input_data.file_name)
            current_hash = self.catalog.sha256(workbook_path)
        except (OSError, ValueError, WorkbookCatalogError) as error:
            raise ModuleExecutionError(str(error)) from error
        if current_hash != input_data.workbook_hash:
            raise ModuleExecutionError(
                "Docling 탐지 이후 Excel 파일이 변경되었습니다. 앞 모듈부터 다시 실행하세요"
            )

        workbook = None
        try:
            workbook = openpyxl.load_workbook(
                workbook_path,
                read_only=False,
                data_only=True,
                keep_vba=workbook_path.suffix.lower() == ".xlsm",
            )
            output_tables: List[Dict[str, Any]] = []
            sheet_layouts: Dict[str, Any] = {}
            for table in input_data.tables:
                if table.sheet_name not in workbook.sheetnames:
                    raise ModuleExecutionError(
                        f"Excel 시트를 찾을 수 없습니다: {table.sheet_name}"
                    )
                worksheet = workbook[table.sheet_name]
                if not worksheet_visible(worksheet):
                    raise ModuleExecutionError(
                        f"숨겨진 Excel 시트는 분석할 수 없습니다: {table.sheet_name}"
                    )
                layout = sheet_layouts.get(table.sheet_name)
                if layout is None:
                    layout = compute_sheet_layout(
                        worksheet,
                        worksheet.max_row or 1,
                        worksheet.max_column or 1,
                    )
                    sheet_layouts[table.sheet_name] = layout

                cell_bounds = table.cell_bounds
                bounds = CellBounds(
                    cell_bounds.min_row,
                    cell_bounds.max_row,
                    cell_bounds.min_column,
                    cell_bounds.max_column,
                )
                classified = self._analyze_table(
                    worksheet,
                    table.table_index,
                    bounds,
                    layout,
                    input_data,
                )
                output_tables.append({"sheet_name": table.sheet_name, **classified})
        finally:
            if workbook is not None:
                workbook.close()

        return {
            "file_name": workbook_path.name,
            "workbook_hash": current_hash,
            "tables": output_tables,
        }
