from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Literal, Optional, Sequence, cast

import openpyxl
from openpyxl.cell.cell import MergedCell
from pydantic import BaseModel, Field

from ..config import PROCESSED_DATA_DIR
from ..spreadsheets.cell_visibility import WorksheetVisibility, worksheet_visible
from ..spreadsheets.structured_cell_text import (
    UNKNOWN_FIELD,
    canonical_sheet_name,
    format_cell_value,
    serialize_structured_cell,
    sheet_code,
)
from ..spreadsheets.workbook_catalog import WorkbookCatalog, WorkbookCatalogError
from .base import ExecutableModule, ModuleDefinition, ModuleExecutionError
from .cell_text_serializer import CellTextSerializerOutput
from .processed_file_selector import WorkbookSelectionDTO


EXHAUSTIVE_SERIALIZATION_VERSION = "exhaustive-cell-v1-visible-only"


class ExhaustiveCellTextSerializerInput(WorkbookSelectionDTO):
    variant_mode: Literal["header_only", "header_with_value", "both"] = Field(
        default="header_only",
        description=(
            "각 헤더 조합에서 생성할 검색 문서 형태. header_only는 Cell Value를 ?로 "
            "두고, header_with_value는 실제 셀 값을 포함합니다"
        ),
    )
    deduplicate_header_values: bool = Field(
        default=True,
        description=(
            "같은 방향에서 값이 동일한 여러 셀은 같은 직렬화 문장을 만들므로 한 후보로 "
            "합칩니다. 끄면 셀 좌표별 조합을 모두 생성합니다"
        ),
    )
    max_documents: int = Field(
        default=2_000_000,
        ge=1,
        le=10_000_000,
        description=(
            "조합 폭증으로 프로세스가 종료되는 것을 막는 안전 한도. 초과 시 일부만 "
            "저장하지 않고 실행 전체를 실패시킵니다"
        ),
    )


@dataclass(frozen=True)
class PopulatedCell:
    row: int
    column: int
    coord: str
    value: str


@dataclass(frozen=True)
class SheetCells:
    sheet_name: str
    cells: Sequence[PopulatedCell]
    rows: Dict[int, Sequence[PopulatedCell]]
    columns: Dict[int, Sequence[PopulatedCell]]


class ExhaustiveCellTextSerializerModule(ExecutableModule):
    """Serialize every distinct left × above header combination without a model."""

    definition = ModuleDefinition(
        type="exhaustive_cell_text_serializer",
        label="Exhaustive Cell Header Serializer",
        category="Transform",
        description=(
            "모든 표시 값 셀을 데이터 셀로 보고 왼쪽 행 후보와 위쪽 열 후보의 "
            "전체 조합을 결정적으로 직렬화합니다."
        ),
        inputs=["input"],
        outputs=["output"],
        config_fields=[
            "variant_mode",
            "deduplicate_header_values",
            "max_documents",
        ],
        raw_output=True,
        version=EXHAUSTIVE_SERIALIZATION_VERSION,
    )
    input_model = ExhaustiveCellTextSerializerInput
    output_model = CellTextSerializerOutput

    def __init__(
        self,
        catalog: WorkbookCatalog | None = None,
        processed_dir: Path = PROCESSED_DATA_DIR,
    ) -> None:
        self.catalog = catalog or WorkbookCatalog(processed_dir)

    @staticmethod
    def _ordered_unique(cells: Iterable[PopulatedCell]) -> List[PopulatedCell]:
        seen: set[str] = set()
        unique: List[PopulatedCell] = []
        for cell in cells:
            if cell.value in seen:
                continue
            seen.add(cell.value)
            unique.append(cell)
        return unique

    @classmethod
    def _candidates(
        cls,
        cells: Iterable[PopulatedCell],
        deduplicate: bool,
    ) -> Sequence[Optional[PopulatedCell]]:
        candidates = list(cells)
        if deduplicate:
            candidates = cls._ordered_unique(candidates)
        return candidates or [None]

    @staticmethod
    def _read_sheet(worksheet) -> SheetCells:
        visibility = WorksheetVisibility.from_worksheet(worksheet)
        cells: List[PopulatedCell] = []
        rows: Dict[int, List[PopulatedCell]] = {}
        columns: Dict[int, List[PopulatedCell]] = {}

        for worksheet_row in worksheet.iter_rows():
            for cell in worksheet_row:
                if isinstance(cell, MergedCell):
                    continue
                if not visibility.cell_visible(cell.row, cell.column):
                    continue
                value = format_cell_value(cell.value)
                if value is None:
                    continue
                populated = PopulatedCell(
                    row=cell.row,
                    column=cell.column,
                    coord=cell.coordinate,
                    value=value,
                )
                cells.append(populated)
                rows.setdefault(cell.row, []).append(populated)
                columns.setdefault(cell.column, []).append(populated)

        return SheetCells(
            sheet_name=worksheet.title,
            cells=cells,
            rows=rows,
            columns=columns,
        )

    @classmethod
    def _header_candidates(
        cls,
        sheet: SheetCells,
        target: PopulatedCell,
        deduplicate: bool,
    ) -> tuple[
        Sequence[Optional[PopulatedCell]],
        Sequence[Optional[PopulatedCell]],
    ]:
        left = cls._candidates(
            (
                cell
                for cell in sheet.rows.get(target.row, ())
                if cell.column < target.column
            ),
            deduplicate,
        )
        above = cls._candidates(
            (
                cell
                for cell in sheet.columns.get(target.column, ())
                if cell.row < target.row
            ),
            deduplicate,
        )
        return left, above

    @staticmethod
    def _variant_values(
        mode: Literal["header_only", "header_with_value", "both"],
        cell_value: str,
    ) -> Sequence[tuple[Literal["header_only", "header_with_value"], str]]:
        if mode == "header_only":
            return [("header_only", UNKNOWN_FIELD)]
        if mode == "header_with_value":
            return [("header_with_value", cell_value)]
        return [
            ("header_only", UNKNOWN_FIELD),
            ("header_with_value", cell_value),
        ]

    @classmethod
    def _document_count(
        cls,
        sheets: Sequence[SheetCells],
        input_data: ExhaustiveCellTextSerializerInput,
    ) -> int:
        variants = 2 if input_data.variant_mode == "both" else 1
        return sum(
            len(left) * len(above) * variants
            for sheet in sheets
            for target in sheet.cells
            for left, above in [
                cls._header_candidates(
                    sheet,
                    target,
                    input_data.deduplicate_header_values,
                )
            ]
        )

    @classmethod
    def _documents(
        cls,
        sheets: Sequence[SheetCells],
        input_data: ExhaustiveCellTextSerializerInput,
    ) -> List[Dict[str, Any]]:
        documents: List[Dict[str, Any]] = []
        for sheet in sheets:
            canonical_name = canonical_sheet_name(sheet.sheet_name)
            code = sheet_code(sheet.sheet_name)
            for target in sheet.cells:
                left_candidates, above_candidates = cls._header_candidates(
                    sheet,
                    target,
                    input_data.deduplicate_header_values,
                )
                for left in left_candidates:
                    row_headers = [left.value] if left is not None else []
                    for above in above_candidates:
                        column_headers = [above.value] if above is not None else []
                        for variant, serialized_value in cls._variant_values(
                            input_data.variant_mode,
                            target.value,
                        ):
                            documents.append(
                                {
                                    "cell_id": f"{code} Cell {target.coord}",
                                    "sheet_name": canonical_name,
                                    "cell_coord": target.coord,
                                    "row_header": row_headers,
                                    "column_header": column_headers,
                                    "cell_value": target.value,
                                    "variant": variant,
                                    "text": serialize_structured_cell(
                                        canonical_name,
                                        row_headers,
                                        column_headers,
                                        serialized_value,
                                    ),
                                }
                            )
        return documents

    def execute(self, payload: BaseModel) -> Dict[str, Any]:
        input_data = cast(ExhaustiveCellTextSerializerInput, payload)
        try:
            workbook_path = self.catalog.resolve(input_data.file_name)
            current_hash = self.catalog.sha256(workbook_path)
        except (OSError, ValueError, WorkbookCatalogError) as error:
            raise ModuleExecutionError(str(error)) from error
        if current_hash != input_data.workbook_hash:
            raise ModuleExecutionError(
                "파일 선택 이후 Excel 파일이 변경되었습니다. 앞 모듈부터 다시 실행하세요"
            )

        workbook = None
        try:
            workbook = openpyxl.load_workbook(
                workbook_path,
                read_only=False,
                data_only=True,
                keep_vba=workbook_path.suffix.lower() == ".xlsm",
            )
            sheets: List[SheetCells] = []
            for sheet_name in input_data.sheet_names:
                if sheet_name not in workbook.sheetnames:
                    raise ModuleExecutionError(
                        f"Excel 시트를 찾을 수 없습니다: {sheet_name}"
                    )
                worksheet = workbook[sheet_name]
                if not worksheet_visible(worksheet):
                    raise ModuleExecutionError(
                        f"숨겨진 Excel 시트는 직렬화할 수 없습니다: {sheet_name}"
                    )
                sheets.append(self._read_sheet(worksheet))

            document_count = self._document_count(sheets, input_data)
            if document_count == 0:
                raise ModuleExecutionError("직렬화할 표시 값 셀이 없습니다")
            if document_count > input_data.max_documents:
                raise ModuleExecutionError(
                    "전체 헤더 조합 문서가 안전 한도를 초과합니다: "
                    f"{document_count:,}개 > {input_data.max_documents:,}개. "
                    "중복 헤더 값 병합을 켜거나 max_documents를 명시적으로 늘리세요"
                )
            items = self._documents(sheets, input_data)
        finally:
            if workbook is not None:
                workbook.close()

        return {
            "file_name": workbook_path.name,
            "workbook_hash": current_hash,
            "items": items,
        }
