"""Persist workbook sheet metadata after the pgvector index exists."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any, ClassVar, Dict, List, Union

from pydantic import Field

from ..core.settings import PROCESSED_DATA_DIR
from ..spreadsheets.workbook_catalog import WorkbookCatalog, WorkbookCatalogError
from .base import (
    EmptyModuleConfigDTO,
    ExecutableModule,
    ModuleDefinition,
    ModuleDTO,
    ModuleExecutionError,
)
from .processed_file_selector import WorkbookSelectionDTO
from .spreadsheet_structure import SpreadsheetStructureOutput
from .vector_index_writer import VectorIndexDTO


logger = logging.getLogger(__name__)
StructureSourceDTO = Union[SpreadsheetStructureOutput, WorkbookSelectionDTO]


class SheetMetadataPersistenceInputDTO(ModuleDTO):
    structure_input: StructureSourceDTO = Field(
        description="구조 감지 결과 또는 전수 직렬화용 워크북 선택 결과",
    )
    index_input: VectorIndexDTO = Field(
        description="source_files 저장이 완료된 pgvector 인덱스 결과",
    )


class SheetMetadataPersistenceOutputDTO(ModuleDTO):
    sheets_saved: int = Field(description="DB에 저장된 시트 수")
    sheet_details: List[Dict[str, Any]] = Field(default_factory=list)


_SHEET_META_DEFINITION = ModuleDefinition(
    type="sheet_metadata_persistence",
    label="Sheet Metadata Persistence",
    category="Storage / DB",
    description="pgvector 적재 후 워크북 시트 크기와 감지 테이블을 DB에 저장합니다.",
    inputs=["structure_input", "index_input"],
    outputs=["output"],
    config_fields=[],
    raw_output=True,
    cacheable=False,
    version="2",
)


class SheetMetadataPersistenceModule(ExecutableModule):
    """Persist sheet dimensions and detected table metadata."""

    definition: ClassVar[ModuleDefinition] = _SHEET_META_DEFINITION
    input_model = SheetMetadataPersistenceInputDTO
    config_model = EmptyModuleConfigDTO
    execution_model = SheetMetadataPersistenceInputDTO
    output_model = SheetMetadataPersistenceOutputDTO

    def __init__(
        self,
        db_manager: Any = None,
        catalog: WorkbookCatalog | None = None,
        processed_dir: Path = PROCESSED_DATA_DIR,
    ) -> None:
        """Initialize the module with an optional database manager and workbook catalog.
        
        Parameters:
            db_manager (Any): Database manager used for persistence.
            catalog (WorkbookCatalog | None): Workbook catalog to use. A catalog for
                `processed_dir` is created when omitted.
            processed_dir (Path): Directory containing processed workbook data.
        """
        self._db_manager = db_manager
        self.catalog = catalog or WorkbookCatalog(processed_dir)

    def execute(self, payload: SheetMetadataPersistenceInputDTO) -> Dict[str, Any]:
        """
        Persist workbook sheet metadata and detected tables for the selected visible sheets.
        
        Parameters:
            payload (SheetMetadataPersistenceInputDTO): Workbook structure and index data used to identify and describe the sheets.
        
        Returns:
            Dict[str, Any]: The number of saved sheets and per-sheet summaries containing dimensions and detected-table counts.
        
        Raises:
            ModuleExecutionError: If workbook inputs reference different workbooks, the database is unavailable, workbook metadata cannot be loaded, or persistence fails.
        """
        from ..storage.db_manager import DatabaseManager

        structure = payload.structure_input
        index = payload.index_input
        if structure.workbook_hash != index.workbook_hash:
            raise ModuleExecutionError("구조 분석과 인덱스의 workbook_hash가 다릅니다")

        database = self._db_manager or DatabaseManager()
        if not database.is_connected():
            raise ModuleExecutionError("시트 메타데이터를 저장할 DB에 연결할 수 없습니다")

        try:
            workbook_path = self.catalog.resolve(structure.file_name)
        except (OSError, ValueError, WorkbookCatalogError) as error:
            raise ModuleExecutionError(str(error)) from error

        detected_tables = (
            []
            if isinstance(structure, WorkbookSelectionDTO)
            else structure.tables
        )

        if structure.sheet_names:
            visible_sheets = list(structure.sheet_names)
        elif detected_tables:
            catalog_sheets = self.catalog.sheet_names(workbook_path)
            table_sheets = {t.sheet_name for t in detected_tables}
            visible_sheets = [s for s in catalog_sheets if s in table_sheets]
            for s in table_sheets:
                if s not in visible_sheets:
                    visible_sheets.append(s)
        else:
            raise ModuleExecutionError(
                "저장할 시트 목록(sheet_names) 또는 감지된 테이블(tables)이 지정되지 않았습니다"
            )
        sheet_dimensions: Dict[str, tuple[int, int]] = {}
        try:
            import openpyxl

            workbook = openpyxl.load_workbook(
                workbook_path,
                read_only=True,
                data_only=True,
            )
            try:
                for sheet_name in workbook.sheetnames:
                    worksheet = workbook[sheet_name]
                    if worksheet.max_row is None or worksheet.max_column is None:
                        worksheet.calculate_dimension(force=True)
                    sheet_dimensions[sheet_name] = (
                        worksheet.max_row or 0,
                        worksheet.max_column or 0,
                    )
            finally:
                workbook.close()
        except Exception as error:
            raise ModuleExecutionError(f"시트 크기 측정 실패: {error}") from error

        sheets_data: List[Dict[str, Any]] = []
        for sheet_index, sheet_name in enumerate(visible_sheets):
            rows, columns = sheet_dimensions.get(sheet_name, (0, 0))
            sheet_tables = [
                table.model_dump(mode="json")
                for table in detected_tables
                if table.sheet_name == sheet_name
            ]
            sheets_data.append(
                {
                    "sheet_name": sheet_name,
                    "sheet_index": sheet_index,
                    "is_visible": True,
                    "row_count": rows,
                    "column_count": columns,
                    "detected_tables": sheet_tables,
                }
            )

        try:
            database.save_sheets(
                file_id=structure.workbook_hash,
                sheets_info=sheets_data,
            )
        except Exception as error:
            raise ModuleExecutionError(f"시트 메타데이터 저장 실패: {error}") from error

        return {
            "sheets_saved": len(sheets_data),
            "sheet_details": [
                {
                    "sheet_name": sheet["sheet_name"],
                    "row_count": sheet["row_count"],
                    "column_count": sheet["column_count"],
                    "table_count": len(sheet["detected_tables"]),
                }
                for sheet in sheets_data
            ],
        }
