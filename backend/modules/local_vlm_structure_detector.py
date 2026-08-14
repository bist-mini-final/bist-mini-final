from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Dict, List, Optional, Protocol, cast

import openpyxl
from openpyxl.utils.cell import range_boundaries
from pydantic import BaseModel, Field

from ..config import PROCESSED_DATA_DIR, SPREADSHEET_ARTIFACT_DIR
from ..ollama_vision import OllamaVisionClient, OllamaVisionError
from ..spreadsheets.cell_semantics import collect_non_empty_cells, compact_sheet_context
from ..spreadsheets.cell_type_overlay import render_cell_type_overlay
from ..spreadsheets.cell_visibility import WorksheetVisibility, worksheet_visible
from ..spreadsheets.grid_structure import build_column_header_tree
from ..spreadsheets.prompt_guidance import TEXT_CELL_ROLE_GUIDANCE
from ..spreadsheets.sheet_renderer import ExcelSheetRenderer
from ..spreadsheets.table_geometry import CellBounds, cell_bounds_bbox
from ..spreadsheets.workbook_catalog import WorkbookCatalog, WorkbookCatalogError
from .base import ExecutableModule, ModuleConfigDTO, ModuleDefinition, ModuleDTO, ModuleExecutionError
from .docling_table_detector import _safe_name
from .processed_file_selector import WorkbookSelectionDTO
from .spreadsheet_structure import SpreadsheetStructureOutput


LOCAL_VLM_SYSTEM_PROMPT = f"""You analyze structured Excel worksheets using both an image and exact cell facts.
Hidden sheets, rows, and columns are intentionally absent. Never infer or restore them.
The image preserves the worksheet layout. Every populated cell is tinted by value type and labeled with its Excel coordinate. Colors: text amber, number green, date blue, boolean purple, error red, unresolved formula gray. A magenta inner border means the source cell is a formula.

{TEXT_CELL_ROLE_GUIDANCE}

Find every independent rectangular table. For each table, return exact Excel ranges for:
- excel_range: the whole table only, excluding unrelated notes and disclaimers.
- title_range: optional title rows describing the whole table.
- column_header_range: optional header rows above the data values.
- row_header_range: optional label columns to the left of data values.
- data_range: the value matrix only.

The column_header_range must include every header level and span every data column. The row_header_range must span every data row; it may contain only the merged-cell anchor column when the supplied merged_range shows that the label visually spans more columns. Include all period rows (for example Actuals/LTM and the exact dates), not only the nearest header row.

Use the supplied coordinates, types, values, formulas, and merged ranges as factual authority. Use the image for spatial grouping and visual style. A data cell's parent headers are the row-header cells on its row plus column-header cells above its column, including merged ancestors. Do not invent coordinates or values. Detect multiple tables when blank separation or distinct headers indicate separate structures. Return only JSON conforming to the schema."""


LOCAL_VLM_USER_TEMPLATE = """Analyze sheet {sheet_name}.
The compact tuple format is [excel_coord, value_type, value, optional_flags].
Here is the exact coordinate context:
{sheet_context}"""


class LocalVisionClient(Protocol):
    def complete_structured(
        self,
        model: str,
        system_prompt: str,
        user_prompt: str,
        image_path: Path,
        json_schema: Dict[str, Any],
        context_window: int,
        timeout_seconds: int,
    ) -> str: ...


class LocalVlmStructureDetectorInputDTO(WorkbookSelectionDTO):
    """Workbook identity and visible sheet selection."""


class LocalVlmStructureDetectorConfigDTO(ModuleConfigDTO):
    model: str = Field(
        default="qwen3-vl:4b-instruct",
        min_length=1,
        description="Ollama에 설치된 로컬 멀티모달 모델 ID",
    )
    max_rows: int = Field(default=400, ge=1, le=2000, description="시트에서 분석할 최대 행 수")
    max_columns: int = Field(default=60, ge=1, le=200, description="시트에서 분석할 최대 열 수")
    max_context_cells: int = Field(
        default=15000,
        ge=100,
        le=50000,
        description="한 시트 요청에 포함할 최대 값 셀 수(초과 시 생략하지 않고 실패)",
    )
    context_window: int = Field(
        default=65536,
        ge=8192,
        le=262144,
        description="Ollama 추론 컨텍스트 길이",
    )
    timeout_seconds: int = Field(
        default=900,
        ge=30,
        le=3600,
        description="시트별 로컬 VLM 요청 제한 시간(초)",
    )
    validation_retries: int = Field(
        default=1,
        ge=0,
        le=2,
        description="좌표 규칙을 위반한 로컬 VLM 응답의 교정 재시도 횟수",
    )
    system_prompt: str = Field(default=LOCAL_VLM_SYSTEM_PROMPT, min_length=1, description="시트 구조 식별 시스템 프롬프트")
    user_prompt_template: str = Field(
        default=LOCAL_VLM_USER_TEMPLATE,
        min_length=1,
        description="{sheet_name}, {sheet_context} 변수를 지원하는 사용자 프롬프트",
    )


class LocalVlmStructureDetectorExecutionDTO(
    LocalVlmStructureDetectorInputDTO,
    LocalVlmStructureDetectorConfigDTO,
):
    """Internal union of workbook data and VLM settings."""


class LocalVlmTableDecisionDTO(ModuleDTO):
    excel_range: str = Field(description="전체 테이블 Excel 범위")
    title_range: Optional[str] = Field(default=None, description="테이블 제목 범위")
    column_header_range: Optional[str] = Field(default=None, description="열 헤더 범위")
    row_header_range: Optional[str] = Field(default=None, description="행 헤더 범위")
    data_range: str = Field(description="데이터 값 행렬 범위")


class LocalVlmSheetDecisionDTO(ModuleDTO):
    sheet_name: str
    tables: List[LocalVlmTableDecisionDTO]


class LocalVlmStructureDetectorOutput(SpreadsheetStructureOutput):
    pass


def _bounds(value: str, field_name: str) -> CellBounds:
    try:
        min_column, min_row, max_column, max_row = range_boundaries(value)
    except (TypeError, ValueError) as error:
        raise ModuleExecutionError(f"VLM {field_name} 범위가 올바르지 않습니다: {value}") from error
    if None in (min_column, min_row, max_column, max_row):
        raise ModuleExecutionError(f"VLM {field_name}은 셀 사각형 범위여야 합니다: {value}")
    return CellBounds(int(min_row), int(max_row), int(min_column), int(max_column))


def _contains(outer: CellBounds, inner: CellBounds) -> bool:
    return (
        outer.min_row <= inner.min_row <= inner.max_row <= outer.max_row
        and outer.min_column <= inner.min_column <= inner.max_column <= outer.max_column
    )


def _visible_bounds(
    bounds: CellBounds,
    visibility: WorksheetVisibility,
    field_name: str,
) -> CellBounds:
    rows = [
        row
        for row in range(bounds.min_row, bounds.max_row + 1)
        if not visibility.row_hidden(row)
    ]
    columns = [
        column
        for column in range(bounds.min_column, bounds.max_column + 1)
        if not visibility.column_hidden(column)
    ]
    if not rows or not columns:
        raise ModuleExecutionError(
            f"VLM {field_name}에 표시된 셀이 없습니다: {bounds.excel_range}"
        )
    return CellBounds(rows[0], rows[-1], columns[0], columns[-1])


class LocalVlmStructureDetectorModule(ExecutableModule):
    definition = ModuleDefinition(
        type="local_vlm_structure_detector",
        label="Local VLM Table Structure Detector",
        category="Logic",
        description="셀 타입 오버레이 이미지와 좌표·타입·값 컨텍스트를 로컬 VLM에 함께 전달해 표와 계층 헤더 영역을 식별합니다.",
        inputs=["input"],
        outputs=["output"],
        config_fields=[
            "model",
            "max_rows",
            "max_columns",
            "max_context_cells",
            "context_window",
            "timeout_seconds",
            "validation_retries",
            "system_prompt",
            "user_prompt_template",
        ],
        raw_output=True,
        version="4",
    )
    input_model = LocalVlmStructureDetectorInputDTO
    config_model = LocalVlmStructureDetectorConfigDTO
    execution_model = LocalVlmStructureDetectorExecutionDTO
    output_model = LocalVlmStructureDetectorOutput

    def __init__(
        self,
        vision_client: LocalVisionClient | None = None,
        catalog: WorkbookCatalog | None = None,
        renderer: ExcelSheetRenderer | None = None,
        processed_dir: Path = PROCESSED_DATA_DIR,
        artifact_dir: Path = SPREADSHEET_ARTIFACT_DIR,
    ) -> None:
        self.vision_client = vision_client or OllamaVisionClient()
        self.catalog = catalog or WorkbookCatalog(processed_dir)
        self.renderer = renderer or ExcelSheetRenderer()
        self.artifact_dir = artifact_dir

    @staticmethod
    def _region(
        region_id: str,
        region_type: str,
        bounds: CellBounds,
        layout,
        parent_ids: List[str],
    ) -> Dict[str, Any]:
        return {
            "region_id": region_id,
            "type": region_type,
            "excel_range": bounds.excel_range,
            "bbox_px": cell_bounds_bbox(bounds, layout),
            "rows": (bounds.min_row, bounds.max_row),
            "columns": (bounds.min_column, bounds.max_column),
            "parent_ids": parent_ids,
        }

    @staticmethod
    def _validate_table(
        table: LocalVlmTableDecisionDTO,
        layout,
        visibility: WorksheetVisibility,
    ) -> Dict[str, CellBounds | None]:
        named = {
            "excel_range": _bounds(table.excel_range, "excel_range"),
            "title_range": _bounds(table.title_range, "title_range") if table.title_range else None,
            "column_header_range": _bounds(table.column_header_range, "column_header_range") if table.column_header_range else None,
            "row_header_range": _bounds(table.row_header_range, "row_header_range") if table.row_header_range else None,
            "data_range": _bounds(table.data_range, "data_range"),
        }
        whole = cast(CellBounds, named["excel_range"])
        data = cast(CellBounds, named["data_range"])
        sheet = CellBounds(1, layout.max_row, 1, layout.max_column)
        if not _contains(sheet, whole):
            raise ModuleExecutionError(f"VLM 테이블 범위가 분석 시트를 벗어났습니다: {whole.excel_range}")
        for field_name, bounds in named.items():
            if field_name == "excel_range" or bounds is None:
                continue
            if not _contains(whole, bounds):
                raise ModuleExecutionError(f"VLM {field_name}이 테이블 범위를 벗어났습니다: {bounds.excel_range}")
        named = {
            field_name: (
                _visible_bounds(bounds, visibility, field_name)
                if bounds is not None
                else None
            )
            for field_name, bounds in named.items()
        }
        whole = cast(CellBounds, named["excel_range"])
        data = cast(CellBounds, named["data_range"])
        title = cast(Optional[CellBounds], named["title_range"])
        column_header = cast(Optional[CellBounds], named["column_header_range"])
        row_header = cast(Optional[CellBounds], named["row_header_range"])
        # Reconcile vertical overlaps between title, column_header, and data_range.
        # Vision models sometimes include header rows inside data_range or title span inside column_header.
        if title and column_header:
            if title.max_row >= column_header.min_row:
                if title.min_row < column_header.min_row:
                    title = CellBounds(
                        title.min_row,
                        column_header.min_row - 1,
                        title.min_column,
                        title.max_column,
                    )
                else:
                    title = None
                named["title_range"] = title

        header_bottom = None
        if column_header:
            header_bottom = column_header.max_row
        elif title:
            header_bottom = title.max_row

        if header_bottom is not None and header_bottom >= data.min_row:
            if data.max_row > header_bottom:
                data = CellBounds(
                    header_bottom + 1,
                    data.max_row,
                    data.min_column,
                    data.max_column,
                )
                named["data_range"] = data
            else:
                if column_header and column_header.min_row < data.min_row:
                    column_header = CellBounds(
                        column_header.min_row,
                        data.min_row - 1,
                        column_header.min_column,
                        column_header.max_column,
                    )
                    named["column_header_range"] = column_header
                else:
                    column_header = None
                    named["column_header_range"] = None
                if title and title.max_row >= data.min_row:
                    if title.min_row < data.min_row:
                        title = CellBounds(
                            title.min_row,
                            data.min_row - 1,
                            title.min_column,
                            title.max_column,
                        )
                    else:
                        title = None
                    named["title_range"] = title

        if title and title.max_row >= data.min_row:
            if title.min_row < data.min_row:
                title = CellBounds(
                    title.min_row,
                    data.min_row - 1,
                    title.min_column,
                    title.max_column,
                )
            else:
                title = None
            named["title_range"] = title

        if column_header and column_header.max_row >= data.min_row:
            if column_header.min_row < data.min_row:
                column_header = CellBounds(
                    column_header.min_row,
                    data.min_row - 1,
                    column_header.min_column,
                    column_header.max_column,
                )
            else:
                column_header = None
            named["column_header_range"] = column_header

        if row_header and row_header.max_column >= data.min_column:
            # Excel ranges are inclusive. Vision models sometimes include the
            # boundary label column in data_range, or include the first data
            # column in row_header_range. Preserve the side that already has
            # an unambiguous non-overlapping span and remove only the overlap.
            if row_header.min_column < data.min_column:
                row_header = CellBounds(
                    row_header.min_row,
                    row_header.max_row,
                    row_header.min_column,
                    data.min_column - 1,
                )
            elif row_header.max_column < data.max_column:
                data = CellBounds(
                    data.min_row,
                    data.max_row,
                    row_header.max_column + 1,
                    data.max_column,
                )
                named["data_range"] = data
            else:
                # A range fully contained in the data matrix provides no
                # coordinate evidence for a separate left-side row header.
                row_header = None
            named["row_header_range"] = row_header
        if row_header:
            row_header = CellBounds(
                data.min_row,
                data.max_row,
                row_header.min_column,
                row_header.max_column,
            )
            named["row_header_range"] = row_header
        if column_header:
            # The model often returns only the visible group-label span (for
            # example Actuals over F:I) even though the declared header rows
            # semantically apply to the complete E:J data matrix.  Preserve
            # its chosen first header row and data boundary, then normalize
            # the rectangular shared DTO across all data columns.
            column_header = CellBounds(
                column_header.min_row,
                data.min_row - 1,
                data.min_column,
                data.max_column,
            )
            named["column_header_range"] = column_header
        return named

    def _table_output(
        self,
        table_index: int,
        table: LocalVlmTableDecisionDTO,
        worksheet,
        layout,
    ) -> Dict[str, Any]:
        visibility = WorksheetVisibility.from_worksheet(worksheet)
        named = self._validate_table(table, layout, visibility)
        whole = cast(CellBounds, named["excel_range"])
        data = cast(CellBounds, named["data_range"])
        prefix = f"table_{table_index}"
        regions: List[Dict[str, Any]] = []
        parents: List[str] = []

        for key, region_type in (
            ("title_range", "title"),
            ("column_header_range", "column_header"),
        ):
            bounds = cast(Optional[CellBounds], named[key])
            if bounds is None:
                continue
            region_id = f"{prefix}_{region_type}"
            regions.append(self._region(region_id, region_type, bounds, layout, list(parents)))
            parents.append(region_id)

        row_header = cast(Optional[CellBounds], named["row_header_range"])
        if row_header is not None:
            row_header_id = f"{prefix}_row_header"
            regions.append(self._region(row_header_id, "row_header", row_header, layout, list(parents)))
            parents.append(row_header_id)
        regions.append(self._region(f"{prefix}_data", "data", data, layout, list(parents)))

        column_header = cast(Optional[CellBounds], named["column_header_range"])
        return {
            "sheet_name": worksheet.title,
            "table_index": table_index,
            "excel_range": whole.excel_range,
            "regions": regions,
            "header_tree": build_column_header_tree(
                worksheet,
                whole,
                column_header.min_row if column_header else data.min_row,
                column_header.max_row if column_header else data.min_row - 1,
                data.min_column,
            ),
        }

    def execute(self, payload: BaseModel) -> Dict[str, Any]:
        settings = cast(LocalVlmStructureDetectorExecutionDTO, payload)
        try:
            workbook_path = self.catalog.resolve(settings.file_name)
            current_hash = self.catalog.sha256(workbook_path)
        except (OSError, ValueError, WorkbookCatalogError) as error:
            raise ModuleExecutionError(str(error)) from error
        if current_hash != settings.workbook_hash:
            raise ModuleExecutionError("선택 이후 Excel 파일이 변경되었습니다. 파일 선택 모듈을 다시 실행하세요")

        formula_workbook = None
        value_workbook = None
        try:
            options = {"read_only": False, "keep_vba": workbook_path.suffix.lower() == ".xlsm"}
            formula_workbook = openpyxl.load_workbook(workbook_path, data_only=False, **options)
            value_workbook = openpyxl.load_workbook(workbook_path, data_only=True, **options)
            outputs: List[Dict[str, Any]] = []
            output_root = self.artifact_dir / current_hash[:16]
            response_schema = LocalVlmSheetDecisionDTO.model_json_schema()

            for sheet_name in settings.sheet_names:
                if sheet_name not in formula_workbook.sheetnames:
                    raise ModuleExecutionError(f"Excel 시트를 찾을 수 없습니다: {sheet_name}")
                formula_sheet = formula_workbook[sheet_name]
                value_sheet = value_workbook[sheet_name]
                if not worksheet_visible(formula_sheet) or not worksheet_visible(value_sheet):
                    raise ModuleExecutionError(
                        f"숨겨진 Excel 시트는 분석할 수 없습니다: {sheet_name}"
                    )
                safe_sheet = _safe_name(sheet_name)
                rendered_path = output_root / "rendered" / f"{safe_sheet}.png"
                typed_path = output_root / "typed" / f"{safe_sheet}.png"
                layout = self.renderer.render(value_sheet, rendered_path, settings.max_rows, settings.max_columns)
                try:
                    cells = collect_non_empty_cells(
                        formula_sheet,
                        value_sheet,
                        layout,
                        settings.max_context_cells,
                    )
                    sheet_context = compact_sheet_context(sheet_name, layout, cells)
                except ValueError as error:
                    raise ModuleExecutionError(str(error)) from error
                render_cell_type_overlay(rendered_path, typed_path, layout, cells)
                prompt = settings.user_prompt_template.replace("{sheet_name}", sheet_name).replace(
                    "{sheet_context}",
                    json.dumps(sheet_context, ensure_ascii=False, separators=(",", ":")),
                )
                previous_response = ""
                validation_error = ""
                for attempt in range(settings.validation_retries + 1):
                    correction = ""
                    if attempt:
                        correction = (
                            "\n\nYour previous JSON violated a coordinate rule: "
                            f"{validation_error}\nPrevious JSON: {previous_response}\n"
                            "Return a corrected complete JSON object."
                        )
                    try:
                        response = self.vision_client.complete_structured(
                            settings.model,
                            settings.system_prompt,
                            prompt + correction,
                            typed_path,
                            response_schema,
                            settings.context_window,
                            settings.timeout_seconds,
                        )
                        previous_response = response
                        decision = LocalVlmSheetDecisionDTO.model_validate_json(response)
                        if decision.sheet_name != sheet_name:
                            raise ModuleExecutionError(
                                f"로컬 VLM 응답 시트가 요청과 다릅니다: {decision.sheet_name} != {sheet_name}"
                            )
                        ordered = sorted(
                            decision.tables,
                            key=lambda item: (
                                _bounds(item.excel_range, "excel_range").min_row,
                                _bounds(item.excel_range, "excel_range").min_column,
                            ),
                        )
                        sheet_outputs = [
                            self._table_output(index, table, value_sheet, layout)
                            for index, table in enumerate(ordered, start=1)
                        ]
                        outputs.extend(sheet_outputs)
                        break
                    except OllamaVisionError as error:
                        raise ModuleExecutionError(str(error)) from error
                    except (ValueError, TypeError, json.JSONDecodeError, ModuleExecutionError) as error:
                        validation_error = str(error)
                        if attempt >= settings.validation_retries:
                            raise ModuleExecutionError(
                                f"로컬 VLM 응답을 좌표 규칙에 맞게 교정하지 못했습니다: {validation_error}"
                            ) from error
        finally:
            if formula_workbook is not None:
                formula_workbook.close()
            if value_workbook is not None:
                value_workbook.close()

        return {
            "file_name": workbook_path.name,
            "workbook_hash": current_hash,
            "tables": outputs,
        }
