"""Merge overlapping table decisions produced by exhaustive image tiles."""

from __future__ import annotations

from dataclasses import dataclass
from typing import Iterable, List, Optional, Sequence

from openpyxl.utils.cell import range_boundaries

from .table_geometry import CellBounds


@dataclass(frozen=True)
class TableFragment:
    excel_range: str
    title_range: Optional[str]
    column_header_range: Optional[str]
    row_header_range: Optional[str]
    data_range: str


def parse_excel_range(value: str) -> CellBounds:
    min_column, min_row, max_column, max_row = range_boundaries(value)
    if min_column is None or min_row is None or max_column is None or max_row is None:
        raise ValueError(f"셀 사각형 범위가 아닙니다: {value}")
    return CellBounds(min_row, max_row, min_column, max_column)


def _union(bounds: Sequence[CellBounds]) -> CellBounds:
    return CellBounds(
        min(item.min_row for item in bounds),
        max(item.max_row for item in bounds),
        min(item.min_column for item in bounds),
        max(item.max_column for item in bounds),
    )


def _intersects(left: CellBounds, right: CellBounds) -> bool:
    return not (
        left.max_row < right.min_row
        or right.max_row < left.min_row
        or left.max_column < right.min_column
        or right.max_column < left.min_column
    )


def _optional_union(values: Iterable[Optional[str]]) -> Optional[CellBounds]:
    parsed = [parse_excel_range(value) for value in values if value]
    return _union(parsed) if parsed else None


def _merge_group(group: Sequence[TableFragment]) -> TableFragment:
    whole = _union([parse_excel_range(item.excel_range) for item in group])
    data = _union([parse_excel_range(item.data_range) for item in group])

    # Header fragments repeated in lower tiles must not be stretched through
    # the data matrix. Only regions that are valid relative to the globally
    # earliest data row/column participate in the canonical header ranges.
    title = _optional_union(
        value
        for item in group
        for value in [item.title_range]
        if value and parse_excel_range(value).max_row < data.min_row
    )
    column_header = _optional_union(
        value
        for item in group
        for value in [item.column_header_range]
        if value and parse_excel_range(value).max_row < data.min_row
    )
    row_header = _optional_union(
        value
        for item in group
        for value in [item.row_header_range]
        if value and parse_excel_range(value).max_column < data.min_column
    )
    return TableFragment(
        excel_range=whole.excel_range,
        title_range=title.excel_range if title else None,
        column_header_range=column_header.excel_range if column_header else None,
        row_header_range=row_header.excel_range if row_header else None,
        data_range=data.excel_range,
    )


def merge_table_fragments(fragments: Sequence[TableFragment]) -> List[TableFragment]:
    """Merge fragments only when their declared table rectangles overlap."""

    if not fragments:
        return []
    parents = list(range(len(fragments)))
    bounds = [parse_excel_range(item.excel_range) for item in fragments]

    def find(index: int) -> int:
        while parents[index] != index:
            parents[index] = parents[parents[index]]
            index = parents[index]
        return index

    def union(left: int, right: int) -> None:
        left_root = find(left)
        right_root = find(right)
        if left_root != right_root:
            parents[right_root] = left_root

    for left in range(len(fragments)):
        for right in range(left + 1, len(fragments)):
            if _intersects(bounds[left], bounds[right]):
                union(left, right)

    groups: dict[int, List[TableFragment]] = {}
    for index, fragment in enumerate(fragments):
        groups.setdefault(find(index), []).append(fragment)
    merged = [_merge_group(group) for group in groups.values()]
    return sorted(
        merged,
        key=lambda item: (
            parse_excel_range(item.excel_range).min_row,
            parse_excel_range(item.excel_range).min_column,
        ),
    )
