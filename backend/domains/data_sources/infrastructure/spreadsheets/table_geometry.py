"""Spreadsheet table geometry helpers."""

from dataclasses import dataclass
from typing import List, Sequence, Tuple

from openpyxl.utils import get_column_letter

from .cell_visibility import WorksheetVisibility

COL_UNIT_PX = 7.5
ROW_POINT_TO_PX = 1.33
DEFAULT_COLUMN_WIDTH = 11.0
DEFAULT_ROW_HEIGHT = 15.0
MIN_COLUMN_PX = 35.0
MIN_ROW_PX = 14.0


@dataclass(frozen=True)
class SheetLayout:
    max_row: int
    max_column: int
    column_widths: List[float]
    row_heights: List[float]
    x_offsets: List[float]
    y_offsets: List[float]

    @property
    def width(self) -> int:
        return max(1, round(self.x_offsets[-1]))

    @property
    def height(self) -> int:
        return max(1, round(self.y_offsets[-1]))


@dataclass(frozen=True)
class CellBounds:
    min_row: int
    max_row: int
    min_column: int
    max_column: int

    @property
    def excel_range(self) -> str:
        return (
            f"{get_column_letter(self.min_column)}{self.min_row}:"
            f"{get_column_letter(self.max_column)}{self.max_row}"
        )


def compute_sheet_layout(worksheet, max_rows: int, max_columns: int) -> SheetLayout:
    max_row = min(worksheet.max_row or 1, max_rows)
    max_column = min(worksheet.max_column or 1, max_columns)

    visibility = WorksheetVisibility.from_worksheet(worksheet)
    column_widths: List[float] = []
    for column in range(1, max_column + 1):
        dimension = worksheet.column_dimensions.get(get_column_letter(column))
        if visibility.column_hidden(column):
            column_widths.append(0.0)
        else:
            width = (
                dimension.width
                if dimension and dimension.width
                else DEFAULT_COLUMN_WIDTH
            ) or DEFAULT_COLUMN_WIDTH
            column_widths.append(max(float(width) * COL_UNIT_PX, MIN_COLUMN_PX))

    row_heights: List[float] = []
    for row in range(1, max_row + 1):
        dimension = worksheet.row_dimensions.get(row)
        if visibility.row_hidden(row):
            row_heights.append(0.0)
        else:
            height = (
                dimension.height
                if dimension and dimension.height
                else DEFAULT_ROW_HEIGHT
            ) or DEFAULT_ROW_HEIGHT
            row_heights.append(max(float(height) * ROW_POINT_TO_PX, MIN_ROW_PX))

    x_offsets = [0.0]
    for width in column_widths:
        x_offsets.append(x_offsets[-1] + width)
    y_offsets = [0.0]
    for height in row_heights:
        y_offsets.append(y_offsets[-1] + height)

    return SheetLayout(
        max_row=max_row,
        max_column=max_column,
        column_widths=column_widths,
        row_heights=row_heights,
        x_offsets=x_offsets,
        y_offsets=y_offsets,
    )


def _nearest_boundary(
    value: float,
    offsets: Sequence[float],
    sizes: Sequence[float],
    end: bool,
) -> int:
    if len(offsets) < 2:
        return 1
    candidates = [
        index
        for index in range(1, len(offsets))
        if sizes[index - 1] > 0
    ]
    if not candidates:
        raise ValueError("표시된 행 또는 열이 없어 픽셀 영역을 셀로 변환할 수 없습니다")
    return min(
        candidates,
        key=lambda index: abs(offsets[index if end else index - 1] - value),
    )


def bbox_to_cell_bounds(
    bbox: Sequence[float],
    layout: SheetLayout,
) -> CellBounds:
    if len(bbox) != 4:
        raise ValueError("테이블 bbox는 [x1, y1, x2, y2] 네 값이어야 합니다")
    x1, y1, x2, y2 = (float(value) for value in bbox)
    min_column = _nearest_boundary(
        max(0.0, x1), layout.x_offsets, layout.column_widths, end=False
    )
    max_column = max(
        min_column,
        _nearest_boundary(
            min(float(layout.width), x2),
            layout.x_offsets,
            layout.column_widths,
            end=True,
        ),
    )
    min_row = _nearest_boundary(
        max(0.0, y1), layout.y_offsets, layout.row_heights, end=False
    )
    max_row = max(
        min_row,
        _nearest_boundary(
            min(float(layout.height), y2),
            layout.y_offsets,
            layout.row_heights,
            end=True,
        ),
    )
    return CellBounds(min_row, max_row, min_column, max_column)


def cell_bounds_bbox(
    bounds: CellBounds, layout: SheetLayout
) -> Tuple[float, float, float, float]:
    return (
        layout.x_offsets[bounds.min_column - 1],
        layout.y_offsets[bounds.min_row - 1],
        layout.x_offsets[bounds.max_column],
        layout.y_offsets[bounds.max_row],
    )
