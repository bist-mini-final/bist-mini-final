"""Resolve an indexed spreadsheet cell to its visual sheet artifact."""


from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils.cell import coordinate_to_tuple
from PIL import Image

from modules.common.config import (
    DEFAULT_STRUCTURE_MAX_COLUMNS,
    DEFAULT_STRUCTURE_MAX_ROWS,
)

from .table_geometry import CellBounds, cell_bounds_bbox, compute_sheet_layout
from .workbook_catalog import WorkbookCatalog, WorkbookCatalogError

_CELL_COORD_PATTERN = re.compile(r"^[A-Za-z]{1,3}[1-9][0-9]*$")
_WORKBOOK_HASH_PATTERN = re.compile(r"^[0-9a-f]{64}$")


def safe_sheet_artifact_name(value: str) -> str:
    """Return the filename-safe worksheet name used by the renderer."""
    return "".join(
        "_" if character in '\\/*?:\"<>| ' else character
        for character in value
    ).strip("_")


def resolve_workbook_path(
    processed_dir: Path,
    *,
    file_name: str,
    workbook_hash: str,
) -> Path | None:
    """Resolve a source workbook and verify it still matches the indexed hash."""
    catalog = WorkbookCatalog(processed_dir)
    try:
        candidate = catalog.resolve(file_name)
    except WorkbookCatalogError:
        candidate = None
    if candidate is not None and catalog.sha256(candidate).casefold() == workbook_hash.casefold():
        return candidate

    for available_name in catalog.file_names():
        path = catalog.resolve(available_name)
        if catalog.sha256(path).casefold() == workbook_hash.casefold():
            return path
    return None


def _prepare_artifact_results(
    *,
    artifact_dir: Path,
    workbook_hash: str,
    sheet_name: str,
    cell_coords: list[str],
) -> tuple[dict[str, dict[str, Any]], list[str], int, int, bool]:
    normalized_hash = workbook_hash.strip().lower()
    safe_sheet = safe_sheet_artifact_name(sheet_name)
    valid_hash = bool(_WORKBOOK_HASH_PATTERN.fullmatch(normalized_hash))
    artifact_key = normalized_hash[:16] if valid_hash else "__invalid_workbook_hash__"
    rendered_path = artifact_dir / artifact_key / "rendered" / f"{safe_sheet}.png"
    typed_path = artifact_dir / artifact_key / "typed" / f"{safe_sheet}.png"
    rendered_available = valid_hash and rendered_path.is_file()
    typed_available = valid_hash and typed_path.is_file()
    normalized_coords = list(dict.fromkeys(coord.strip().upper() for coord in cell_coords))
    base: dict[str, Any] = {
        "rendered_available": rendered_available,
        "typed_available": typed_available,
        "image_width": None,
        "image_height": None,
        "cell_bbox_px": None,
        "unavailable_reason": None,
    }
    results = {coord: dict(base) for coord in normalized_coords}
    if not rendered_available:
        for result in results.values():
            result["unavailable_reason"] = "이 인덱스에는 원본 시트 이미지가 생성되어 있지 않습니다."
        return results, [], 0, 0, False

    with Image.open(rendered_path) as image:
        image_width, image_height = image.size
    for coord, result in results.items():
        result["image_width"] = image_width
        result["image_height"] = image_height
        if not _CELL_COORD_PATTERN.fullmatch(coord):
            result["unavailable_reason"] = "셀 좌표 형식이 올바르지 않아 강조 영역을 계산할 수 없습니다."
    valid_coords = [coord for coord in normalized_coords if _CELL_COORD_PATTERN.fullmatch(coord)]
    return results, valid_coords, image_width, image_height, True


def _populate_cell_bboxes(
    *,
    worksheet: Any,
    cell_coords: list[str],
    results: dict[str, dict[str, Any]],
    image_width: int,
    image_height: int,
) -> None:
    layout = compute_sheet_layout(
        worksheet,
        DEFAULT_STRUCTURE_MAX_ROWS,
        DEFAULT_STRUCTURE_MAX_COLUMNS,
    )
    scale_x = image_width / layout.width
    scale_y = image_height / layout.height
    merged_ranges = list(worksheet.merged_cells.ranges)
    for coord in cell_coords:
        result = results[coord]
        row_index, column_index = coordinate_to_tuple(coord)
        if row_index > layout.max_row or column_index > layout.max_column:
            result["unavailable_reason"] = "시트 이미지 렌더링 범위 밖의 셀입니다."
            continue
        bounds = CellBounds(row_index, row_index, column_index, column_index)
        for merged_range in merged_ranges:
            if coord in merged_range:
                bounds = CellBounds(
                    merged_range.min_row,
                    merged_range.max_row,
                    merged_range.min_col,
                    merged_range.max_col,
                )
                break
        raw_bbox = cell_bounds_bbox(bounds, layout)
        result["cell_bbox_px"] = [
            min(float(image_width), max(0.0, round(raw_bbox[0] * scale_x, 2))),
            min(float(image_height), max(0.0, round(raw_bbox[1] * scale_y, 2))),
            min(float(image_width), max(0.0, round(raw_bbox[2] * scale_x, 2))),
            min(float(image_height), max(0.0, round(raw_bbox[3] * scale_y, 2))),
        ]


def locate_cell_artifacts(
    *,
    processed_dir: Path,
    artifact_dir: Path,
    file_name: str,
    workbook_hash: str,
    sheet_name: str,
    cell_coords: list[str],
) -> dict[str, dict[str, Any]]:
    """Return pixel bounds for many cells while opening the workbook only once."""
    normalized_hash = workbook_hash.strip().lower()
    results, valid_coords, image_width, image_height, rendered_available = (
        _prepare_artifact_results(
            artifact_dir=artifact_dir,
            workbook_hash=workbook_hash,
            sheet_name=sheet_name,
            cell_coords=cell_coords,
        )
    )
    if not rendered_available:
        return results
    if not valid_coords:
        return results
    workbook_path = resolve_workbook_path(
        processed_dir,
        file_name=file_name,
        workbook_hash=normalized_hash,
    )
    if workbook_path is None:
        for coord in valid_coords:
            results[coord]["unavailable_reason"] = "인덱싱에 사용한 원본 Excel 파일을 찾을 수 없습니다."
        return results

    workbook = openpyxl.load_workbook(
        workbook_path,
        read_only=False,
        data_only=True,
        keep_vba=workbook_path.suffix.lower() == ".xlsm",
    )
    try:
        if sheet_name not in workbook.sheetnames:
            for coord in valid_coords:
                results[coord]["unavailable_reason"] = "원본 Excel 파일에서 해당 시트를 찾을 수 없습니다."
            return results
        worksheet = workbook[sheet_name]
        _populate_cell_bboxes(
            worksheet=worksheet,
            cell_coords=valid_coords,
            results=results,
            image_width=image_width,
            image_height=image_height,
        )
        return results
    finally:
        workbook.close()


def locate_cell_artifact(
    *,
    processed_dir: Path,
    artifact_dir: Path,
    file_name: str,
    workbook_hash: str,
    sheet_name: str,
    cell_coord: str,
) -> dict[str, Any]:
    """Return image availability and exact pixel bounds for one source cell."""
    normalized_coord = cell_coord.strip().upper()
    return locate_cell_artifacts(
        processed_dir=processed_dir,
        artifact_dir=artifact_dir,
        file_name=file_name,
        workbook_hash=workbook_hash,
        sheet_name=sheet_name,
        cell_coords=[normalized_coord],
    )[normalized_coord]
