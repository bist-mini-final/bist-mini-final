"""Spreadsheet cell semantics."""

from __future__ import annotations

from datetime import date, datetime
from typing import Any, Dict, List, Tuple

from openpyxl.cell.cell import MergedCell
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import get_column_letter

from .cell_visibility import WorksheetVisibility
from .table_geometry import SheetLayout

CELL_TYPE_COLORS = {
    "text": "#F59E0B",
    "number": "#22C55E",
    "date": "#3B82F6",
    "boolean": "#8B5CF6",
    "error": "#EF4444",
    "formula": "#64748B",
}


def _display_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _merged_origins(worksheet) -> Dict[Tuple[int, int], str]:
    return {
        (cell_range.min_row, cell_range.min_col): str(cell_range)
        for cell_range in worksheet.merged_cells.ranges
    }


def _semantic_type(formula_cell, value_cell) -> str:
    if formula_cell.data_type == "e" or value_cell.data_type == "e":
        return "error"
    value = value_cell.value
    number_format = str(value_cell.number_format or formula_cell.number_format or "")
    if isinstance(value, (date, datetime)) or (
        isinstance(value, (int, float))
        and not isinstance(value, bool)
        and is_date_format(number_format)
    ):
        return "date"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "number"
    if value is not None and str(value).strip():
        return "text"
    if formula_cell.data_type == "f":
        return "formula"
    return "text"


def _cell_record(
    formula_worksheet,
    value_worksheet,
    merged: Dict[Tuple[int, int], str],
    row: int,
    column: int,
) -> Dict[str, Any] | None:
    formula_cell = formula_worksheet.cell(row=row, column=column)
    if isinstance(formula_cell, MergedCell):
        return None
    value_cell = value_worksheet.cell(row=row, column=column)
    raw_value = formula_cell.value
    value = value_cell.value
    if value is None and formula_cell.data_type != "f":
        value = raw_value
    if value is None and formula_cell.data_type != "f":
        return None
    if isinstance(value, str) and not value.strip() and formula_cell.data_type != "f":
        return None
    record: Dict[str, Any] = {
        "coord": f"{get_column_letter(column)}{row}",
        "row": row,
        "column": column,
        "type": _semantic_type(formula_cell, value_cell),
        "value": _display_value(value),
    }
    if formula_cell.data_type == "f":
        record.update(is_formula=True, formula=str(raw_value))
    if merged_range := merged.get((row, column)):
        record["merged_range"] = merged_range
    return record


def collect_non_empty_cells(
    formula_worksheet,
    value_worksheet,
    layout: SheetLayout,
    max_context_cells: int,
) -> List[Dict[str, Any]]:
    """Return compact coordinate-backed cell facts for one rendered sheet."""

    merged = _merged_origins(formula_worksheet)
    formula_visibility = WorksheetVisibility.from_worksheet(formula_worksheet)
    value_visibility = WorksheetVisibility.from_worksheet(value_worksheet)
    records: List[Dict[str, Any]] = []
    for row in range(1, layout.max_row + 1):
        if (
            layout.row_heights[row - 1] <= 0
            or formula_visibility.row_hidden(row)
            or value_visibility.row_hidden(row)
        ):
            continue
        for column in range(1, layout.max_column + 1):
            if (
                layout.column_widths[column - 1] <= 0
                or formula_visibility.column_hidden(column)
                or value_visibility.column_hidden(column)
            ):
                continue
            record = _cell_record(
                formula_worksheet,
                value_worksheet,
                merged,
                row,
                column,
            )
            if record is None:
                continue
            records.append(record)
            if len(records) > max_context_cells:
                raise ValueError(
                    f"시트 {formula_worksheet.title}의 값 셀이 {max_context_cells}개를 넘습니다. "
                    "max_context_cells를 늘리거나 분석 행·열 범위를 줄이세요"
                )
    return records


def compact_sheet_context(
    sheet_name: str,
    layout: SheetLayout,
    cells: List[Dict[str, Any]],
) -> Dict[str, Any]:
    """Group cell facts by row to reduce local-model prompt tokens."""

    visible_rows = [index for index, height in enumerate(layout.row_heights, start=1) if height > 0]
    visible_columns = [
        index for index, width in enumerate(layout.column_widths, start=1) if width > 0
    ]
    if not visible_rows or not visible_columns:
        raise ValueError(f"시트 {sheet_name}에 표시된 셀 영역이 없습니다")
    rows: Dict[int, List[List[Any]]] = {}
    for cell in cells:
        flags: Dict[str, Any] = {}
        for key in ("is_formula", "formula", "merged_range"):
            if key in cell:
                flags[key] = cell[key]
        item: List[Any] = [cell["coord"], cell["type"], cell["value"]]
        if flags:
            item.append(flags)
        rows.setdefault(int(cell["row"]), []).append(item)
    return {
        "sheet_name": sheet_name,
        "sheet_range": (
            f"{get_column_letter(visible_columns[0])}{visible_rows[0]}:"
            f"{get_column_letter(visible_columns[-1])}{visible_rows[-1]}"
        ),
        "cell_count": len(cells),
        "cell_tuple": ["excel_coord", "value_type", "value", "optional_flags"],
        "rows": [{"row": row, "cells": row_cells} for row, row_cells in sorted(rows.items())],
    }
