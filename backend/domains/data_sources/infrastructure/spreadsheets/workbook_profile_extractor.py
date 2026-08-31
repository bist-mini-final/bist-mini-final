"""Deterministic workbook profile extraction from original spreadsheet cells."""

from __future__ import annotations

import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from hashlib import sha256
from pathlib import Path
from typing import Any, Iterable, Literal

import openpyxl

from backend.domains.data_sources.domain.workbook_profiles import (
    WorkbookAmountScale,
    WorkbookPeriodKind,
    WorkbookPeriodProfile,
    WorkbookProfile,
    WorkbookProfileEvidence,
    WorkbookSheetProfile,
    WorkbookSheetRole,
)
from modules.structure.luna_vlm_structure_detector import SpreadsheetStructureOutput

WORKBOOK_PROFILE_VERSION = "1"

_CURRENCY_PATTERNS: tuple[tuple[re.Pattern[str], str], ...] = (
    (re.compile(r"\b(?:u\.?s\.?\s*dollars?|usd)\b", re.I), "USD"),
    (re.compile(r"\b(?:euros?|eur)\b", re.I), "EUR"),
    (re.compile(r"\b(?:korean\s+won|krw)\b|원화", re.I), "KRW"),
    (re.compile(r"\b(?:japanese\s+yen|jpy)\b|엔화", re.I), "JPY"),
    (re.compile(r"\b(?:british\s+pounds?|pound\s+sterling|gbp)\b", re.I), "GBP"),
    (re.compile(r"(?:data\s+in|amounts?\s+in)\s*\(\s*\$", re.I), "USD"),
    (re.compile(r"\$\s*(?:k|m|mm|bn|b)\b", re.I), "USD"),
)
_SCALE_PATTERNS: tuple[tuple[re.Pattern[str], WorkbookAmountScale], ...] = (
    (re.compile(r"\b(?:billions?|bn)\b|\$\s*(?:bn|b)\b", re.I), WorkbookAmountScale.BILLIONS),
    (re.compile(r"\b(?:millions?|mm)\b|\$\s*(?:m|mm)\b", re.I), WorkbookAmountScale.MILLIONS),
    (re.compile(r"\b(?:thousands?|000s)\b|\$\s*k\b|천\s*단위", re.I), WorkbookAmountScale.THOUSANDS),
    (re.compile(r"\b(?:ones)\b|원\s*단위", re.I), WorkbookAmountScale.ONES),
)
_ROLE_TERMS: dict[WorkbookSheetRole, tuple[str, ...]] = {
    WorkbookSheetRole.INCOME_STATEMENT: (
        "income statement", "statement of operations", "profit and loss", "revenue",
        "operating income", "net income", "손익계산서", "매출액", "영업이익",
    ),
    WorkbookSheetRole.BALANCE_SHEET: (
        "balance sheet", "financial position", "total assets", "total liabilities",
        "total equity", "재무상태표", "총자산", "총부채", "자본총계",
    ),
    WorkbookSheetRole.CASH_FLOW: (
        "cash flow", "operating cash flow", "capital expenditure", "free cash flow",
        "investing activities", "현금흐름표", "영업활동현금흐름",
    ),
    WorkbookSheetRole.KEY_STATS: (
        "key stats", "key statistics", "enterprise value", "tev / ebitda",
        "valuation multiples", "주요 지표", "핵심 지표",
    ),
}
_DATE_TEXT = re.compile(r"(?<!\d)(20\d{2}|19\d{2})[-./](0?[1-9]|1[0-2])[-./](0?[1-9]|[12]\d|3[01])(?!\d)")
_FY_TEXT = re.compile(r"\bFY\s*[-_ ]?(20\d{2}|19\d{2})\b|(?<!\d)(20\d{2}|19\d{2})\s*(?:A|FY)\b", re.I)


@dataclass(frozen=True, slots=True)
class _PeriodCandidate:
    sheet_name: str
    cell_coord: str
    end_date: date | None
    year: int
    source_label: str
    kind: WorkbookPeriodKind


class WorkbookProfileExtractor:
    """Read explicit semantic facts without relying on fixed sheets or cells."""

    def extract(
        self,
        *,
        workbook_path: Path,
        structure: SpreadsheetStructureOutput,
        index_id: str,
    ) -> WorkbookProfile:
        workbook = openpyxl.load_workbook(
            workbook_path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        try:
            selected_sheets = set(structure.sheet_names) or set(workbook.sheetnames)
            cells_by_sheet = {
                worksheet.title: tuple(self._populated_cells(worksheet))
                for worksheet in workbook.worksheets
                if worksheet.sheet_state == "visible"
                and worksheet.title in selected_sheets
            }
        finally:
            workbook.close()

        evidence: dict[str, WorkbookProfileEvidence] = {}
        sheets = tuple(
            self._sheet_profile(sheet_name, cells, evidence)
            for sheet_name, cells in cells_by_sheet.items()
        )
        currency = self._discover_currency(cells_by_sheet, evidence)
        amount_scale = self._discover_scale(cells_by_sheet, evidence)
        periods = self._discover_periods(cells_by_sheet, structure, evidence)
        diagnostics = tuple(
            message
            for missing, message in (
                (not periods, "periods_missing"),
                (currency is None, "currency_missing"),
                (amount_scale is None, "amount_scale_missing"),
            )
            if missing
        )
        return WorkbookProfile(
            profile_version=WORKBOOK_PROFILE_VERSION,
            file_name=structure.file_name,
            workbook_hash=structure.workbook_hash,
            index_id=index_id,
            status="partial" if diagnostics else "ready",
            currency=currency,
            amount_scale=amount_scale,
            periods=periods,
            sheets=sheets,
            evidence=tuple(evidence.values()),
            diagnostics=diagnostics,
        )

    @staticmethod
    def _populated_cells(worksheet: Any) -> Iterable[tuple[str, int, int, Any]]:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value not in (None, ""):
                    yield cell.coordinate, cell.row, cell.column, cell.value

    def _sheet_profile(
        self,
        sheet_name: str,
        cells: tuple[tuple[str, int, int, Any], ...],
        evidence: dict[str, WorkbookProfileEvidence],
    ) -> WorkbookSheetProfile:
        searchable = self._normalize_text(
            " ".join([sheet_name, *(str(value) for _, _, _, value in cells[:400])])
        )
        normalized_name = self._normalize_text(sheet_name)
        scores = {
            role: sum(
                4 if self._normalize_text(term) in normalized_name else 1
                for term in terms
                if self._normalize_text(term) in searchable
            )
            for role, terms in _ROLE_TERMS.items()
        }
        role, score = max(scores.items(), key=lambda item: item[1])
        if score == 0:
            return WorkbookSheetProfile(
                sheet_name=sheet_name,
                role=WorkbookSheetRole.UNKNOWN,
                confidence=0,
            )
        matched = next(
            (
                (coord, str(value))
                for coord, _, _, value in cells
                if any(self._normalize_text(term) in self._normalize_text(str(value)) for term in _ROLE_TERMS[role])
            ),
            ("A1", sheet_name),
        )
        evidence_id = self._add_evidence(
            evidence,
            kind="sheet_role",
            sheet_name=sheet_name,
            cell_coord=matched[0],
            source_text=matched[1],
        )
        confidence = min(0.98, 0.55 + (score * 0.1))
        return WorkbookSheetProfile(
            sheet_name=sheet_name,
            role=role,
            confidence=confidence,
            evidence_ids=(evidence_id,),
        )

    def _discover_currency(
        self,
        cells_by_sheet: dict[str, tuple[tuple[str, int, int, Any], ...]],
        evidence: dict[str, WorkbookProfileEvidence],
    ) -> str | None:
        matches: list[tuple[str, str, str]] = []
        for sheet_name, cells in cells_by_sheet.items():
            for coord, _, _, value in cells:
                text = str(value).strip()
                for pattern, currency in _CURRENCY_PATTERNS:
                    if pattern.search(text):
                        matches.append((currency, sheet_name, coord))
                        self._add_evidence(evidence, "currency", sheet_name, coord, text)
                        break
        return self._unambiguous_majority(item[0] for item in matches)

    def _discover_scale(
        self,
        cells_by_sheet: dict[str, tuple[tuple[str, int, int, Any], ...]],
        evidence: dict[str, WorkbookProfileEvidence],
    ) -> WorkbookAmountScale | None:
        matches: list[WorkbookAmountScale] = []
        for sheet_name, cells in cells_by_sheet.items():
            for coord, _, _, value in cells:
                text = str(value).strip()
                for pattern, scale in _SCALE_PATTERNS:
                    if pattern.search(text):
                        matches.append(scale)
                        self._add_evidence(evidence, "scale", sheet_name, coord, text)
                        break
        return self._unambiguous_majority(matches)

    def _discover_periods(
        self,
        cells_by_sheet: dict[str, tuple[tuple[str, int, int, Any], ...]],
        structure: SpreadsheetStructureOutput,
        evidence: dict[str, WorkbookProfileEvidence],
    ) -> tuple[WorkbookPeriodProfile, ...]:
        header_cells = self._header_cell_keys(structure)
        candidates = tuple(
            candidate
            for sheet_name, cells in cells_by_sheet.items()
            for candidate in self._sheet_period_candidates(sheet_name, cells, header_cells)
        )
        grouped = self._group_period_candidates(candidates, evidence)
        return self._period_profiles(grouped)

    def _sheet_period_candidates(
        self,
        sheet_name: str,
        cells: tuple[tuple[str, int, int, Any], ...],
        header_cells: set[tuple[str, int, int]],
    ) -> tuple[_PeriodCandidate, ...]:
        row_periods: dict[int, list[tuple[str, int, int, date | None, int, str]]] = defaultdict(list)
        cell_values = {(row, column): value for _, row, column, value in cells}
        for coord, row, column, value in cells:
            parsed = self._period_value(value)
            if parsed is not None:
                end_date, year, source_label = parsed
                row_periods[row].append((coord, row, column, end_date, year, source_label))
        return tuple(
            candidate
            for row, row_candidates in row_periods.items()
            for candidate in self._row_period_candidates(
                sheet_name,
                row,
                row_candidates,
                cell_values,
                header_cells,
            )
        )

    def _row_period_candidates(
        self,
        sheet_name: str,
        row: int,
        row_candidates: list[tuple[str, int, int, date | None, int, str]],
        cell_values: dict[tuple[int, int], Any],
        header_cells: set[tuple[str, int, int]],
    ) -> tuple[_PeriodCandidate, ...]:
        unique_dates = {item[3] for item in row_candidates if item[3] is not None}
        has_semantics = self._row_has_period_semantics(cell_values, row)
        if len(row_candidates) >= 2 and len(unique_dates) == 1 and not has_semantics:
            return ()
        is_axis = len(unique_dates) >= 2 and has_semantics
        accepted: list[_PeriodCandidate] = []
        for coord, _, column, end_date, year, source_label in row_candidates:
            if not is_axis and (sheet_name, row, column) not in header_cells:
                continue
            if self._column_is_forecast(cell_values, row, column):
                continue
            kind = (
                WorkbookPeriodKind.LTM
                if self._column_has_ltm(cell_values, row, column)
                else WorkbookPeriodKind.FY
            )
            accepted.append(
                _PeriodCandidate(
                    sheet_name=sheet_name,
                    cell_coord=coord,
                    end_date=end_date,
                    year=year,
                    source_label=source_label,
                    kind=kind,
                )
            )
        return tuple(accepted)

    def _group_period_candidates(
        self,
        candidates: tuple[_PeriodCandidate, ...],
        evidence: dict[str, WorkbookProfileEvidence],
    ) -> dict[tuple[WorkbookPeriodKind, str], dict[str, Any]]:
        grouped: dict[tuple[WorkbookPeriodKind, str], dict[str, Any]] = {}
        for candidate in candidates:
            canonical_key = (
                candidate.end_date.isoformat()
                if candidate.end_date
                else str(candidate.year)
            )
            key = (candidate.kind, canonical_key)
            evidence_id = self._add_evidence(
                evidence,
                "period",
                candidate.sheet_name,
                candidate.cell_coord,
                candidate.source_label,
            )
            current = grouped.setdefault(
                key,
                {
                    "kind": candidate.kind,
                    "end_date": candidate.end_date,
                    "year": candidate.year,
                    "source_label": candidate.source_label,
                    "evidence_ids": [],
                },
            )
            if evidence_id not in current["evidence_ids"]:
                current["evidence_ids"].append(evidence_id)
        return grouped

    @staticmethod
    def _period_profiles(
        grouped: dict[tuple[WorkbookPeriodKind, str], dict[str, Any]],
    ) -> tuple[WorkbookPeriodProfile, ...]:
        dated_years = {
            (item["kind"], item["year"])
            for item in grouped.values()
            if item["end_date"] is not None
        }
        return tuple(
            WorkbookPeriodProfile(
                period_id=f"{item['kind'].value}-{key[1]}",
                kind=item["kind"],
                label=("LTM " if item["kind"] is WorkbookPeriodKind.LTM else "FY")
                + str(item["year"]),
                source_label=item["source_label"][:64],
                end_date=item["end_date"],
                ordinal=item["year"],
                evidence_ids=tuple(item["evidence_ids"][:8]),
            )
            for key, item in sorted(
                grouped.items(),
                key=lambda pair: (pair[1]["year"], pair[0][0].value),
            )
            if item["end_date"] is not None
            or (item["kind"], item["year"]) not in dated_years
        )

    @staticmethod
    def _header_cell_keys(
        structure: SpreadsheetStructureOutput,
    ) -> set[tuple[str, int, int]]:
        keys: set[tuple[str, int, int]] = set()
        for table in structure.tables:
            for region in table.regions:
                if region.type != "column_header":
                    continue
                for row in range(region.rows[0], region.rows[1] + 1):
                    for column in range(region.columns[0], region.columns[1] + 1):
                        keys.add((table.sheet_name, row, column))
        return keys

    @staticmethod
    def _column_has_ltm(
        values: dict[tuple[int, int], Any],
        row: int,
        column: int,
    ) -> bool:
        return any(
            "ltm" in str(values.get((nearby_row, column), "")).casefold()
            or "last twelve months" in str(values.get((nearby_row, column), "")).casefold()
            for nearby_row in range(max(1, row - 10), row + 5)
        )

    @staticmethod
    def _row_has_period_semantics(
        values: dict[tuple[int, int], Any],
        row: int,
    ) -> bool:
        row_text = " ".join(
            str(value)
            for (cell_row, _), value in values.items()
            if cell_row == row and isinstance(value, str)
        ).casefold()
        return bool(
            re.search(
                r"period\s*(?:ended|end|date)|fiscal\s*year|as\s+of|"
                r"reporting\s+date|기간|결산일|기준일|iq_period_end",
                row_text,
            )
        )

    @staticmethod
    def _column_is_forecast(
        values: dict[tuple[int, int], Any],
        row: int,
        column: int,
    ) -> bool:
        context = " ".join(
            str(values.get((nearby_row, column), ""))
            for nearby_row in range(max(1, row - 6), row + 1)
        ).casefold()
        return bool(re.search(r"estimates?|forecasts?|projections?|fy\s*\+|ntm", context))

    @staticmethod
    def _period_value(value: Any) -> tuple[date | None, int, str] | None:
        if isinstance(value, datetime):
            return value.date(), value.year, value.date().isoformat()
        if isinstance(value, date):
            return value, value.year, value.isoformat()
        if not isinstance(value, str):
            return None
        text = " ".join(value.split())
        match = _DATE_TEXT.search(text)
        if match:
            try:
                parsed = date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
            except ValueError:
                return None
            return parsed, parsed.year, match.group(0)
        match = _FY_TEXT.search(text)
        if match:
            year = int(match.group(1) or match.group(2))
            return None, year, match.group(0)
        return None

    @staticmethod
    def _unambiguous_majority(values: Iterable[Any]) -> Any | None:
        counts = Counter(values)
        if not counts:
            return None
        most_common = counts.most_common()
        if len(most_common) > 1 and most_common[0][1] == most_common[1][1]:
            return None
        return most_common[0][0]

    @staticmethod
    def _normalize_text(value: str) -> str:
        return re.sub(r"[_\W]+", " ", value.casefold()).strip()

    @staticmethod
    def _add_evidence(
        evidence: dict[str, WorkbookProfileEvidence],
        kind: Literal["currency", "scale", "period", "sheet_role"],
        sheet_name: str,
        cell_coord: str,
        source_text: str,
    ) -> str:
        identity = f"{kind}:{sheet_name}:{cell_coord}:{source_text}"
        evidence_id = "wpe-" + sha256(identity.encode("utf-8")).hexdigest()[:24]
        evidence.setdefault(
            evidence_id,
            WorkbookProfileEvidence(
                evidence_id=evidence_id,
                kind=kind,
                sheet_name=sheet_name,
                cell_coord=cell_coord,
                source_text=source_text[:2_000],
            ),
        )
        return evidence_id


__all__ = ["WORKBOOK_PROFILE_VERSION", "WorkbookProfileExtractor"]
