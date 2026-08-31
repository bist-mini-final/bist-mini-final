"""Session-scoped chat attachment storage and safe text extraction."""

from __future__ import annotations

import csv
import json
import zipfile
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from xml.etree import ElementTree

from openpyxl import load_workbook

from backend.domains.chatbot.application.attachments import StoredChatAttachment
from backend.domains.chatbot.domain.errors import (
    ChatAttachmentTooLargeError,
    ChatAttachmentValidationError,
    UnsupportedChatAttachmentError,
)

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_EXTRACTED_CHARS = 60_000
MAX_SHEET_EXTRACTED_CHARS = 6_000
SUPPORTED_SUFFIXES = frozenset({".txt", ".md", ".csv", ".json", ".xlsx", ".xlsm"})


def _limit(text: str) -> str:
    normalized = text.replace("\x00", "").strip()
    if len(normalized) <= MAX_EXTRACTED_CHARS:
        return normalized
    return (
        normalized[:MAX_EXTRACTED_CHARS]
        + "\n\n[첨부 파일 내용이 길어 처음 60,000자만 사용했습니다.]"
    )


def _xml_text(element: ElementTree.Element | None) -> str:
    if element is None:
        return ""
    return "".join(element.itertext()).strip()


def _extract_xlsx_xml(content: bytes) -> str:
    """Fallback reader for valid xlsx containers openpyxl cannot fully deserialize."""
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    lines: list[str] = []
    with zipfile.ZipFile(BytesIO(content)) as archive:
        shared: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
            shared = [_xml_text(item) for item in root.findall(f"{namespace}si")]
        sheets = sorted(
            (
                name
                for name in archive.namelist()
                if name.startswith("xl/worksheets/") and name.endswith(".xml")
            ),
            key=str.casefold,
        )
        for sheet in sheets:
            lines.append(f"[시트: {Path(sheet).stem}]")
            root = ElementTree.fromstring(archive.read(sheet))
            for row in root.findall(f".//{namespace}row"):
                values: list[str] = []
                for cell in row.findall(f"{namespace}c"):
                    value = _xml_text(cell.find(f"{namespace}v"))
                    cell_type = cell.get("t")
                    if cell_type == "s" and value.isdigit() and int(value) < len(shared):
                        value = shared[int(value)]
                    elif cell_type == "inlineStr":
                        value = _xml_text(cell.find(f"{namespace}is"))
                    if value:
                        values.append(value)
                if values:
                    lines.append(" | ".join(values))
                if sum(len(line) + 1 for line in lines) > MAX_EXTRACTED_CHARS:
                    return _limit("\n".join(lines))
    return _limit("\n".join(lines))


def _extract_json(content: bytes) -> str:
    try:
        document = json.loads(content.decode("utf-8-sig"))
        return _limit(json.dumps(document, ensure_ascii=False, indent=2))
    except (UnicodeDecodeError, json.JSONDecodeError):
        return _limit(content.decode("utf-8-sig", errors="replace"))


def _extract_csv(content: bytes) -> str:
    decoded = content.decode("utf-8-sig", errors="replace")
    rows = csv.reader(decoded.splitlines())
    return _limit("\n".join(" | ".join(cell.strip() for cell in row) for row in rows))


def _worksheet_lines(worksheet) -> list[str]:
    lines = [f"[시트: {worksheet.title}]"]
    for row in worksheet.iter_rows(values_only=True):
        values = [str(value).strip() for value in row if value is not None and str(value).strip()]
        if values:
            lines.append(" | ".join(values))
        if sum(len(line) + 1 for line in lines) > MAX_SHEET_EXTRACTED_CHARS:
            lines.append("[이 시트는 처음 6,000자만 사용했습니다.]")
            break
    return lines


def _extract_workbook(content: bytes) -> str:
    workbook = load_workbook(
        BytesIO(content),
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    lines: list[str] = []
    try:
        for worksheet in workbook.worksheets:
            if worksheet.sheet_state != "visible":
                continue
            lines.extend(_worksheet_lines(worksheet))
            if sum(len(line) + 1 for line in lines) > MAX_EXTRACTED_CHARS:
                return _limit("\n".join(lines))
    finally:
        workbook.close()
    return _limit("\n".join(lines))


def extract_text(file_name: str, content: bytes) -> str:
    """Extract a bounded, plain-text evidence context from supported local files."""
    suffix = Path(file_name).suffix.casefold()
    if suffix in {".txt", ".md"}:
        return _limit(content.decode("utf-8-sig", errors="replace"))
    if suffix == ".json":
        return _extract_json(content)
    if suffix == ".csv":
        return _extract_csv(content)
    if suffix in {".xlsx", ".xlsm"}:
        try:
            return _extract_workbook(content)
        except (TypeError, ValueError):
            return _extract_xlsx_xml(content)
    raise ValueError("지원하지 않는 파일 형식입니다")


class LocalChatAttachmentStorage:
    """Store bounded chat attachments and extract their text evidence locally."""

    def __init__(self, upload_dir: Path) -> None:
        self._upload_dir = upload_dir

    def save(
        self,
        *,
        file_name: str,
        content_type: str | None,
        content: bytes,
    ) -> StoredChatAttachment:
        if not file_name:
            raise ChatAttachmentValidationError("유효한 파일명이 필요합니다")
        safe_name = Path(file_name).name
        suffix = Path(safe_name).suffix.casefold()
        if suffix not in SUPPORTED_SUFFIXES:
            allowed = ", ".join(sorted(SUPPORTED_SUFFIXES))
            raise UnsupportedChatAttachmentError(f"지원 파일: {allowed}")
        if len(content) > MAX_UPLOAD_BYTES:
            raise ChatAttachmentTooLargeError("첨부 파일은 20MB를 초과할 수 없습니다")
        try:
            extracted = extract_text(safe_name, content)
            if not extracted:
                raise ChatAttachmentValidationError(
                    "질문 근거로 사용할 텍스트를 추출하지 못했습니다"
                )
            attachment_id = f"attachment-{uuid4().hex}"
            self._upload_dir.mkdir(parents=True, exist_ok=True)
            destination = self._upload_dir / f"{attachment_id}{suffix}"
            destination.write_bytes(content)
        except ChatAttachmentValidationError:
            raise
        except Exception as error:
            raise ChatAttachmentValidationError(
                f"첨부 파일을 읽지 못했습니다: {error}"
            ) from error
        return StoredChatAttachment(
            attachment_id=attachment_id,
            file_name=safe_name,
            content_type=content_type,
            file_size=len(content),
            storage_path=str(destination.resolve()),
            extracted_text=extracted,
        )


__all__ = [
    "LocalChatAttachmentStorage",
    "MAX_UPLOAD_BYTES",
    "SUPPORTED_SUFFIXES",
    "extract_text",
]
