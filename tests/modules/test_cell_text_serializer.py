from __future__ import annotations

from pathlib import Path

import openpyxl

from backend.domains.data_sources.infrastructure.spreadsheets.structured_cell_text import (
    UNKNOWN_FIELD,
    extract_resolved_cell_value,
    generate_header_combinations,
    resolved_cell_value,
    serialize_structured_cell,
)
from backend.domains.data_sources.infrastructure.spreadsheets.workbook_catalog import (
    WorkbookCatalog,
)
from modules.structure.cell_text_serializer import (
    CellTextSerializerConfigDTO,
    CellTextSerializerInputDTO,
    CellTextSerializerModule,
    CellTextSerializerOutput,
)


def test_resolved_cell_value_distinguishes_search_placeholder_from_evidence() -> None:
    assert resolved_cell_value("62753") == "62753"
    assert resolved_cell_value("?") is None
    assert resolved_cell_value("NA") is None
    assert (
        extract_resolved_cell_value(
            "Company: IBM | Sheet: Income_Statement | Row Header: Total Revenue | "
            "Column Header: 2024-12-31 | Cell Value: 62753"
        )
        == "62753"
    )
    assert extract_resolved_cell_value("Cell Value: ?") is None


def test_header_hierarchy_contract_keeps_search_and_evidence_variants() -> None:
    combinations = generate_header_combinations(
        ["Assets", "Current Assets", "Cash"],
        ["Fiscal Year", "2024-12-31"],
    )
    assert combinations[0] == (
        ["Assets", "Current Assets", "Cash"],
        ["Fiscal Year", "2024-12-31"],
    )
    assert (["Cash"], ["2024-12-31"]) in combinations
    assert (["Current Assets", "Cash"], ["Fiscal Year"]) in combinations

    row_headers, column_headers = combinations[0]
    search_text = serialize_structured_cell(
        "Balance_Sheet",
        row_headers,
        column_headers,
        UNKNOWN_FIELD,
        company_name="Nexora Labs",
    )
    evidence_text = serialize_structured_cell(
        "Balance_Sheet",
        row_headers,
        column_headers,
        "10081",
        company_name="Nexora Labs",
    )
    assert search_text.endswith("Cell Value: ?")
    assert extract_resolved_cell_value(search_text) is None
    assert extract_resolved_cell_value(evidence_text) == "10081"


def test_cell_text_serializer_execution(tmp_path: Path) -> None:
    wb_path = tmp_path / "test.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "IS"
    ws["B1"] = "FY2024"
    ws["C1"] = "FY2025"
    ws["A2"] = "Revenue"
    ws["B2"] = 100
    ws["C2"] = 120
    wb.save(wb_path)

    module = CellTextSerializerModule(catalog=WorkbookCatalog(tmp_path))
    current_hash = module.catalog.sha256(wb_path)

    input_dto = CellTextSerializerInputDTO.model_validate(
        {
            "file_name": "test.xlsx",
            "workbook_hash": current_hash,
            "sheet_names": ["IS"],
            "tables": [
                {
                    "sheet_name": "IS",
                    "table_index": 1,
                    "excel_range": "A1:C2",
                    "regions": [
                        {
                            "region_id": "table_1_column_header",
                            "type": "column_header",
                            "excel_range": "B1:C1",
                            "bbox_px": [0, 0, 1, 1],
                            "rows": [1, 1],
                            "columns": [2, 3],
                            "parent_ids": [],
                        },
                        {
                            "region_id": "table_1_row_header",
                            "type": "row_header",
                            "excel_range": "A2:A2",
                            "bbox_px": [0, 0, 1, 1],
                            "rows": [2, 2],
                            "columns": [1, 1],
                            "parent_ids": ["table_1_column_header"],
                        },
                        {
                            "region_id": "table_1_data",
                            "type": "data",
                            "excel_range": "B2:C2",
                            "bbox_px": [0, 0, 1, 1],
                            "rows": [2, 2],
                            "columns": [2, 3],
                            "parent_ids": [
                                "table_1_column_header",
                                "table_1_row_header",
                            ],
                        },
                    ],
                    "header_tree": [],
                }
            ],
        }
    )

    result = module.execute(input_dto, config=CellTextSerializerConfigDTO())
    validated = CellTextSerializerOutput.model_validate(result)

    assert validated.file_name == "test.xlsx"
    # Default variant_mode="both" creates 2 variants per cell (2 cells * 2 = 4 documents)
    assert len(validated.items) == 4
    assert validated.items[0].sheet_name == "IS"
    assert "Revenue" in validated.items[0].row_header


def test_cell_text_serializer_keeps_title_region_out_of_row_header(tmp_path: Path) -> None:
    wb_path = tmp_path / "ibm.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Income_Statement"
    ws["A1"] = "International Business Machines Corporation"
    ws["A2"] = "Source: S&P Capital IQ Pro"
    ws["B3"] = "2024-12-31"
    ws["A4"] = "Total Revenue"
    ws["B4"] = 62753
    wb.save(wb_path)

    module = CellTextSerializerModule(catalog=WorkbookCatalog(tmp_path))
    current_hash = module.catalog.sha256(wb_path)
    result = module.run(
        {
            "file_name": "ibm.xlsx",
            "workbook_hash": current_hash,
            "company_name": "IBM",
            "sheet_names": ["Income_Statement"],
            "tables": [
                {
                    "sheet_name": "Income_Statement",
                    "table_index": 1,
                    "excel_range": "A1:B4",
                    "regions": [
                        {
                            "region_id": "title",
                            "type": "title",
                            "excel_range": "A1:A2",
                            "bbox_px": [0, 0, 1, 1],
                            "rows": [1, 2],
                            "columns": [1, 1],
                            "parent_ids": [],
                        },
                        {
                            "region_id": "columns",
                            "type": "column_header",
                            "excel_range": "B3:B3",
                            "bbox_px": [0, 0, 1, 1],
                            "rows": [3, 3],
                            "columns": [2, 2],
                            "parent_ids": [],
                        },
                        {
                            "region_id": "rows",
                            "type": "row_header",
                            "excel_range": "A4:A4",
                            "bbox_px": [0, 0, 1, 1],
                            "rows": [4, 4],
                            "columns": [1, 1],
                            "parent_ids": [],
                        },
                        {
                            "region_id": "data",
                            "type": "data",
                            "excel_range": "B4:B4",
                            "bbox_px": [0, 0, 1, 1],
                            "rows": [4, 4],
                            "columns": [2, 2],
                            "parent_ids": [],
                        },
                    ],
                }
            ],
        }
    )

    assert len(result["items"]) == 2
    for item in result["items"]:
        assert item["row_header"] == ["Total Revenue"]
        assert item["text"].startswith("Company: IBM | Sheet: Income_Statement")
        assert "International Business Machines Corporation" not in item["text"]
        assert "Source: S&P Capital IQ Pro" not in item["text"]
