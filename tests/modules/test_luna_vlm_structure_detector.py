from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock

import openpyxl

from backend.providers.openai_responses import OpenAIResponseResult
from backend.storage.spreadsheets.cell_visibility import WorksheetVisibility
from backend.storage.spreadsheets.sheet_renderer import ExcelSheetRenderer
from backend.storage.spreadsheets.table_geometry import compute_sheet_layout
from backend.storage.spreadsheets.workbook_catalog import WorkbookCatalog
from modules.structure.luna_vlm_structure_detector import (
    LunaVlmStructureDetectorConfigDTO,
    LunaVlmStructureDetectorModule,
    SpreadsheetStructureOutput,
    VlmTableDecisionDTO,
)


def test_luna_detector_executes_on_tiny_workbook(tmp_path: Path) -> None:
    workbook_path = tmp_path / "tiny.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Sheet1"
    ws["A1"] = "Category"
    ws["B1"] = "2023"
    ws["A2"] = "Revenue"
    ws["B2"] = 1000
    wb.save(workbook_path)

    catalog = WorkbookCatalog(tmp_path)
    vision_client = MagicMock()
    vision_client.complete_vision_structured.return_value = OpenAIResponseResult(
        response_id="resp_vision",
        content=(
            '{"tables":[{"excel_range":"A1:B2","title_range":null,'
            '"column_header_range":"B1:B1","row_header_range":"A2:A2",'
            '"data_range":"B2:B2"}]}'
        ),
        usage={
            "prompt_tokens": 120,
            "completion_tokens": 30,
            "cached_tokens": 20,
            "total_tokens": 150,
        },
        latency_seconds=0.25,
    )

    module = LunaVlmStructureDetectorModule(
        vision_client=vision_client,
        catalog=catalog,
        renderer=ExcelSheetRenderer(),
        artifact_dir=tmp_path / "luna",
    )

    input_data = {
        "file_name": workbook_path.name,
        "workbook_hash": catalog.sha256(workbook_path),
        "sheet_names": ["Sheet1"],
    }

    result = module.run(
        input_data,
        LunaVlmStructureDetectorConfigDTO(validation_retries=0, max_concurrency=1),
    )
    validated = SpreadsheetStructureOutput.model_validate(result)

    assert len(validated.tables) == 1
    assert validated.tables[0].sheet_name == "Sheet1"
    assert module.last_usage == {
        "prompt_tokens": 120,
        "completion_tokens": 30,
        "cached_tokens": 20,
        "total_tokens": 150,
    }
    assert module.last_duration_seconds == 0.25


def test_table_validation_separates_overlapping_semantic_regions() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["F10"] = 1
    layout = compute_sheet_layout(worksheet, max_rows=10, max_columns=6)
    visibility = WorksheetVisibility.from_worksheet(worksheet)
    decision = VlmTableDecisionDTO(
        excel_range="A1:F10",
        title_range="A1:F2",
        column_header_range="A2:F4",
        row_header_range="A4:B10",
        data_range="B4:F10",
    )

    ranges = LunaVlmStructureDetectorModule._validate_table(
        decision,
        layout,
        visibility,
    )

    assert ranges["title_range"] is not None
    assert ranges["title_range"].excel_range == "A1:F1"
    assert ranges["column_header_range"] is not None
    assert ranges["column_header_range"].excel_range == "B2:F4"
    assert ranges["row_header_range"] is not None
    assert ranges["row_header_range"].excel_range == "A5:A10"
    assert ranges["data_range"] is not None
    assert ranges["data_range"].excel_range == "B5:F10"
    workbook.close()
