"""Regression contracts for refactored feature-level policy helpers."""

from __future__ import annotations

from datetime import UTC, date, datetime
from io import BytesIO

from openpyxl import Workbook

from backend.domains.bi.application.snapshot_builder import BiSnapshotBuilder
from backend.domains.bi.domain.catalog import METRIC_CATALOG, SourceMetricDefinition
from backend.domains.bi.domain.extraction_models import BiMetricExtractionResult
from backend.domains.bi.domain.materialization_models import (
    BiDocumentProfile,
    BiSnapshotBuildInput,
)
from backend.domains.bi.domain.models import (
    AmountScale,
    BiMaterializationRequest,
    BiMaterializationSource,
    BiPeriod,
    CompanyId,
    IndexId,
    JobId,
    MetricStatus,
    PeriodId,
    PeriodKind,
    SnapshotStatus,
    UnavailableObservation,
)
from backend.domains.bi.infrastructure.integrations.fast_rag_adapter import FastRagPipelineAdapter
from backend.domains.chatbot.infrastructure.filesystem import extract_text


def test_ranked_cell_aliases_cover_financial_sheet_abbreviations() -> None:
    keys = FastRagPipelineAdapter._evidence_keys(
        {
            "cell_id": "Income_Statement Cell E16",
            "sheet_name": "Income_Statement",
            "cell_coord": "E16",
        }
    )

    assert "Income_Statement Cell E16" in keys
    assert "Income_Statement:E16" in keys
    assert "IS Cell E16" in keys


def test_workbook_text_extraction_skips_hidden_vendor_sheets() -> None:
    workbook = Workbook()
    visible = workbook.active
    assert visible is not None
    visible.title = "Income_Statement"
    visible.append(["Revenue", 120])
    hidden = workbook.create_sheet("VendorPayload")
    hidden.sheet_state = "veryHidden"
    hidden.append(["binary-like-payload"])
    content = BytesIO()
    workbook.save(content)
    workbook.close()

    extracted = extract_text("financials.xlsx", content.getvalue())

    assert "Revenue | 120" in extracted
    assert "binary-like-payload" not in extracted


def test_workbook_text_extraction_keeps_rows_after_previous_per_sheet_limit() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Key_Stats"
    for index in range(500):
        worksheet.append([f"Filler metric {index}", "x" * 20])
    worksheet.append(["Total Enterprise Value (TEV)", 35_532])
    content = BytesIO()
    workbook.save(content)
    workbook.close()

    extracted = extract_text("financials.xlsx", content.getvalue())

    assert len(extracted) > 6_000
    assert "Total Enterprise Value (TEV) | 35532" in extracted


def test_snapshot_builder_projects_source_failures_into_a_partial_snapshot() -> None:
    period_id = PeriodId("fy-2025")
    period = BiPeriod(
        period_id=period_id,
        kind=PeriodKind.FY,
        label="FY2025",
        source_label="FY2025",
        end_date=date(2025, 12, 31),
        ordinal=1,
    )
    extracted = tuple(
        BiMetricExtractionResult(
            metric_id=metric_id,
            period_id=period_id,
            value_kind=definition.value_kind,
            currency="USD",
            scale=AmountScale.MILLIONS,
            observation=UnavailableObservation(
                period_id=period_id,
                status=MetricStatus.MISSING,
                reason="source_missing",
            ),
        )
        for metric_id, definition in METRIC_CATALOG.items()
        if isinstance(definition, SourceMetricDefinition)
    )
    source = BiMaterializationSource(
        file_name="sample.xlsx",
        workbook_hash="a" * 64,
        index_id=IndexId("index-sample"),
    )
    snapshot = BiSnapshotBuilder().build(
        BiSnapshotBuildInput(
            request=BiMaterializationRequest(
                company_id=CompanyId("sample"),
                display_name="Sample",
                source=source,
            ),
            job_id=JobId("job-sample"),
            profile=BiDocumentProfile(
                periods=(period,),
                currency="USD",
                scale=AmountScale.MILLIONS,
            ),
            extracted=extracted,
            generated_at=datetime(2026, 1, 1, tzinfo=UTC),
        )
    )

    assert snapshot.snapshot.status is SnapshotStatus.PARTIAL
    assert snapshot.periods == (period,)
    assert snapshot.issues
