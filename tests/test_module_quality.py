from __future__ import annotations

from pathlib import Path
from typing import Any, cast
from unittest.mock import MagicMock

import openpyxl
import pytest

from backend.engine.runtime.registry import ModuleRegistry
from backend.storage.answer_cache import AnswerCacheRepository
from backend.storage.connection_pool import PooledConnectionWrapper
from backend.storage.embedding_artifacts import EmbeddingArtifactStore
from backend.storage.pgvector_store import PgVectorStore, PgVectorStoreError
from modules.common.base_embedder import get_expected_dimension
from modules.common.base_module import DocumentContextDTO, QueryContextDTO
from modules.common.exceptions import ModuleExecutionError, StorageError
from modules.embedding.cell_text_embedder import (
    CellTextEmbedderConfigDTO,
    CellTextEmbedderInputDTO,
    CellTextEmbedderModule,
)
from modules.embedding.query_embedder import EmbeddingsDTO
from modules.query.decomposer import SubqueriesDTO
from modules.query.llm_query_router import (
    LlmQueryRouterInputDTO,
    LlmQueryRouterModule,
)
from modules.reader.reader import LookupCellMetadataInput
from modules.retrieval.context_expander import PgContextExpanderModule
from modules.retrieval.pgvector_retriever import (
    PgVectorRetrieverConfigDTO,
    PgVectorRetrieverInputDTO,
    PgVectorRetrieverModule,
    RankedSearchCandidateDTO,
    RankedSearchResultDTO,
)
from modules.retrieval.postgres_native_keyword_retriever import (
    PostgresNativeKeywordRetrieverConfigDTO,
    PostgresNativeKeywordRetrieverModule,
)
from modules.retrieval.rrf_fusion import RrfFusionInputDTO, RrfFusionModule
from modules.storage.company_entity_extractor import CompanyEntityExtractorModule
from modules.storage.pgvector_collection_loader import (
    IndexOutputDTO,
    PgVectorCollectionLoaderModule,
)
from modules.storage.pgvector_index_writer import PgVectorIndexWriterModule, VectorIndexDTO
from modules.storage.processed_file_selector import ProcessedFileSelectorModule
from modules.storage.qa_example_loader import QaExampleLoaderModule
from modules.storage.sheet_metadata_persistence import SheetMetadataPersistenceModule
from modules.structure.cell_text_serializer import (
    CellTextSerializerConfigDTO,
    CellTextSerializerModule,
)
from modules.structure.luna_vlm_structure_detector import (
    LunaVlmStructureDetectorConfigDTO,
    LunaVlmStructureDetectorModule,
    SpreadsheetStructureOutput,
    VlmTableDecisionDTO,
    unify_sheet_tables,
)


def _write_workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "IS"
    sheet["B1"] = "FY2024"
    sheet["C1"] = "FY2025"
    sheet["A2"] = "Revenue"
    sheet["B2"] = 100
    sheet["C2"] = 120
    workbook.save(path)
    workbook.close()


def _structure(file_name: str, workbook_hash: str) -> SpreadsheetStructureOutput:
    return SpreadsheetStructureOutput.model_validate(
        {
            "file_name": file_name,
            "workbook_hash": workbook_hash,
            "sheet_names": ["IS"],
            "tables": [
                {
                    "sheet_name": "IS",
                    "table_index": 1,
                    "excel_range": "A1:C2",
                    "regions": [
                        {
                            "region_id": "table_1_column_header",
                            "type": "column_header",
                            "excel_range": "B1:C1",
                            "bbox_px": [0, 0, 1, 1],
                            "rows": [1, 1],
                            "columns": [2, 3],
                            "parent_ids": [],
                        },
                        {
                            "region_id": "table_1_row_header",
                            "type": "row_header",
                            "excel_range": "A2:A2",
                            "bbox_px": [0, 0, 1, 1],
                            "rows": [2, 2],
                            "columns": [1, 1],
                            "parent_ids": ["table_1_column_header"],
                        },
                        {
                            "region_id": "table_1_data",
                            "type": "data",
                            "excel_range": "B2:C2",
                            "bbox_px": [0, 0, 1, 1],
                            "rows": [2, 2],
                            "columns": [2, 3],
                            "parent_ids": [
                                "table_1_column_header",
                                "table_1_row_header",
                            ],
                        },
                    ],
                    "header_tree": [],
                }
            ],
        }
    )


def _serialized_items() -> list[dict[str, Any]]:
    return [
        {
            "cell_id": "IS Cell B2",
            "sheet_name": "IS",
            "cell_coord": "B2",
            "row_header": ["Revenue"],
            "column_header": ["FY2024"],
            "cell_value": "100",
            "variant": "header_with_value",
            "text": "Company: ? | Sheet: IS | Row Header: Revenue | "
            "Column Header: FY2024 | Cell Value: 100",
        },
        {
            "cell_id": "IS Cell C2",
            "sheet_name": "IS",
            "cell_coord": "C2",
            "row_header": ["Revenue"],
            "column_header": ["FY2025"],
            "cell_value": "120",
            "variant": "header_with_value",
            "text": "Company: ? | Sheet: IS | Row Header: Revenue | "
            "Column Header: FY2025 | Cell Value: 120",
        },
    ]


def test_embedding_dimensions_are_model_specific() -> None:
    assert get_expected_dimension("text-embedding-3-large") == 3072
    assert get_expected_dimension("text-embedding-3-small") == 1536
    assert get_expected_dimension("BAAI/bge-base-en-v1.5") == 768


def test_registry_reuses_model_clients_and_storage(tmp_path: Path) -> None:
    completion_client = MagicMock()
    pgvector_store = MagicMock()
    db_manager = MagicMock()
    registry = ModuleRegistry(
        repository=AnswerCacheRepository(tmp_path / "answer-cache.sqlite"),
        completion_client=completion_client,
        embedding_artifact_store=EmbeddingArtifactStore(tmp_path / "artifacts"),
        pgvector_store=pgvector_store,
        db_manager=db_manager,
        processed_dir=tmp_path,
        spreadsheet_artifact_dir=tmp_path / "sheets",
    )

    for module_name in ("decomposer", "llm_query_router", "reader", "company_entity_extractor"):
        assert cast(Any, registry.get(module_name)).completion_client is completion_client
    assert cast(Any, registry.get("reader")).pgvector_store is pgvector_store


def test_cell_embedder_streams_and_reuses_artifact(tmp_path: Path) -> None:
    encoder = MagicMock()
    encoder.encode.side_effect = [
        [[0.1, 0.2], [0.3, 0.4]],
        [[0.5, 0.6]],
    ]
    artifact_store = EmbeddingArtifactStore(tmp_path)
    module = CellTextEmbedderModule(encoder=encoder, artifact_store=artifact_store)
    input_data = CellTextEmbedderInputDTO.model_validate(
        {
            "file_name": "sample.xlsx",
            "workbook_hash": "w" * 64,
            "items": [*_serialized_items(), {**_serialized_items()[0], "cell_id": "IS Cell D2", "cell_coord": "D2"}],
        }
    )
    config = CellTextEmbedderConfigDTO(model="custom-2d", dimension=2, batch_size=2)

    first = module.run(input_data, config)
    second = module.run(input_data, config)

    assert first["cache_hit"] is False
    assert second["cache_hit"] is True
    assert encoder.encode.call_count == 2
    assert max(len(call.args[0]) for call in encoder.encode.call_args_list) <= 2
    assert artifact_store.is_valid(first["artifact_id"], 3, 2)


def test_pgvector_writer_uses_lazy_artifact_batches(tmp_path: Path) -> None:
    artifact_store = EmbeddingArtifactStore(tmp_path / "vectors")
    artifact_id = "a" * 64
    artifact_store.put(artifact_id, [[0.1, 0.2], [0.3, 0.4]])
    db_manager = MagicMock()
    pgvector_store = MagicMock()
    module = PgVectorIndexWriterModule(
        artifact_store=artifact_store,
        db_manager=db_manager,
        pgvector_store=pgvector_store,
        processed_dir=tmp_path,
    )

    result = module.run(
        {
            "file_name": "sample.xlsx",
            "workbook_hash": "b" * 64,
            "model": "custom-2d",
            "artifact_id": artifact_id,
            "dimension": 2,
            "items": [
                {**item, "embedding_index": index}
                for index, item in enumerate(_serialized_items())
            ],
        }
    )

    assert result["index_id"] == f"idx_{artifact_id}"
    vectors = pgvector_store.put_documents.call_args.kwargs["vectors"]
    assert vectors[0] == pytest.approx([0.1, 0.2])
    assert len(vectors) == 2
    db_manager.save_source_file.assert_called_once()


def test_pgvector_store_publishes_staging_collection_atomically(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from langchain_core.documents import Document

    import backend.storage.pgvector_store as store_module

    vector_backend = MagicMock()
    factory = MagicMock(return_value=vector_backend)
    monkeypatch.setattr(store_module, "get_vector_store", factory)

    connection = MagicMock()
    cursor = MagicMock()
    cursor.rowcount = 0
    cursor.fetchone.side_effect = [None, ("new-uuid",)]
    connection.cursor.return_value.__enter__.return_value = cursor
    store = PgVectorStore("postgresql://unused")
    cast(Any, store)._raw_connection = MagicMock(return_value=connection)
    cast(Any, store).ensure_optimized_indexes = MagicMock()
    cast(Any, store).ensure_collection_vector_index = MagicMock()

    store.put_documents(
        "idx_target",
        [Document(page_content="value", metadata={"cell_id": "IS:B2"})],
        model_name="custom-2d",
        vectors=[[0.1, 0.2]],
        metadata={"dimension": 2},
    )

    staging_name = factory.call_args.kwargs["collection_name"]
    assert staging_name.startswith("idx_target__staging__")
    vector_backend.add_embeddings.assert_called_once()
    executed_sql = "\n".join(call.args[0] for call in cursor.execute.call_args_list)
    assert "pg_advisory_xact_lock" in executed_sql
    assert "SET name = %s, cmetadata = %s" in executed_sql
    connection.commit.assert_called_once()


def test_collection_local_hnsw_index_is_quantized_partial_and_concurrent() -> None:
    store = PgVectorStore("postgresql://unused")
    collection_uuid = "11111111-1111-1111-1111-111111111111"
    cast(Any, store)._collection_uuid = MagicMock(return_value=collection_uuid)
    connection = MagicMock()
    cursor = MagicMock()
    connection.cursor.return_value.__enter__.return_value = cursor
    cast(Any, store)._read_connection = MagicMock(return_value=connection)

    index_name = store.ensure_collection_vector_index("idx_target", 3072)

    sql = cursor.execute.call_args.args[0]
    assert index_name.endswith("_3072")
    assert "CREATE INDEX CONCURRENTLY" in sql
    assert f"collection_id = '{collection_uuid}'::uuid" in sql
    assert "binary_quantize(embedding)::bit(3072)" in sql
    assert "bit_hamming_ops" in sql


def test_pgvector_store_delete_cleans_collection_local_index() -> None:
    connection = MagicMock()
    cursor = MagicMock()
    cursor.fetchone.return_value = (
        "11111111-1111-1111-1111-111111111111",
        {"dimension": 3072},
    )
    connection.cursor.return_value.__enter__.return_value = cursor
    store = PgVectorStore("postgresql://unused")
    cast(Any, store)._raw_connection = MagicMock(return_value=connection)
    cast(Any, store)._drop_collection_vector_index = MagicMock()

    assert store.delete("idx_target") is True
    cast(Any, store)._drop_collection_vector_index.assert_called_once_with(
        "11111111-1111-1111-1111-111111111111",
        3072,
    )


def test_pgvector_store_failed_stage_does_not_publish(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from langchain_core.documents import Document

    import backend.storage.pgvector_store as store_module

    vector_backend = MagicMock()
    vector_backend.add_embeddings.side_effect = RuntimeError("batch failed")
    monkeypatch.setattr(store_module, "get_vector_store", MagicMock(return_value=vector_backend))
    store = PgVectorStore("postgresql://unused")
    cast(Any, store)._raw_connection = MagicMock()

    with pytest.raises(PgVectorStoreError, match="배치 적재 실패"):
        store.put_documents(
            "idx_target",
            [Document(page_content="value", metadata={})],
            vectors=[[0.1, 0.2]],
        )
    vector_backend.delete_collection.assert_called_once()
    cast(Any, store)._raw_connection.assert_not_called()


def test_pooled_connection_rolls_back_before_reuse() -> None:
    pool = MagicMock()
    connection = MagicMock()
    connection.get_transaction_status.return_value = 2
    connection.autocommit = False
    wrapper = PooledConnectionWrapper(pool, connection)
    wrapper.close()
    connection.rollback.assert_called_once()
    pool.putconn.assert_called_once_with(connection)


def test_collection_loader_flattens_metadata_and_rejects_mixed_models() -> None:
    store = MagicMock()
    store.list_indexes.return_value = [
        {
            "index_id": "idx_a",
            "file_name": "a.xlsx",
            "workbook_hash": "hash-a",
            "model": "model-a",
            "dimension": 2,
            "document_count": 4,
        }
    ]
    module = PgVectorCollectionLoaderModule(store, MagicMock())
    result = module.run({"collection_name": "a.xlsx"})
    assert result["index_output"]["index_id"] == "idx_a"

    store.list_indexes.return_value.append(
        {
            "index_id": "idx_b",
            "file_name": "b.xlsx",
            "workbook_hash": "hash-b",
            "model": "model-b",
            "dimension": 3,
            "document_count": 5,
        }
    )
    with pytest.raises(ModuleExecutionError, match="서로 다른 임베딩"):
        module.run({"collection_names": ["idx_a", "idx_b"]})


def test_dense_retrieval_preserves_top_k_per_subquery() -> None:
    from langchain_core.documents import Document

    store = MagicMock()
    store.similarity_search_by_vector_with_score.side_effect = [
        [(Document(page_content="one", metadata={"cell_id": "A:IS:B2"}), 0.1)],
        [(Document(page_content="two", metadata={"cell_id": "A:IS:C2"}), 0.2)],
    ]
    module = PgVectorRetrieverModule(store)
    result = module.run(
        PgVectorRetrieverInputDTO(
            query_input=EmbeddingsDTO(
                query_context=QueryContextDTO(question_id="q", question_text="question"),
                items={"query one": [0.1, 0.2], "query two": [0.3, 0.4]},
            ),
            index_input=IndexOutputDTO(
                index_id="idx_a",
                file_name="a.xlsx",
                workbook_hash="hash-a",
                model="custom-2d",
                dimension=2,
                document_count=2,
            ),
        ),
        PgVectorRetrieverConfigDTO(top_k=1),
    )
    assert len(result["items"]) == 2
    assert [item["rank"] for item in result["items"]] == [1, 1]
    assert result["document_context"]["index_id"] == "idx_a"


def test_keyword_retrieval_resets_rank_for_each_subquery() -> None:
    store = MagicMock()
    connection = MagicMock()
    cursor = MagicMock()
    cursor.fetchall.side_effect = [
        [("id-1", "one", {"cell_id": "A:IS:B2"}, 0.9, "uuid-a")],
        [("id-2", "two", {"cell_id": "A:IS:C2"}, 0.8, "uuid-a")],
    ]
    connection.cursor.return_value.__enter__.return_value = cursor
    store._read_connection.return_value = connection
    module = PostgresNativeKeywordRetrieverModule(store)
    result = module.run(
        {
            "query_input": SubqueriesDTO(
                query_context=QueryContextDTO(question_id="q", question_text="question"),
                subqueries=["Revenue FY2024", "Revenue FY2025"],
            ),
            "index_input": IndexOutputDTO(
                index_id="idx_a",
                file_name="a.xlsx",
                workbook_hash="hash-a",
                model="custom-2d",
                dimension=2,
                document_count=2,
            ),
        },
        PostgresNativeKeywordRetrieverConfigDTO(top_k=1),
    )
    assert [item["rank"] for item in result["items"]] == [1, 1]
    assert [item["matched_subquery"] for item in result["items"]] == [
        "Revenue FY2024",
        "Revenue FY2025",
    ]


def test_rrf_rejects_cross_document_fusion() -> None:
    query = QueryContextDTO(question_id="q", question_text="question")
    candidate = RankedSearchCandidateDTO(
        rank=1,
        cell_id="A:IS:B2",
        score=1,
        text="value",
        matched_subquery="query",
    )
    first = RankedSearchResultDTO(
        query_context=query,
        document_context=DocumentContextDTO(
            file_name="a.xlsx", workbook_hash="hash-a", index_id="idx_a"
        ),
        items=[candidate],
    )
    second = RankedSearchResultDTO(
        query_context=query,
        document_context=DocumentContextDTO(
            file_name="b.xlsx", workbook_hash="hash-b", index_id="idx_b"
        ),
        items=[candidate],
    )
    with pytest.raises(ModuleExecutionError, match="문서 또는 인덱스"):
        RrfFusionModule().run(RrfFusionInputDTO(bm25_result=first, dense_result=second))


def test_context_expander_batches_rows_per_sheet() -> None:
    store = MagicMock()
    store.fetch_rows_cells.return_value = {
        2: [
            {
                "col_index": 2,
                "column_header": ["FY2024"],
                "cell_value": "100",
                "cell_coord": "B2",
                "row_header": ["Revenue"],
            }
        ]
    }
    module = PgContextExpanderModule(store)
    retrieval = {
        "query_context": {"question_id": "q", "question_text": "Revenue?"},
        "document_context": {
            "file_name": "a.xlsx",
            "workbook_hash": "hash-a",
            "index_id": "idx_a",
        },
        "items": [
            {
                "rank": 1,
                "cell_id": "Company:IS:B2",
                "rrf_score": 0.1,
                "text": "candidate",
                "matched_subquery": "Revenue",
            }
        ],
    }
    result = module.run(
        {"retrieval_json": retrieval},
        {"top_k": 1, "adjacent_radius": 0, "max_blocks": 10},
    )
    assert store.fetch_rows_cells.call_count == 1
    assert "FY2024" in result["context_blocks"][1]


def test_router_preserves_low_model_confidence() -> None:
    from backend.providers.llm.chat_completion import ChatCompletionResult
    from backend.semantic_matching.catalog import QueryExample

    client = MagicMock()
    client.complete_with_metadata.return_value = ChatCompletionResult(
        content='{"target":"IS","confidence":0.01,"company_name":null,'
        '"company_scopes":[],"reason":"weak"}',
        usage={},
        latency_seconds=0,
    )
    module = LlmQueryRouterModule(
        client,
        examples=[
            QueryExample(
                example_id="1",
                question="Revenue?",
                target="IS",
                sheets=("IS",),
                query_type=1,
            )
        ],
    )
    result = module.run(
        LlmQueryRouterInputDTO(
            query_context=QueryContextDTO(question_id="q", question_text="maybe revenue")
        )
    )
    assert result["semantic_match"]["confidence"] == 0.01


def test_processed_selector_and_serializer_honor_config(tmp_path: Path) -> None:
    workbook_path = tmp_path / "sample.xlsx"
    _write_workbook(workbook_path)
    selected = ProcessedFileSelectorModule(processed_dir=tmp_path).run(
        {"file_name": workbook_path.name}
    )
    assert selected["sheet_names"] == ["IS"]

    structure = _structure(workbook_path.name, selected["workbook_hash"])
    serialized = CellTextSerializerModule(processed_dir=tmp_path).run(
        structure,
        CellTextSerializerConfigDTO(variant_mode="header_only"),
    )
    assert len(serialized["items"]) == 2
    assert {item["variant"] for item in serialized["items"]} == {"header_only"}
    assert all(item["text"].startswith("Company: ? | Sheet: IS") for item in serialized["items"])
    assert all(item["text"].endswith("Cell Value: ?") for item in serialized["items"])


def test_luna_detector_executes_on_tiny_workbook(tmp_path: Path) -> None:
    from backend.storage.spreadsheets.workbook_catalog import WorkbookCatalog

    workbook_path = tmp_path / "sample.xlsx"
    _write_workbook(workbook_path)
    catalog = WorkbookCatalog(tmp_path)

    vision_client = MagicMock()
    vision_client.complete_structured.return_value = (
        '{"tables":[{"excel_range":"A1:C2","title_range":null,'
        '"column_header_range":"B1:C1","row_header_range":"A2:A2",'
        '"data_range":"B2:C2"}]}'
    )
    module = LunaVlmStructureDetectorModule(
        vision_client=vision_client,
        catalog=catalog,
        processed_dir=tmp_path,
        artifact_dir=tmp_path / "luna",
    )
    result = module.run(
        {
            "file_name": workbook_path.name,
            "workbook_hash": catalog.sha256(workbook_path),
            "sheet_names": ["IS"],
        },
        LunaVlmStructureDetectorConfigDTO(validation_retries=0, max_concurrency=1),
    )
    assert len(result["tables"]) == 1
    assert result["tables"][0]["sheet_name"] == "IS"


def test_luna_does_not_merge_independent_tables_with_own_headers() -> None:
    tables = [
        VlmTableDecisionDTO(
            excel_range="A1:C3",
            column_header_range="B1:C1",
            row_header_range="A2:A3",
            data_range="B2:C3",
        ),
        VlmTableDecisionDTO(
            excel_range="A5:C7",
            column_header_range="B5:C5",
            row_header_range="A6:A7",
            data_range="B6:C7",
        ),
    ]
    assert unify_sheet_tables(tables) == tables


def test_sheet_metadata_persistence_and_storage_failure(tmp_path: Path) -> None:
    workbook_path = tmp_path / "sample.xlsx"
    _write_workbook(workbook_path)
    from backend.storage.spreadsheets.workbook_catalog import WorkbookCatalog

    catalog = WorkbookCatalog(tmp_path)
    workbook_hash = catalog.sha256(workbook_path)
    structure = _structure(workbook_path.name, workbook_hash)
    index = VectorIndexDTO(
        index_id="idx_" + "a" * 64,
        file_name=workbook_path.name,
        workbook_hash=workbook_hash,
        model="custom-2d",
        dimension=2,
        document_count=2,
    )
    database = MagicMock()
    database.is_connected.return_value = True
    module = SheetMetadataPersistenceModule(db_manager=database, catalog=catalog)
    result = module.run({"structure_input": structure, "index_input": index})
    assert result["sheets_saved"] == 1
    database.save_sheets.assert_called_once()

    database.save_sheets.side_effect = RuntimeError("database unavailable")
    with pytest.raises(StorageError):
        module.run({"structure_input": structure, "index_input": index})


def test_qa_loader_builtin_and_path_traversal() -> None:
    module = QaExampleLoaderModule()
    result = module.run({})
    assert result["total_count"] > 0
    with pytest.raises(ModuleExecutionError, match="JSON 파일명"):
        module.run({"file_name": "../secret.json"})


def test_company_metadata_persistence_failure_is_not_silenced() -> None:
    module = CompanyEntityExtractorModule(
        catalog=MagicMock(),
        completion_client=MagicMock(),
        pgvector_store=MagicMock(),
    )
    module._sample_workbook = MagicMock(return_value=[])
    cast(Any, module.pgvector_store).update_index_company.side_effect = RuntimeError("db down")
    with pytest.raises(StorageError):
        module.run({"file_name": "acme.xlsx", "index_id": "idx_a"})


def test_reader_tool_input_bounds() -> None:
    with pytest.raises(ValueError):
        LookupCellMetadataInput(cell_coords=[])
    with pytest.raises(ValueError):
        LookupCellMetadataInput(cell_coords=["../../etc/passwd"])
