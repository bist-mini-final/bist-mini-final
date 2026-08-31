from __future__ import annotations

from datetime import datetime
from hashlib import sha256
from pathlib import Path

import openpyxl

from backend.domains.data_sources.domain.workbook_profiles import (
    WorkbookAmountScale,
    WorkbookSheetRole,
)
from backend.domains.data_sources.infrastructure.spreadsheets.workbook_profile_extractor import (
    WorkbookProfileExtractor,
)
from modules.structure.luna_vlm_structure_detector import (
    ClassifiedRegionDTO,
    ClassifiedTableDTO,
    SpreadsheetStructureOutput,
)


def _workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    income = workbook.active
    assert income is not None
    income.title = "Operations A"
    income["A1"] = "Statement of Operations"
    income["A2"] = "Data in ($M)"
    income["A4"] = "Period Ended"
    for column, year in enumerate(range(2021, 2026), start=2):
        income.cell(4, column, datetime(year, 12, 31))
    income["G3"] = "LTM"
    income["G4"] = datetime(2025, 12, 31)
    income["H3"] = "Estimates"
    income["H4"] = datetime(2026, 12, 31)
    income["A5"] = "Total Revenue"
    income["B5"] = 100
    income["A8"] = "Valuation date"
    for column in range(2, 8):
        income.cell(8, column, datetime(2026, 2, 6))

    balance = workbook.create_sheet("Position 9")
    balance["A1"] = "Statement of Financial Position"
    balance["A2"] = "Total Assets"
    balance["A3"] = "Total Liabilities"

    cash = workbook.create_sheet("Movement Z")
    cash["A1"] = "Cash Flow"
    cash["A2"] = "Operating Cash Flow"
    cash["A3"] = "Capital Expenditure"
    workbook.save(path)


def _structure(path: Path) -> SpreadsheetStructureOutput:
    digest = sha256(path.read_bytes()).hexdigest()
    return SpreadsheetStructureOutput(
        file_name=path.name,
        workbook_hash=digest,
        sheet_names=["Operations A", "Position 9", "Movement Z"],
        tables=[
            ClassifiedTableDTO(
                sheet_name="Operations A",
                table_index=0,
                excel_range="A4:H5",
                regions=[
                    ClassifiedRegionDTO(
                        region_id="header",
                        type="column_header",
                        excel_range="B4:H4",
                        bbox_px=(0, 0, 1, 1),
                        rows=(4, 4),
                        columns=(2, 8),
                        parent_ids=[],
                    )
                ],
            )
        ],
    )


def test_extracts_common_profile_without_fixed_sheet_names_or_cells(tmp_path: Path) -> None:
    path = tmp_path / "arbitrary-layout.xlsx"
    _workbook(path)
    structure = _structure(path)

    profile = WorkbookProfileExtractor().extract(
        workbook_path=path,
        structure=structure,
        index_id="idx_" + structure.workbook_hash,
    )

    assert profile.status == "ready"
    assert profile.currency == "USD"
    assert profile.amount_scale is WorkbookAmountScale.MILLIONS
    assert [(period.label, period.kind.value) for period in profile.periods] == [
        ("FY2021", "fy"),
        ("FY2022", "fy"),
        ("FY2023", "fy"),
        ("FY2024", "fy"),
        ("FY2025", "fy"),
        ("LTM 2025", "ltm"),
    ]
    roles = {sheet.sheet_name: sheet.role for sheet in profile.sheets}
    assert roles == {
        "Operations A": WorkbookSheetRole.INCOME_STATEMENT,
        "Position 9": WorkbookSheetRole.BALANCE_SHEET,
        "Movement Z": WorkbookSheetRole.CASH_FLOW,
    }
    assert {item.cell_coord for item in profile.evidence if item.kind == "scale"} >= {"A2"}


def test_profile_is_partial_when_units_are_not_explicit(tmp_path: Path) -> None:
    path = tmp_path / "unitless.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Metrics"
    sheet["A1"] = "Period Ended"
    sheet["B1"] = datetime(2024, 12, 31)
    sheet["C1"] = datetime(2025, 12, 31)
    workbook.save(path)
    structure = _structure_for_unitless(path)

    profile = WorkbookProfileExtractor().extract(
        workbook_path=path,
        structure=structure,
        index_id="idx_" + structure.workbook_hash,
    )

    assert profile.status == "partial"
    assert profile.currency is None
    assert profile.amount_scale is None
    assert profile.diagnostics == ("currency_missing", "amount_scale_missing")


def _structure_for_unitless(path: Path) -> SpreadsheetStructureOutput:
    digest = sha256(path.read_bytes()).hexdigest()
    return SpreadsheetStructureOutput(
        file_name=path.name,
        workbook_hash=digest,
        sheet_names=["Metrics"],
        tables=[],
    )
