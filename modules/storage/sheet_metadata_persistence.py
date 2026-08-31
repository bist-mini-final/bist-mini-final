"""pgvector 벡터 인덱싱 완료 후 엑셀 시트 메타데이터 및 감지된 표 구조를 PostgreSQL에 저장하는 모듈.

엑셀 파일의 실제 행/열 차원(dimensions)과 VLM/파서로 감지된 테이블 바운딩 박스를
데이터베이스의 `sheet_metadata` 테이블에 영속화합니다.

Example:
    Input DTO (입력 예시):
    ```json
    {
      "structure_input": {
        "file_name": "samsung_2023.xlsx",
        "workbook_hash": "a1b2c3d4...",
        "sheet_names": ["손익계산서", "재무상태표"]
      },
      "index_input": {
        "collection_name": "rag_cells_a1b2c3d4",
        "workbook_hash": "a1b2c3d4...",
        "vector_count": 1500
      }
    }
    ```

    Output DTO (출력 예시):
    ```json
    {
      "sheets_saved": 2,
      "sheet_details": [
        {
          "sheet_name": "손익계산서",
          "row_count": 85,
          "column_count": 12,
          "table_count": 1
        },
        {
          "sheet_name": "재무상태표",
          "row_count": 110,
          "column_count": 14,
          "table_count": 1
        }
      ]
    }
    ```
"""

from __future__ import annotations

# ==============================================================================
# 1. Imports & Logger Setup
# ==============================================================================
import logging
from typing import Any, ClassVar, Dict, List, Optional, Union, cast

from pydantic import Field

from backend.domains.data_sources.infrastructure.spreadsheets.workbook_catalog import (
    WorkbookCatalog,
)
from modules.common.base_module import (
    BaseModule,
    EmptyModuleConfigDTO,
    ModuleDefinition,
    ModuleDTO,
    ModuleExecutionError,
)
from modules.common.exceptions import DocumentParsingError, StorageError
from modules.storage.pgvector_index_writer import VectorIndexDTO
from modules.storage.ports import SourceFileRepositoryPort
from modules.storage.processed_file_selector import WorkbookSelectionDTO
from modules.structure.luna_vlm_structure_detector import SpreadsheetStructureOutput

logger = logging.getLogger(__name__)
StructureSourceDTO = Union[SpreadsheetStructureOutput, WorkbookSelectionDTO]


# ==============================================================================
# 2. DTOs & Item Models
# ==============================================================================
class SheetMetadataPersistenceInputDTO(ModuleDTO):
    """Input contract containing workbook structure and pgvector index outcome."""

    structure_input: StructureSourceDTO = Field(
        description="구조 감지 결과 또는 전수 직렬화용 워크북 선택 결과",
    )
    index_input: VectorIndexDTO = Field(
        description="source_files 저장이 완료된 pgvector 인덱스 결과",
    )


class SheetMetadataPersistenceOutputDTO(ModuleDTO):
    """Output contract containing saved sheet metadata summary."""

    sheets_saved: int = Field(description="DB에 저장된 시트 수")
    sheet_details: List[Dict[str, Any]] = Field(default_factory=list)


# ==============================================================================
# 3. Module Definition & Implementation
# ==============================================================================
_SHEET_META_DEFINITION = ModuleDefinition(
    type="sheet_metadata_persistence",
    label="Sheet Metadata Persistence",
    category="Storage / DB",
    description="pgvector 적재 후 워크북 시트 크기와 감지 테이블을 DB에 저장합니다.",
    inputs=["structure_input", "index_input"],
    outputs=["output"],
    config_fields=[],
    raw_output=True,
    cacheable=False,
    version="2",
)


class SheetMetadataPersistenceModule(BaseModule):
    """Persist sheet dimensions and detected table metadata."""

    definition: ClassVar[ModuleDefinition] = _SHEET_META_DEFINITION
    input_model = SheetMetadataPersistenceInputDTO
    config_model = EmptyModuleConfigDTO
    output_model = SheetMetadataPersistenceOutputDTO

    def __init__(
        self,
        source_files: SourceFileRepositoryPort,
        catalog: WorkbookCatalog,
    ) -> None:
        """Initialize with dependencies owned by the runtime composition root."""
        super().__init__()
        self._source_files = source_files
        self.catalog = catalog

    @staticmethod
    def _detected_tables(structure: StructureSourceDTO) -> List[Any]:
        if isinstance(structure, WorkbookSelectionDTO):
            return []
        return list(structure.tables)

    def _visible_sheet_names(
        self,
        structure: StructureSourceDTO,
        workbook_path: Any,
        detected_tables: List[Any],
    ) -> List[str]:
        if structure.sheet_names:
            return list(structure.sheet_names)
        if detected_tables:
            table_sheets = {table.sheet_name for table in detected_tables}
            return [
                sheet_name
                for sheet_name in self.catalog.sheet_names(workbook_path)
                if sheet_name in table_sheets
            ]
        raise ModuleExecutionError(
            "저장할 시트 목록(sheet_names) 또는 감지된 테이블(tables)이 지정되지 않았습니다"
        )

    @staticmethod
    def _sheet_dimensions(workbook_path: Any) -> Dict[str, tuple[int, int]]:
        try:
            import openpyxl

            workbook = openpyxl.load_workbook(
                workbook_path,
                read_only=True,
                data_only=True,
            )
            try:
                dimensions: Dict[str, tuple[int, int]] = {}
                for sheet_name in workbook.sheetnames:
                    worksheet = workbook[sheet_name]
                    if worksheet.max_row is None or worksheet.max_column is None:
                        cast(Any, worksheet).calculate_dimension(force=True)
                    dimensions[sheet_name] = (
                        worksheet.max_row or 0,
                        worksheet.max_column or 0,
                    )
                return dimensions
            finally:
                workbook.close()
        except Exception as error:
            raise DocumentParsingError(f"시트 크기 측정 실패: {error}") from error

    @staticmethod
    def _sheet_records(
        visible_sheets: List[str],
        dimensions: Dict[str, tuple[int, int]],
        detected_tables: List[Any],
    ) -> List[Dict[str, Any]]:
        records: List[Dict[str, Any]] = []
        for sheet_index, sheet_name in enumerate(visible_sheets):
            if sheet_name not in dimensions:
                logger.warning("Skipping sheet '%s' not found in workbook dimensions", sheet_name)
                continue
            rows, columns = dimensions[sheet_name]
            records.append(
                {
                    "sheet_name": sheet_name,
                    "sheet_index": sheet_index,
                    "is_visible": True,
                    "row_count": rows,
                    "column_count": columns,
                    "detected_tables": [
                        table.model_dump(mode="json")
                        for table in detected_tables
                        if table.sheet_name == sheet_name
                    ],
                }
            )
        return records

    def _persist_sheets(
        self,
        workbook_hash: str,
        sheets_data: List[Dict[str, Any]],
    ) -> None:
        try:
            self._source_files.save_sheets(
                file_id=workbook_hash,
                sheets_info=sheets_data,
            )
        except Exception as error:
            raise StorageError(f"시트 메타데이터 저장 실패: {error}") from error

    @staticmethod
    def _sheet_summary(sheets_data: List[Dict[str, Any]]) -> Dict[str, Any]:
        return {
            "sheets_saved": len(sheets_data),
            "sheet_details": [
                {
                    "sheet_name": sheet["sheet_name"],
                    "row_count": sheet["row_count"],
                    "column_count": sheet["column_count"],
                    "table_count": len(sheet["detected_tables"]),
                }
                for sheet in sheets_data
            ],
        }

    def execute(
        self,
        input_data: SheetMetadataPersistenceInputDTO,
        config: Optional[EmptyModuleConfigDTO] = None,
    ) -> Dict[str, Any]:
        """Persist workbook sheet metadata and detected tables for the selected visible sheets."""
        structure = input_data.structure_input
        index = input_data.index_input
        if structure.workbook_hash != index.workbook_hash:
            raise ModuleExecutionError("구조 분석과 인덱스의 workbook_hash가 다릅니다")

        database = self._source_files
        if not database.is_connected():
            raise StorageError("시트 메타데이터를 저장할 DB에 연결할 수 없습니다")

        workbook_path = self.catalog.resolve(structure.file_name)
        detected_tables = self._detected_tables(structure)
        visible_sheets = self._visible_sheet_names(
            structure,
            workbook_path,
            detected_tables,
        )
        dimensions = self._sheet_dimensions(workbook_path)
        sheets_data = self._sheet_records(visible_sheets, dimensions, detected_tables)
        self._persist_sheets(structure.workbook_hash, sheets_data)
        return self._sheet_summary(sheets_data)


# ==============================================================================
# 4. Exports
# ==============================================================================
__all__ = [
    "SheetMetadataPersistenceInputDTO",
    "SheetMetadataPersistenceModule",
    "SheetMetadataPersistenceOutputDTO",
]
