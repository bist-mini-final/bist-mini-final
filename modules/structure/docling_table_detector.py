from __future__ import annotations

from pathlib import Path
from typing import Optional, Any, Dict, List, Sequence, Tuple, cast

import openpyxl
from PIL import Image, ImageDraw, ImageFont
from pydantic import BaseModel, Field

from backend.core.settings import PROCESSED_DATA_DIR, SPREADSHEET_ARTIFACT_DIR
from backend.storage.spreadsheets.cell_visibility import worksheet_visible
from backend.storage.spreadsheets.docling_extractor import DoclingTableExtractor, TableExtractor
from backend.storage.spreadsheets.sheet_renderer import ExcelSheetRenderer
from backend.storage.spreadsheets.table_geometry import (
    bbox_to_cell_bounds,
    cell_bounds_bbox,
)
from backend.storage.spreadsheets.workbook_catalog import WorkbookCatalog, WorkbookCatalogError
from modules.common.base_module import BaseModule, ModuleConfigDTO, ModuleDefinition, ModuleDTO, ModuleExecutionError
from modules.storage.processed_file_selector import WorkbookSelectionDTO


class DoclingTableDetectorInputDTO(WorkbookSelectionDTO):
    """Workbook identity and visible sheet selection."""


from modules.common.config import (
    DEFAULT_STRUCTURE_MAX_COLUMNS,
    DEFAULT_STRUCTURE_MAX_ROWS,
)


class DoclingTableDetectorConfigDTO(ModuleConfigDTO):
    max_rows: int = Field(
        default=DEFAULT_STRUCTURE_MAX_ROWS,
        ge=1,
        le=2000,
        description="시트 이미지화 및 탐지에 포함할 최대 행 수",
    )
    max_columns: int = Field(
        default=DEFAULT_STRUCTURE_MAX_COLUMNS,
        ge=1,
        le=200,
        description="시트 이미지화 및 탐지에 포함할 최대 열 수",
    )


class DoclingTableDetectorExecutionDTO(
    DoclingTableDetectorInputDTO,
    DoclingTableDetectorConfigDTO,
):
    """Internal union of workbook data and render limits."""


class TableCellBoundsDTO(ModuleDTO):
    min_row: int
    max_row: int
    min_column: int
    max_column: int


class DoclingTableRegionDTO(ModuleDTO):
    sheet_name: str
    table_index: int
    excel_range: str
    bbox_px: Tuple[float, float, float, float]
    cell_bounds: TableCellBoundsDTO


class DoclingTableDetectorOutput(ModuleDTO):
    file_name: str
    workbook_hash: str
    sheet_names: List[str]
    tables: List[DoclingTableRegionDTO]


def _safe_name(value: str) -> str:
    return "".join("_" if character in '\\/*?:\"<>| ' else character for character in value).strip("_")


def _annotate_tables(
    source_path: Path,
    output_path: Path,
    detections: Sequence[DoclingTableRegionDTO],
) -> None:
    with Image.open(source_path) as source:
        image = source.convert("RGBA")
    overlay = Image.new("RGBA", image.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay)
    font = ImageFont.load_default()
    for detection in detections:
        x1, y1, x2, y2 = detection.bbox_px
        draw.rectangle((x1, y1, x2, y2), fill=(8, 145, 178, 32), outline=(8, 145, 178, 230), width=3)
        label = f"Table {detection.table_index}: {detection.excel_range}"
        text_box = draw.textbbox((x1 + 4, y1 + 4), label, font=font)
        draw.rectangle(text_box, fill=(255, 255, 255, 220))
        draw.text((x1 + 4, y1 + 4), label, fill=(14, 116, 144, 255), font=font)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    Image.alpha_composite(image, overlay).convert("RGB").save(output_path, format="PNG")


class DoclingTableDetectorModule(BaseModule):
    definition = ModuleDefinition(
        type="docling_table_detector",
        label="Docling Table Region Detector",
        category="Logic",
        description="Excel 시트를 PNG로 렌더링하고 Docling으로 테이블 경계를 추출합니다.",
        inputs=["input"],
        outputs=["output"],
        config_fields=["max_rows", "max_columns"],
        raw_output=True,
        version="4",
    )
    input_model = DoclingTableDetectorInputDTO
    config_model = DoclingTableDetectorConfigDTO
    execution_model = DoclingTableDetectorExecutionDTO
    output_model = DoclingTableDetectorOutput

    def __init__(
        self,
        catalog: WorkbookCatalog | None = None,
        renderer: ExcelSheetRenderer | None = None,
        extractor: TableExtractor | None = None,
        processed_dir: Path = PROCESSED_DATA_DIR,
        artifact_dir: Path = SPREADSHEET_ARTIFACT_DIR,
    ) -> None:
        self.catalog = catalog or WorkbookCatalog(processed_dir)
        self.renderer = renderer or ExcelSheetRenderer()
        self.extractor = extractor or DoclingTableExtractor()
        self.artifact_dir = artifact_dir

    def _normalize_bbox(
        self,
        bbox: Sequence[float],
        image_width: int,
        image_height: int,
    ) -> Tuple[float, float, float, float] | None:
        if len(bbox) != 4:
            return None
        x1, y1, x2, y2 = (float(value) for value in bbox)
        normalized = (
            max(0.0, min(x1, x2)),
            max(0.0, min(y1, y2)),
            min(float(image_width), max(x1, x2)),
            min(float(image_height), max(y1, y2)),
        )
        return normalized if normalized[2] > normalized[0] and normalized[3] > normalized[1] else None

    def execute(
        self,
        input_data: DoclingTableDetectorInputDTO,
        config: Optional[DoclingTableDetectorConfigDTO] = None,
    ) -> Dict[str, Any]:
        if config is None and isinstance(input_data, DoclingTableDetectorExecutionDTO):
            cfg = input_data
        else:
            cfg = config or DoclingTableDetectorConfigDTO()
        try:
            workbook_path = self.catalog.resolve(input_data.file_name)
            current_hash = self.catalog.sha256(workbook_path)
        except (OSError, ValueError, WorkbookCatalogError) as error:
            raise ModuleExecutionError(str(error)) from error
        if current_hash != input_data.workbook_hash:
            raise ModuleExecutionError(
                "선택 이후 Excel 파일이 변경되었습니다. 파일 선택 모듈을 다시 실행하세요"
            )

        output_root = self.artifact_dir / current_hash[:16]
        workbook = None
        try:
            workbook = openpyxl.load_workbook(
                workbook_path,
                read_only=False,
                data_only=True,
                keep_vba=workbook_path.suffix.lower() == ".xlsm",
            )
            table_outputs: List[Dict[str, Any]] = []
            for sheet_name in input_data.sheet_names:
                if sheet_name not in workbook.sheetnames:
                    raise ModuleExecutionError(f"Excel 시트를 찾을 수 없습니다: {sheet_name}")
                worksheet = workbook[sheet_name]
                if not worksheet_visible(worksheet):
                    raise ModuleExecutionError(
                        f"숨겨진 Excel 시트는 분석할 수 없습니다: {sheet_name}"
                    )
                safe_sheet = _safe_name(sheet_name)
                image_path = output_root / "rendered" / f"{safe_sheet}.png"
                annotated_path = output_root / "docling" / f"{safe_sheet}.png"
                layout = self.renderer.render(
                    worksheet,
                    image_path,
                    input_data.max_rows,
                    input_data.max_columns,
                )
                try:
                    detected_boxes = self.extractor.detect(image_path)
                except ImportError as error:
                    raise ModuleExecutionError(
                        "Docling이 설치되지 않았습니다. requirements.txt 의존성을 설치하세요"
                    ) from error
                except Exception as error:
                    raise ModuleExecutionError(
                        f"Docling 테이블 탐지에 실패했습니다: {sheet_name}: {error}"
                    ) from error

                detections: List[DoclingTableRegionDTO] = []
                for box in detected_boxes:
                    normalized = self._normalize_bbox(box, layout.width, layout.height)
                    if normalized is None:
                        continue
                    bounds = bbox_to_cell_bounds(normalized, layout)
                    detections.append(
                        DoclingTableRegionDTO(
                            sheet_name=sheet_name,
                            table_index=len(detections) + 1,
                            excel_range=bounds.excel_range,
                            bbox_px=cell_bounds_bbox(bounds, layout),
                            cell_bounds=TableCellBoundsDTO(
                                min_row=bounds.min_row,
                                max_row=bounds.max_row,
                                min_column=bounds.min_column,
                                max_column=bounds.max_column,
                            ),
                        )
                    )
                _annotate_tables(image_path, annotated_path, detections)
                table_outputs.extend(item.model_dump(mode="json") for item in detections)
        finally:
            if workbook is not None:
                workbook.close()

        return {
            "file_name": workbook_path.name,
            "workbook_hash": current_hash,
            "sheet_names": input_data.sheet_names,
            "tables": table_outputs,
        }
