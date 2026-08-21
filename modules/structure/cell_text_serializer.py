from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Literal, Optional, cast

import openpyxl
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from backend.core.settings import PROCESSED_DATA_DIR
from backend.storage.spreadsheets.cell_visibility import worksheet_visible
from backend.storage.spreadsheets.structured_cell_text import (
    SERIALIZATION_VERSION,
    UNKNOWN_FIELD,
    WorksheetValueReader,
    canonical_sheet_name,
    ensure_period_header,
    generate_header_combinations,
    serialize_structured_cell,
    sheet_code,
)
from backend.storage.spreadsheets.workbook_catalog import WorkbookCatalog, WorkbookCatalogError
from modules.common.base_module import (
    ExecutableModule,
    ModuleConfigDTO,
    ModuleDefinition,
    ModuleDTO,
    ModuleExecutionError,
)
from modules.structure.spreadsheet_structure import (
    ClassifiedRegionDTO,
    ClassifiedTableDTO,
    SpreadsheetStructureOutput,
)


class CellTextSerializerInputDTO(SpreadsheetStructureOutput):
    """Classified workbook data received from a structure detector."""


class CellTextSerializerConfigDTO(ModuleConfigDTO):
    variant_mode: Literal["header_only", "header_with_value", "both"] = Field(
        default="both",
        description="생성할 검색 문서 변형 (header_only, header_with_value, 또는 both)",
        json_schema_extra={
            "enum": ["header_only", "header_with_value", "both"],
            "options": ["header_only", "header_with_value", "both"],
        },
    )


class CellTextSerializerExecutionDTO(
    CellTextSerializerInputDTO,
    CellTextSerializerConfigDTO,
):
    """Combined input and config DTO for execution."""


class CellTextDocumentDTO(ModuleDTO):
    cell_id: str = Field(description="시트 코드와 셀 좌표로 만든 검색 문서 ID")
    sheet_name: str = Field(description="정규화된 원본 시트 이름")
    cell_coord: str = Field(description="Excel 셀 좌표")
    row_header: List[str] = Field(description="상위 수준부터 수집한 행 헤더")
    column_header: List[str] = Field(description="상위 수준부터 수집한 열 헤더")
    cell_value: str = Field(description="Excel 데이터 셀의 표시 값")
    variant: Literal["header_only", "header_with_value"]
    text: str = Field(description="4필드 공통 포맷으로 직렬화한 검색 문서")


class CellTextSerializerOutput(ModuleDTO):
    file_name: str
    workbook_hash: str
    items: List[CellTextDocumentDTO]


class CellTextSerializerModule(ExecutableModule):
    definition = ModuleDefinition(
        type="cell_text_serializer",
        label="Structured Cell Text Serializer",
        category="Transform",
        description="분류된 Excel 셀을 Sheet·Row Header·Column Header·Cell Value 포맷으로 직렬화합니다.",
        inputs=["input"],
        outputs=["output"],
        config_fields=["variant_mode"],
        raw_output=True,
        version=SERIALIZATION_VERSION,
    )
    input_model = CellTextSerializerInputDTO
    config_model = CellTextSerializerConfigDTO
    execution_model = CellTextSerializerExecutionDTO
    output_model = CellTextSerializerOutput

    def __init__(
        self,
        catalog: WorkbookCatalog | None = None,
        processed_dir: Path = PROCESSED_DATA_DIR,
    ) -> None:
        self.catalog = catalog or WorkbookCatalog(processed_dir)

    @staticmethod
    def _region(
        table: ClassifiedTableDTO,
        region_type: str,
    ) -> Optional[ClassifiedRegionDTO]:
        return next(
            (region for region in table.regions if region.type == region_type),
            None,
        )

    @staticmethod
    def _header_values(
        reader: WorksheetValueReader,
        row: int,
        column: int,
        title_region: Optional[ClassifiedRegionDTO],
        row_region: Optional[ClassifiedRegionDTO],
        column_region: Optional[ClassifiedRegionDTO],
    ) -> tuple[List[str], List[str]]:
        row_headers: List[str] = []
        if title_region is not None:
            for title_row in range(title_region.rows[0], title_region.rows[1] + 1):
                if reader.row_hidden(title_row):
                    continue
                for title_column in range(
                    title_region.columns[0],
                    title_region.columns[1] + 1,
                ):
                    if reader.column_hidden(title_column):
                        continue
                    value = reader.value(title_row, title_column)
                    if value is not None and value not in row_headers:
                        row_headers.append(value)
        if row_region is not None:
            for header_column in range(
                row_region.columns[0],
                row_region.columns[1] + 1,
            ):
                if reader.column_hidden(header_column):
                    continue
                value = reader.value(row, header_column)
                if value is not None and value not in row_headers:
                    row_headers.append(value)

        column_headers: List[str] = []
        if column_region is not None:
            for header_row in range(
                column_region.rows[0],
                column_region.rows[1] + 1,
            ):
                if reader.row_hidden(header_row):
                    continue
                value = reader.value(header_row, column)
                if value is not None:
                    column_headers.append(value)
        return row_headers, ensure_period_header(column_headers)

    def _table_documents(
        self,
        reader: WorksheetValueReader,
        table: ClassifiedTableDTO,
        seen_cell_ids: set[str],
        variant_mode: str = "both",
    ) -> List[Dict[str, Any]]:
        """
        Generate serialized documents for the visible, non-empty cells in a classified table.
        
        Parameters:
            reader (WorksheetValueReader): Provides worksheet values and visibility information.
            table (ClassifiedTableDTO): Classified table whose data cells and headers are serialized.
            seen_cell_ids (set[str]): Cell identifiers already generated by other tables.
            variant_mode (str): Document variants to generate: ``"header_only"``,
                ``"header_with_value"``, or ``"both"``.
        
        Returns:
            List[Dict[str, Any]]: Serialized cell documents.
        
        Raises:
            ModuleExecutionError: If a cell identifier is generated more than once across tables.
        """
        data_region = self._region(table, "data")
        if data_region is None:
            return []
        row_region = self._region(table, "row_header")
        column_region = self._region(table, "column_header")
        title_region = self._region(table, "title")
        canonical_name = canonical_sheet_name(table.sheet_name)
        code = sheet_code(table.sheet_name)
        documents: List[Dict[str, Any]] = []

        for row in range(data_region.rows[0], data_region.rows[1] + 1):
            if reader.row_hidden(row):
                continue
            for column in range(
                data_region.columns[0],
                data_region.columns[1] + 1,
            ):
                if reader.column_hidden(column):
                    continue
                cell_value = reader.value(row, column)
                if cell_value is None:
                    continue
                row_headers, column_headers = self._header_values(
                    reader,
                    row,
                    column,
                    title_region,
                    row_region,
                    column_region,
                )
                if not row_headers:
                    continue

                cell_coord = f"{get_column_letter(column)}{row}"
                cell_id = f"{code} Cell {cell_coord}"
                if cell_id in seen_cell_ids:
                    raise ModuleExecutionError(
                        f"겹치는 테이블 영역에서 중복 셀이 생성되었습니다: {cell_id}"
                    )
                seen_cell_ids.add(cell_id)

                cell_seen_texts: set[str] = set()
                header_combos = generate_header_combinations(row_headers, column_headers)
                for r_combo, c_combo in header_combos:
                    common = {
                        "cell_id": cell_id,
                        "sheet_name": canonical_name,
                        "cell_coord": cell_coord,
                        "row_header": r_combo,
                        "column_header": c_combo,
                        "cell_value": cell_value,
                    }
                    if variant_mode == "both":
                        active_variants = [("header_only", UNKNOWN_FIELD), ("header_with_value", cell_value)]
                    elif variant_mode == "header_with_value":
                        active_variants = [("header_with_value", cell_value)]
                    else:
                        active_variants = [("header_only", UNKNOWN_FIELD)]

                    for variant, val in active_variants:
                        text = serialize_structured_cell(
                            canonical_name,
                            r_combo,
                            c_combo,
                            val,
                        )
                        if text not in cell_seen_texts:
                            cell_seen_texts.add(text)
                            documents.append(
                                {
                                    **common,
                                    "variant": variant,
                                    "text": text,
                                }
                            )
        return documents

    def execute(self, payload: BaseModel) -> Dict[str, Any]:
        """
        Serialize classified workbook cells into cell-text documents.
        
        Parameters:
            payload (BaseModel): Classification data and serialization configuration.
        
        Returns:
            Dict[str, Any]: Workbook filename, SHA-256 hash, and serialized cell documents.
        
        Raises:
            ModuleExecutionError: If the workbook cannot be resolved or read, has changed since classification, or references a missing or hidden sheet.
        """
        input_data = cast(CellTextSerializerInputDTO, payload)
        try:
            workbook_path = self.catalog.resolve(input_data.file_name)
            current_hash = self.catalog.sha256(workbook_path)
        except (OSError, ValueError, WorkbookCatalogError) as error:
            raise ModuleExecutionError(str(error)) from error
        if current_hash != input_data.workbook_hash:
            raise ModuleExecutionError(
                "영역 분류 이후 Excel 파일이 변경되었습니다. 앞 모듈부터 다시 실행하세요"
            )

        workbook = None
        try:
            workbook = openpyxl.load_workbook(
                workbook_path,
                read_only=False,
                data_only=True,
                keep_vba=workbook_path.suffix.lower() == ".xlsm",
            )
            items: List[Dict[str, Any]] = []
            seen_cell_ids: set[str] = set()
            readers: Dict[str, WorksheetValueReader] = {}
            total_tables = len(input_data.tables)
            self.report_progress(
                {
                    "phase": "serialization_tables",
                    "completed_tables": 0,
                    "total_tables": total_tables,
                    "completed_items": 0,
                }
            )
            for table_index, table in enumerate(input_data.tables, start=1):
                if table.sheet_name not in workbook.sheetnames:
                    raise ModuleExecutionError(
                        f"Excel 시트를 찾을 수 없습니다: {table.sheet_name}"
                    )
                reader = readers.get(table.sheet_name)
                if reader is None:
                    worksheet = workbook[table.sheet_name]
                    if not worksheet_visible(worksheet):
                        raise ModuleExecutionError(
                            f"숨겨진 Excel 시트는 직렬화할 수 없습니다: {table.sheet_name}"
                        )
                    reader = WorksheetValueReader(worksheet)
                    readers[table.sheet_name] = reader
                items.extend(
                    self._table_documents(
                        reader,
                        table,
                        seen_cell_ids,
                        variant_mode=getattr(input_data, "variant_mode", "both"),
                    )
                )
                self.report_progress(
                    {
                        "phase": "serialization_tables",
                        "completed_tables": table_index,
                        "total_tables": total_tables,
                        "completed_items": len(items),
                        "current_sheet": table.sheet_name,
                    }
                )
        finally:
            if workbook is not None:
                workbook.close()

        return {
            "file_name": workbook_path.name,
            "workbook_hash": current_hash,
            "items": items,
        }
